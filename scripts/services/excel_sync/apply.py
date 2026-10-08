# -*- coding: utf-8 -*-
import os
import sys
import re
import datetime
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple

SCRIPTS_DIR = Path(__file__).resolve().parents[2]
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))
from write_guard import ProductionWriteBlocked

try:
    import openpyxl
except ImportError:
    openpyxl = None

from services.excel_sync.identity import (
    QUARANTINE_CONFLICT,
    ROW_IDENTITY_FIELDS,
    ExcelApplyConflict,
    identity_matches,
)
from extractors.ledger.dates import LEDGER_COLUMN_ALIASES

from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from services.excel_sync.service import ExcelSyncService

    _HostBase_ExcelApplyMixin = ExcelSyncService
else:
    _HostBase_ExcelApplyMixin = object


class ExcelApplyMixin(_HostBase_ExcelApplyMixin):  # pyright: ignore[reportGeneralTypeIssues]  # static-only cycle; runtime base is object
    def find_master_excel(self) -> Optional[Path]:
        """마스터 대장 엑셀을 고른다. 탐색 규칙은 `extractors/ledger/discovery.py` 한 곳에만 있다.

        엑셀→DB 파싱(`load_request_ledger`)도 같은 함수를 쓰고, `sync_excel_to_db`는 여기서
        고른 경로를 파서에 그대로 넘긴다. (감사 R4-03)
        """
        from extractors.ledger.discovery import discover_master_excel
        chosen, candidates, pinned = discover_master_excel(self.base_dir)
        # 사본이 폴더에 생기면 수정 시각 최신본으로 정본이 조용히 바뀐다. 화면에서도 알린다. (R4 4.3)
        self._master_candidates = [c.name for c in candidates] if (chosen is not None and not pinned and len(candidates) > 1) else []
        if chosen is None:
            return None
        if not pinned and len(candidates) > 1:
            warn_key = tuple(str(c) for c in candidates)
            if self._master_warn_key != warn_key:
                self._master_warn_key = warn_key
                print(
                    f"⚠️ 대장 엑셀 후보가 {len(candidates)}건입니다. 가장 최근 수정본을 사용합니다: {candidates[0].name}\n"
                    "   고정하려면 config.json의 master_excel 항목에 파일명을 지정하세요."
                )
        return chosen

    def master_excel_ambiguity(self) -> List[str]:
        """`master_excel`로 고정하지 않았는데 대장 후보가 여러 개면 후보 파일명(최신순). 아니면 빈 목록."""
        self.find_master_excel()
        return list(getattr(self, "_master_candidates", []) or [])

    # -------------------------------------------------------------------------
    # 2. Excel Lock & Write Status Helpers
    # -------------------------------------------------------------------------
    def is_file_locked(self, filepath: Path) -> bool:
        """Checks if a file is currently locked by another application (e.g. MS Excel).

        **Windows 전용 의미론이다.** POSIX에서는 다른 프로세스가 파일을 열고 있어도
        자기 자신으로의 rename이 항상 성공하므로 항상 False가 되고, 보류 큐 설계
        전체가 무력화된다. 이 시스템은 Windows 업무 PC를 전제로 한다.
        """
        if not filepath.exists():
            return False
        if sys.platform != "win32":
            # 잠금을 검출할 수 없는 플랫폼에서는 쓰기를 시도하되, 실패는 호출부의
            # PermissionError 처리가 받아 보류 큐로 넘긴다.
            return False
        try:
            # On Windows, os.rename to itself is an atomic check for exclusive locks
            os.rename(filepath, filepath)
            return False
        except (IOError, PermissionError, OSError):
            return True

    # -------------------------------------------------------------------------
    # 3. Web -> Excel Write Operations
    # -------------------------------------------------------------------------
    def sync_item_to_excel(self, item: Dict[str, Any], action: str = "insert") -> Dict[str, Any]:
        """
        Applies an insert, update, or delete operation to the master Excel file.
        If the file is locked, enqueues the operation for deferred application.
        """
        if item.get("_operation_id"):
            return self._sync_durable_operation(item["_operation_id"])
        master_excel = self.find_master_excel()
        if not master_excel or not master_excel.exists():
            return {
                "excel_synced": False,
                "pending": False,
                "message": "마스터 엑셀 파일을 찾을 수 없어 DB에만 반영되었습니다."
            }

        # Check for file lock
        if self.is_file_locked(master_excel):
            with self._pending_lock:
                saved = self._enqueue_pending(item, action)
            return {
                "excel_synced": False,
                "pending": saved,
                "message": "현재 엑셀 파일이 열려 있어 저장 대기 중입니다. 엑셀을 닫으시면 자동 반영됩니다." if saved else "엑셀 파일이 잠겨 있고 대기 큐 저장에도 실패했습니다. 관리자에게 알리세요."
            }

        with self.lock:
            try:
                # 행 위치가 아니라 ID로 행을 찾을 수 있게 먼저 ID를 고정한다.
                # 스탬핑도 워크북을 통째로 다시 쓰므로 반드시 같은 락 안에서 한다.
                self._ensure_ledger_ids(master_excel)
                before_mtime = master_excel.stat().st_mtime
                success = self._apply_to_excel(master_excel, item, action)
                if success:
                    # Record write timestamp to avoid echo watcher trigger
                    self._mark_own_write(before_mtime, master_excel.stat().st_mtime)
                    self._record_history(item.get("ledger_id"), "EXCEL_APPLIED", item)
                    self.bump_version()
                    return {
                        "excel_synced": True,
                        "pending": False,
                        "message": f"✓ 마스터 엑셀({master_excel.name})에 즉시 반영되었습니다."
                    }
                # 반영하지 못했으면 반영했다고 말하지 않는다. 특히 삭제는 대상 행을
                # 못 찾아도 과거에는 성공으로 보고돼 사용자가 엑셀이 정리된 줄 알았다.
                if action == "delete":
                    return {
                        "excel_synced": False,
                        "pending": False,
                        "message": (
                            f"DB에서는 삭제했지만 마스터 엑셀({master_excel.name})에서 "
                            f"해당 대장ID 행을 찾지 못해 엑셀은 그대로입니다."
                        ),
                    }
                reason = "엑셀 반영 대상 시트를 찾지 못했습니다."
            except ExcelApplyConflict as conflict:
                # 같은 ID의 행이 다른 항목이다. 덮어쓰지도, 재시도하지도 않는다.
                self._quarantine_pending_operations(
                    [self._new_operation(item, action)], QUARANTINE_CONFLICT, str(conflict)
                )
                return {
                    "excel_synced": False,
                    "pending": False,
                    "conflict": True,
                    "message": (
                        f"엑셀의 같은 대장ID 행이 다른 요구자료라서 덮어쓰지 않았습니다. "
                        f"DB에만 반영됐으니 엑셀을 확인해 주세요. ({conflict})"
                    ),
                }
            except ProductionWriteBlocked as blocked:
                return {"excel_synced": False, "pending": False, "message": str(blocked)}
            except PermissionError:
                reason = None
            except Exception as e:
                reason = f"엑셀 반영 중 오류 발생: {e}"

            # 권한 오류(엑셀 사용 중)뿐 아니라 저장 검증 실패·디스크 오류도 큐에 남긴다.
            # 큐에 없으면 다음 엑셀→DB 동기화가 이 웹 변경을 엑셀의 옛 값으로 되돌린다.
            # (감사 R3-07)
            with self._pending_lock:
                saved = self._enqueue_pending(item, action)
            if reason is None:
                message = ("엑셀 파일이 사용 중입니다. 닫히는 즉시 자동 반영됩니다." if saved
                           else "엑셀 파일이 사용 중이고 대기 큐 저장에도 실패했습니다. 관리자에게 알리세요.")
            else:
                message = (f"{reason} 자동 재시도 대기열에 넣었습니다." if saved
                           else f"{reason} 대기 큐 저장에도 실패했습니다. 관리자에게 알리세요.")
            return {"excel_synced": False, "pending": saved, "message": message}

    # 헤더 별칭 표. 파서와 반드시 같은 표를 쓴다(정본: extractors/ledger/dates.py). (감사 R4-04)
    COLUMN_ALIASES = LEDGER_COLUMN_ALIASES

    def _save_workbook(self, wb, excel_path: Path):
        """마스터 엑셀을 원자적으로 저장하고 직전 본을 백업한다."""
        from extractors.ledger_parser import save_workbook_atomic
        save_workbook_atomic(wb, excel_path, base_dir=self.base_dir)

    @staticmethod
    def _build_header_map(ws) -> Dict[str, int]:
        header_map = {}
        for c in range(1, ws.max_column + 1):
            val = ws.cell(1, c).value
            if val:
                header_map[str(val).strip()] = c
        return header_map

    @classmethod
    def _resolve_columns(cls, ws) -> Dict[str, Optional[int]]:
        """시트의 1행 헤더를 표준 필드명 → 열번호로 바꾼다."""
        header_map = cls._build_header_map(ws)
        return {
            field: next((header_map[n] for n in names if n in header_map), None)
            for field, names in cls.COLUMN_ALIASES
        }

    def _find_row_by_ledger_id(self, wb, ledger_id: str) -> Optional[Tuple[str, int]]:
        """워크북 **전체 시트**에서 `대장ID` 행을 찾는다. (시트명, 행번호) 또는 None.

        AGENTS.md 관리대장 불변 조건 2. 대상 연도 시트만 뒤지면 항목의 연도가 바뀌었을 때
        원래 행을 찾지 못해 중복 ID 행이 생기고, 다음 ID 스탬핑이 원래 행에 새 ID를
        발급해 유령 항목을 만든다.
        """
        ledger_id = str(ledger_id or "").strip()
        if not ledger_id:
            return None
        for name in wb.sheetnames:
            ws = wb[name]
            c_id = self._resolve_columns(ws)["ledger_id"]
            if not c_id:
                continue
            for r in range(2, ws.max_row + 1):
                if str(ws.cell(r, c_id).value or "").strip() == ledger_id:
                    return name, r
        return None

    def _row_identity(self, wb, located: Tuple[str, int]) -> Dict[str, Any]:
        ws = wb[located[0]]
        cols = self._resolve_columns(ws)
        return {
            field: (ws.cell(located[1], cols[field]).value if cols.get(field) else None)
            for field in ROW_IDENTITY_FIELDS
        }

    def _verify_row_identity(self, wb, located: Tuple[str, int], item: Dict[str, Any], action: str):
        """같은 대장ID로 찾은 엑셀 행이 반영하려는 항목의 행인지 검사한다.

        - insert: 같은 ID 행이 이미 있으면 같은 항목(재시도)일 때만 허용한다.
        - update/delete: 수정 전 지문(`_expected`)이나 수정 후 값 중 하나와 맞아야 한다.
        비교할 값이 양쪽에 하나도 없으면 판단할 근거가 없으므로 허용한다.
        """
        row = self._row_identity(wb, located)
        if action == "insert":
            verdict = identity_matches(row, item)
            if verdict is False:
                raise ExcelApplyConflict(
                    f"{located[0]} 시트 {located[1]}행의 대장ID {item.get('ledger_id')}가 "
                    f"'{row.get('title') or row.get('requester') or '-'}' 항목에 이미 쓰이고 있습니다"
                )
            return
        verdicts = [identity_matches(row, item.get("_expected")), identity_matches(row, item)]
        if True in verdicts or all(v is None for v in verdicts):
            return
        raise ExcelApplyConflict(
            f"{located[0]} 시트 {located[1]}행({row.get('title') or row.get('requester') or '-'})이 "
            f"대장ID {item.get('ledger_id')}의 요구자료와 다릅니다"
        )

    def _create_year_sheet(self, wb, year_str: str, template_sheet_name: str):
        """새 연도 시트를 만들고 기준 시트의 헤더·열 너비를 복제한다."""
        from copy import copy as _copy

        src = wb[template_sheet_name]
        try:
            index = wb.sheetnames.index(template_sheet_name)
        except ValueError:
            index = 0
        ws = wb.create_sheet(title=year_str, index=index)
        for c in range(1, src.max_column + 1):
            src_cell = src.cell(1, c)
            new_cell = ws.cell(1, c, value=src_cell.value)
            new_cell.font = _copy(src_cell.font)
            new_cell.fill = _copy(src_cell.fill)
            new_cell.border = _copy(src_cell.border)
            new_cell.alignment = _copy(src_cell.alignment)
        for key, dim in src.column_dimensions.items():
            if dim.width:
                ws.column_dimensions[key].width = dim.width
        ws.row_dimensions[1].height = src.row_dimensions[1].height or 24
        print(f"✓ 마스터 엑셀에 신규 연도 시트를 생성했습니다: {year_str}")
        return ws

    def _apply_to_excel(self, excel_path: Path, item: Dict[str, Any], action: str) -> bool:
        """Internal worker that loads workbook, modifies target row, and saves.

        **호출자는 반드시 `self.lock`을 쥐고 있어야 한다.** openpyxl은 워크북 전체를
        읽어 전체를 다시 쓰므로, 두 스레드가 동시에 들어오면 나중에 저장한 쪽이
        먼저 저장한 쪽의 변경을 통째로 덮어쓴다.
        """
        if openpyxl is None:
            return False

        is_xlsm = excel_path.suffix.lower() == ".xlsm"
        wb = openpyxl.load_workbook(str(excel_path), keep_vba=is_xlsm)
        # 성공·실패·충돌 어느 경로든 워크북은 여기서 한 번만 닫는다. (감사 R4-15f)
        try:
            return self._apply_to_loaded_workbook(wb, excel_path, item, action)
        finally:
            try:
                wb.close()
            except Exception:
                pass

    def _apply_to_loaded_workbook(self, wb, excel_path: Path, item: Dict[str, Any], action: str) -> bool:
        from openpyxl.styles import Alignment, Border, Font, Side
        """`_apply_to_excel`의 본문. 워크북을 열고 닫는 일은 호출자가 한다."""
        target_seq = str(item.get("seq_no", "")).strip()
        target_title = str(item.get("title", "")).strip()
        target_ledger_id = str(item.get("ledger_id", "")).strip()
        target_requester = str(item.get("requester", "")).strip()

        # 1. `대장ID`는 워크북 **전체 시트**에서 먼저 찾는다. 항목의 연도가 바뀌어
        #    다른 연도 시트에 있어도 원래 행을 정확히 집는다. (AGENTS.md 불변 조건 2)
        located = self._find_row_by_ledger_id(wb, target_ledger_id)

        # 1-1. 찾은 행이 정말 이 항목의 행인지 확인한다. (감사 R3-02)
        #      보류 큐 작업은 적재 당시의 ID만 들고 있어서, 그사이 ID 스탬핑이 같은 ID를
        #      다른 행에 매겼거나 테스트·오염된 큐가 남아 있으면 실제 대장 행을 덮어쓴다.
        if located is not None:
            self._verify_row_identity(wb, located, item, action)

        # 2. 삭제는 행이 실제로 있는 시트에서 처리한다. 대상 연도 시트를 새로 만들
        #    필요가 없고, 못 찾았으면 반영하지 않았다고 정직하게 보고한다.
        if action == "delete":
            if located is None:
                return False
            ws = wb[located[0]]
            del_cols = self._resolve_columns(ws)
            del_row = located[1]
            if del_cols["status"]:
                ws.cell(del_row, del_cols["status"], value="[삭제]")
            if del_cols["note"]:
                old_note = str(ws.cell(del_row, del_cols["note"]).value or "")
                ws.cell(del_row, del_cols["note"], value=f"{old_note} (웹 삭제됨)".strip())
            self._save_workbook(wb, excel_path)
            return True

        year_str = str(item.get("year", datetime.datetime.now().year)).strip()
        target_sheet = None

        if year_str in wb.sheetnames:
            target_sheet = wb[year_str]
        else:
            year_sheets = [s for s in wb.sheetnames if re.fullmatch(r"20\d{2}", str(s).strip())]
            if year_sheets and re.fullmatch(r"20\d{2}", year_str):
                # 해당 연도 시트가 없으면 다른 연도 시트에 섞어 넣지 않고 새로 만든다.
                # 과거에는 첫 시트에 기록해 항목의 연도가 그 시트 연도로 바뀌었다.
                target_sheet = self._create_year_sheet(wb, year_str, max(year_sheets))
            elif wb.sheetnames:
                # Fallback to active or first non-search sheet
                candidates = [s for s in wb.sheetnames if "검색" not in s]
                target_sheet = wb[candidates[0]] if candidates else wb.active
                assert target_sheet is not None

        if target_sheet is None:
            return False

        ws = target_sheet
        cols = self._resolve_columns(ws)
        c_seq, c_party, c_req, c_aide = cols["seq_no"], cols["party"], cols["requester"], cols["aide"]
        c_title, c_detail = cols["title"], cols["details"]
        c_req_date, c_deadline, c_submit_date = cols["request_date"], cols["deadline"], cols["submit_date"]
        c_dept, c_status, c_note = cols["department"], cols["status"], cols["note"]
        c_type, c_id = cols["request_type"], cols["ledger_id"]

        target_row = None
        move_from = None  # 다른 연도 시트에서 옮겨 오는 경우 (시트명, 행번호)
        if located is not None:
            if located[0] == ws.title:
                target_row = located[1]
            else:
                # 연도가 바뀌었다. 새 시트에 전체 필드를 기록하고 원래 행은 제거한다.
                # 남겨 두면 중복 ID가 되고, 다음 ID 스탬핑이 원래 행에 새 ID를
                # 발급해 실재하지 않는 유령 항목을 만든다.
                move_from = located

        # 3. 보조 키 매칭은 ID를 확정할 수 없을 때만 쓴다. ID 열이 있으면 ID가
        #    비어 있는 행만 후보로 삼는다. insert는 ID로 못 찾으면 항상 새 행이다.
        if target_row is None and move_from is None and (action == "update" or not c_id):
            for r in range(2, ws.max_row + 1):
                if c_id and str(ws.cell(r, c_id).value or "").strip():
                    continue
                if c_seq and target_seq:
                    if str(ws.cell(r, c_seq).value or "").strip() == target_seq:
                        target_row = r
                        break
                if c_title and target_title:
                    curr_title = str(ws.cell(r, c_title).value or "").strip()
                    if curr_title == target_title:
                        if c_req and target_requester:
                            if str(ws.cell(r, c_req).value or "").strip() == target_requester:
                                target_row = r
                                break
                        else:
                            target_row = r
                            break

        thin_side = Side(style='thin', color='D3D3D3')
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        regular_font = Font(name="맑은 고딕", size=10)

        if action in ("insert", "update"):
            is_new_row = target_row is None
            if is_new_row:
                # Insert at bottom
                target_row = ws.max_row + 1
                # If row is empty, ensure correct row index
                while target_row > 2 and not any(ws.cell(target_row - 1, col).value for col in range(1, ws.max_column + 1)):
                    target_row -= 1

            # Auto-calculate sequence number if not given
            seq_val = item.get("seq_no")
            if not seq_val and c_seq:
                try:
                    prev_seq = 0
                    for r in range(2, target_row):
                        v = ws.cell(r, c_seq).value
                        if v and str(v).isdigit():
                            prev_seq = max(prev_seq, int(v))
                    seq_val = str(prev_seq + 1)
                except Exception:
                    seq_val = str(target_row - 1)

            def set_val(col_idx, val, align="left"):
                if not col_idx or val is None:
                    return
                c = ws.cell(row=target_row, column=col_idx, value=str(val))
                c.font = regular_font
                c.border = cell_border
                c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=(align == "left"))

            field_cols = [
                ("seq_no", c_seq, seq_val, "center"),
                ("party", c_party, item.get("party", ""), "center"),
                ("requester", c_req, item.get("requester", ""), "left"),
                ("aide", c_aide, item.get("aide", ""), "left"),
                ("title", c_title, item.get("title", ""), "left"),
                ("details", c_detail, item.get("details", ""), "left"),
                ("request_date", c_req_date, item.get("request_date", ""), "center"),
                ("deadline", c_deadline, item.get("deadline", ""), "center"),
                ("submit_date", c_submit_date, item.get("submit_date", ""), "center"),
                ("department", c_dept, item.get("department", ""), "left"),
                ("status", c_status, item.get("status", "제출"), "center"),
                ("note", c_note, item.get("note", ""), "left"),
                ("request_type", c_type, item.get("request_type", "시스템"), "center"),
            ]
            # 기존 행을 고칠 때만 "바뀐 필드"로 좁힌다. 보류 큐 재생 때 오래된 전체
            # 스냅숏으로 사용자의 엑셀 수기 수정을 덮어쓰지 않기 위한 장치다.
            # 반대로 **새 행을 만들 때는 반드시 전체 필드를 기록한다.** 일부만 쓰면
            # 의원·제목 칸이 빈 행이 되고, 다음 Excel→DB 동기화가 그 빈 값으로
            # DB의 요구자명을 지워 버린다.
            changed = item.get("_changed_fields")
            allowed = set(changed) if (action == "update" and changed and not is_new_row) else None
            for field, col, value, align in field_cols:
                if allowed is not None and field not in allowed:
                    continue
                set_val(col, value, align)
            set_val(c_id, item.get("ledger_id", ""), "center")

            ws.row_dimensions[target_row].height = 24

            # 다른 연도 시트에 있던 원래 행을 제거해 중복 `대장ID`를 남기지 않는다.
            if move_from is not None:
                try:
                    wb[move_from[0]].delete_rows(move_from[1])
                    print(f"✓ 관리대장 항목을 {move_from[0]} → {ws.title} 시트로 이동했습니다: {target_ledger_id}")
                except Exception as e:
                    print(f"[관리대장 이동 경고] 원래 행({move_from[0]}!{move_from[1]}) 제거 실패: {e}")

        self._save_workbook(wb, excel_path)
        return True
