# -*- coding: utf-8 -*-
import sys
import re
import sqlite3
import datetime
from pathlib import Path
from typing import List
import openpyxl

SCRIPTS_DIR = Path(__file__).resolve().parents[2]
DEFAULT_SYSTEM_DIR = SCRIPTS_DIR.parent
DEFAULT_DB_PATH = DEFAULT_SYSTEM_DIR / "data_requests.db"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))
from extractors.text_extractor import mask_pii
from extractors.ledger.paths import _get_active_system_dir, _get_active_db_path, DEFAULT_SYSTEM_DIR
from extractors.ledger.workbook import connect_readonly
from extractors.ledger.dates import (
    normalize_ledger_date,
    year_sheet_names,
    build_header_map,
    cell_text,
    is_ledger_data_row,
    LEDGER_ID_RE,
    ID_HEADER_NAMES,
    LEDGER_VALUE_FIELDS,
    LEDGER_COLUMN_ALIAS_MAP,
)
from extractors.ledger.discovery import discover_master_excel

def load_request_ledger(root_dir: Path, documents: list, merge_db_records: bool = True,
                        protected_ids=None, db_path=None, system_dir=None, excel_path=None) -> list:
    """Finds and parses *국회 요구자료 목록*.xlsm/.xlsx, merges DB-only records, and links records to documents.

    `protected_ids`: 아직 엑셀에 반영되지 않은 웹 변경의 대장 ID(보류 큐·격리 작업).
    이 ID들만 엑셀 값 대신 DB 값을 쓴다.
    `db_path`/`system_dir`: 병합할 DB와 설정 폴더. 생략하면 root_dir 기준으로 정한다.
    `excel_path`: 파싱할 대장 엑셀. 호출자가 이미 고른 파일이 있으면 반드시 넘긴다.
    쓰기와 읽기가 서로 다른 파일을 고르지 않게 하기 위해서다. (감사 R4-03)
    """
    protected = set(protected_ids or ())

    # 0. Check existing SQLite database to preserve user-added or modified ledger items
    existing_user_records = {}
    if db_path is not None:
        active_db = Path(db_path)
    elif (Path(root_dir) / "data_requests.db").exists():
        active_db = Path(root_dir) / "data_requests.db"
    else:
        active_db = _get_active_db_path(root_dir)
    if merge_db_records and active_db.exists():
        tmp_conn = None
        try:
            # 읽기 전용으로 연다. 병합용 조회가 DB 파일을 만들거나 체크포인트로 바꾸지 않게 한다.
            tmp_conn = connect_readonly(active_db)
            tmp_conn.row_factory = sqlite3.Row
            tmp_cur = tmp_conn.cursor()
            tbl_chk = tmp_cur.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='request_ledger'").fetchone()
            if tbl_chk:
                for r in tmp_cur.execute("SELECT * FROM request_ledger").fetchall():
                    existing_user_records[r["ledger_id"]] = dict(r)
        except Exception as e:
            print(f"기존 대장 DB 조회 알림: {e}")
        finally:
            if tmp_conn:
                tmp_conn.close()

    root_dir = Path(root_dir)
    sys_dir = Path(system_dir) if system_dir is not None else _get_active_system_dir(root_dir)

    if excel_path is not None and Path(excel_path).exists():
        target_excel = Path(excel_path)
    else:
        target_excel, candidates, pinned = discover_master_excel(sys_dir, root_dir)
        if target_excel is None:
            print("알림: 요구자료 목록 엑셀 파일을 찾을 수 없습니다. (신규 등록 및 공문서 데이터만 반영)")
            return list(existing_user_records.values())
        if not pinned and len(candidates) > 1:
            print(
                f"⚠️ 대장 엑셀 후보가 {len(candidates)}건입니다. 가장 최근 수정본을 사용합니다: {target_excel.name}\n"
                "   고정하려면 config.json의 master_excel 항목에 파일명을 지정하세요."
            )
    print(f"국회 요구자료 마스터 목록 파싱 시작: {target_excel.name}")

    # Build doc index for linking
    doc_map_by_num = {}
    doc_map_by_req_date = {}
    for d in documents:
        did = d.get("doc_id")
        yr = d.get("year", "")
        dnum = d.get("doc_number", "")
        req = d.get("requester", "")
        req_date = d.get("request_date", "")
        if dnum:
            clean_num = re.sub(r'\D', '', str(dnum)).lstrip('0')
            if clean_num:
                doc_map_by_num[(yr, clean_num)] = did
        if req and req_date:
            req_clean = re.sub(r'\s*(의원|처|실|팀|관)?$', '', req).strip()
            doc_map_by_req_date[(yr, req_clean, req_date)] = did

    wb = openpyxl.load_workbook(str(target_excel), data_only=True)
    all_sheet_names = wb.sheetnames
    # Do not cap the ledger at a hard-coded calendar year.  Parse every
    # year-named sheet (newest first); if none are named by year, preserve the
    # legacy fallback of parsing all sheets.
    target_sheets = year_sheet_names(all_sheet_names)

    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ledger_items = []

    # Track existing doc links
    doc_id_to_ledger = {}
    seen_ids = set()
    skipped_unstamped = 0
    skipped_duplicate = 0

    for y in target_sheets:
        ws = wb[y]
        header_map = build_header_map(ws)

        # 명시 ID가 하나라도 있는 시트에서는 위치 기반 ID를 쓰지 않는다.
        # 위치 기반 번호는 기존 명시 ID와 충돌하기 때문이다.
        sheet_has_explicit_id = any(
            LEDGER_ID_RE.match(cell_text(ws, r, header_map, ID_HEADER_NAMES) or "")
            for r in range(2, ws.max_row + 1)
            if is_ledger_data_row(ws, r, header_map)
        )

        row_seq = 0
        for r in range(2, ws.max_row + 1):
            def gv(*names):
                for name in names:
                    c = header_map.get(name)
                    if c:
                        v = ws.cell(r, c).value
                        if v is not None and str(v).strip():
                            return str(v).strip()
                return ""
            
            A = LEDGER_COLUMN_ALIAS_MAP
            seq = gv(*A["seq_no"])
            title = gv(*A["title"])
            detail = gv(*A["details"])
            party = gv(*A["party"])
            req = gv(*A["requester"])
            if not any([seq, title, detail, req]):
                continue
            
            c_id = gv(*ID_HEADER_NAMES)
            if c_id and LEDGER_ID_RE.match(c_id):
                if c_id in seen_ids:
                    # 행을 복사하면 ID까지 복제된다. 중복 ID는 스탬핑이 새 ID를 줄 때까지 건너뛴다.
                    skipped_duplicate += 1
                    continue
                ledger_id = c_id
            elif sheet_has_explicit_id:
                # 아직 ID가 스탬핑되지 않은 신규 수기 행. 다음 스탬핑에서 ID를 받는다.
                skipped_unstamped += 1
                continue
            else:
                row_seq += 1
                ledger_id = f"REQ-{y}-{row_seq:03d}"
                if ledger_id in seen_ids:
                    skipped_duplicate += 1
                    continue
            seen_ids.add(ledger_id)


            req_d = normalize_ledger_date(gv(*A["request_date"]), y)
            dline = normalize_ledger_date(gv(*A["deadline"]), y)
            sub_d = normalize_ledger_date(gv(*A["submit_date"]), y)
            
            # Linking logic
            linked_did = ""
            if seq:
                clean_seq = re.sub(r'\D', '', str(seq)).lstrip('0')
                if clean_seq and (y, clean_seq) in doc_map_by_num:
                    linked_did = doc_map_by_num[(y, clean_seq)]
            if not linked_did and req and req_d:
                req_clean = re.sub(r'\s*(의원|처|실|팀|관)?$', '', req).strip()
                if (y, req_clean, req_d) in doc_map_by_req_date:
                    linked_did = doc_map_by_req_date[(y, req_clean, req_d)]

            status_val = gv(*A["status"])
            if not status_val:
                # 운영 대장은 제출한 행에 '제출'을 직접 적는다(2019~2025 시트 전 행). 칸이 비어 있고
                # 제출일도 없는 행은 아직 진행 중인 요구자료다. 예전에는 이것도 '제출'로 채워서
                # 마감 경고·미제출 필터에서 빠졌다.
                status_val = "완료" if sub_d else "미제출"

            item = {
                "ledger_id": ledger_id,
                "year": y,
                "seq_no": seq,
                "party": party,
                "requester": req,
                "aide": mask_pii(gv(*A["aide"])),
                "title": title,
                "details": mask_pii(detail),
                "request_date": req_d,
                "deadline": dline,
                "submit_date": sub_d,
                "department": gv(*A["department"]),
                "status": status_val,
                "note": mask_pii(gv(*A["note"])),
                "request_type": gv(*A["request_type"]),
                "linked_doc_id": linked_did,
                "created_at": now_str,
                "updated_at": now_str
            }

            # 필드 값의 정본은 마스터 엑셀이다(README 「저장소별 정본」). DB 값을 쓰는 경우는
            # 둘뿐이다: 아직 엑셀에 반영되지 못한 웹 변경(protected_ids)과 삭제 tombstone.
            # 예전에는 "updated_at != created_at"이면 DB 값을 우선해, 서버가 꺼진 동안의
            # 엑셀 수기 수정이 재빌드 때마다 조용히 버려졌다. (감사 R3-06)
            if ledger_id in existing_user_records:
                old = existing_user_records[ledger_id]
                db_wins = (
                    ledger_id in protected
                    or str(old.get("status") or "").strip() == "[삭제]"
                )
                if db_wins:
                    for k in LEDGER_VALUE_FIELDS:
                        if k in old and old[k] is not None:
                            item[k] = old[k]
                changed = any(
                    str(item.get(k) or "").strip() != str(old.get(k) or "").strip()
                    for k in LEDGER_VALUE_FIELDS
                )
                item["created_at"] = old.get("created_at") or item["created_at"]
                if not changed and old.get("updated_at"):
                    item["updated_at"] = old["updated_at"]
                if old.get("linked_doc_id"):
                    item["linked_doc_id"] = old["linked_doc_id"]
                    linked_did = old["linked_doc_id"]

            if linked_did:
                doc_id_to_ledger[linked_did] = ledger_id

            ledger_items.append(item)

    try:
        wb.close()
    except Exception:
        pass

    if skipped_unstamped:
        print(
            f"⚠️ 대장ID가 비어 있는 신규 행 {skipped_unstamped}건은 이번 파싱에서 제외했습니다. "
            "엑셀 잠금이 풀린 뒤 ID 고정(스탬핑)이 완료되면 자동 반영됩니다."
        )
    if skipped_duplicate:
        print(f"⚠️ 중복된 대장ID {skipped_duplicate}건을 건너뛰었습니다. 복사된 행의 ID를 비워 두면 새 ID가 부여됩니다.")

    # Merge custom user-created records not present in Excel (e.g. registered via web UI or API)
    excel_ids = {item["ledger_id"] for item in ledger_items}
    custom_added_count = 0
    for lid, old in existing_user_records.items():
        if lid not in excel_ids:
            # Re-link if not yet linked
            if not old.get("linked_doc_id"):
                seq = old.get("seq_no", "")
                y = str(old.get("year") or "")
                req = old.get("requester", "")
                req_d = old.get("request_date", "")
                clean_seq = re.sub(r'\D', '', str(seq)).lstrip('0') if seq else ""
                if clean_seq and (y, clean_seq) in doc_map_by_num:
                    old["linked_doc_id"] = doc_map_by_num[(y, clean_seq)]
                elif req and req_d:
                    req_clean = re.sub(r'\s*(의원|처|실|팀|관)?$', '', req).strip()
                    if (y, req_clean, req_d) in doc_map_by_req_date:
                        old["linked_doc_id"] = doc_map_by_req_date[(y, req_clean, req_d)]
            if old.get("linked_doc_id"):
                doc_id_to_ledger[old["linked_doc_id"]] = lid
            ledger_items.append(old)
            custom_added_count += 1

    if custom_added_count > 0:
        print(f"✓ 기존 웹 관리화면/DB에서 등록된 신규 요구자료 {custom_added_count}건 보존 및 병합 완료!")

    # Update linked_ledger_id on documents
    for d in documents:
        did = d.get("doc_id")
        if did in doc_id_to_ledger:
            d["linked_ledger_id"] = doc_id_to_ledger[did]
        else:
            d["linked_ledger_id"] = ""

    print(f"국회 요구자료 관리대장 파싱 완료: 총 {len(ledger_items)}건 (답변문서 연동: {len(doc_id_to_ledger)}건)")
    return ledger_items

