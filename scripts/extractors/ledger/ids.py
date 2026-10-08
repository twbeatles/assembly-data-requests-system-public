# -*- coding: utf-8 -*-
import sys
import re
from pathlib import Path
from typing import Any, List, cast
import openpyxl

SCRIPTS_DIR = Path(__file__).resolve().parents[2]
DEFAULT_SYSTEM_DIR = SCRIPTS_DIR.parent
DEFAULT_DB_PATH = DEFAULT_SYSTEM_DIR / "data_requests.db"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))
from extractors.ledger.workbook import save_workbook_atomic
from extractors.ledger.dates import (
    year_sheet_names, build_header_map, find_id_column, cell_text, is_ledger_data_row,
    LEDGER_ID_RE, ID_HEADER_NAMES, TITLE_HEADER_NAMES, REQUESTER_HEADER_NAMES,
)

def _id_year_and_number(ledger_id: str):
    m = re.match(r'^REQ-(\w+)-(\d+)$', str(ledger_id or "").strip())
    return (m.group(1), int(m.group(2))) if m else None


def plan_ledger_ids(wb_vals, db_records=None):
    """`대장ID`가 비어 있는 행에 부여할 안정 ID 계획을 만든다.

    반환값은 (assignments, sheets_without_id_column)이다. assignments는
    (시트명, 행번호, ledger_id) 목록이며 기존 명시 ID는 건드리지 않는다.
    위치 기반 ID는 행 삭제·삽입 때마다 밀리므로, 한 번 계산한 ID를 워크북에
    기록해 고정하는 것이 이 함수의 목적이다.
    """
    db_by_id = {}
    db_by_content = {}
    tombstones = set()
    max_num = {}
    for rec in (db_records or []):
        lid = str(rec.get("ledger_id") or "").strip()
        if not lid:
            continue
        db_by_id[lid] = rec
        if str(rec.get("status") or "").strip() == "[삭제]":
            tombstones.add(lid)
        key = (
            str(rec.get("year") or "").strip(),
            str(rec.get("title") or "").strip(),
            str(rec.get("requester") or "").strip(),
        )
        db_by_content.setdefault(key, []).append(lid)
        parsed = _id_year_and_number(lid)
        if parsed:
            max_num[parsed[0]] = max(max_num.get(parsed[0], 0), parsed[1])

    assignments = []
    sheets_without_id_column = set()
    used = set()

    def _remember(lid: str):
        used.add(lid)
        parsed = _id_year_and_number(lid)
        if parsed:
            max_num[parsed[0]] = max(max_num.get(parsed[0], 0), parsed[1])

    # 1차(워크북 전체): 유효하고 중복되지 않은 명시 ID를 **모든 시트에서** 먼저 확정한다.
    # 시트마다 따로 모으면, 아직 순회하지 않은 시트(다른 연도로 옮긴 항목)의 ID를 새 행에
    # 다시 발급해 두 항목의 정체성이 뒤바뀐다. (감사 R3-05)
    sheets = []
    for sheet in year_sheet_names(wb_vals.sheetnames):
        ws = wb_vals[sheet]
        header_map = build_header_map(ws)
        if find_id_column(header_map) is None:
            sheets_without_id_column.add(sheet)
        data_rows = [r for r in range(2, ws.max_row + 1) if is_ledger_data_row(ws, r, header_map)]
        explicit_rows = set()
        for r in data_rows:
            raw = cell_text(ws, r, header_map, ID_HEADER_NAMES)
            if raw and LEDGER_ID_RE.match(raw) and raw not in used:
                _remember(raw)
                explicit_rows.add(r)
        sheets.append((sheet, ws, header_map, data_rows, explicit_rows))

    # 2차: 비어 있거나 중복인 ID 행에 ID를 배정한다.
    for sheet, ws, header_map, data_rows, explicit_rows in sheets:
        sheet_has_explicit = bool(explicit_rows)
        row_seq = 0
        for r in data_rows:
            if r in explicit_rows:
                continue
            row_seq += 1
            candidate = f"REQ-{sheet}-{row_seq:03d}"
            chosen = None
            row_title = cell_text(ws, r, header_map, TITLE_HEADER_NAMES)
            row_requester = cell_text(ws, r, header_map, REQUESTER_HEADER_NAMES)

            # (1) 명시 ID가 없는 레거시 시트에서 기존 위치 기반 매핑이 DB에 남아 있으면 유지한다.
            #     명시 ID가 있는 시트에서는 위치 순번이 원래 위치와 무관하므로, DB 레코드의
            #     제목·요구자가 이 행과 같고 삭제되지 않은 경우에만 물려준다.
            candidate_ok = candidate in db_by_id and candidate not in used
            if candidate_ok and sheet_has_explicit:
                rec = db_by_id[candidate]
                candidate_ok = (
                    candidate not in tombstones
                    and str(rec.get("title") or "").strip() == row_title
                    and str(rec.get("requester") or "").strip() == row_requester
                )
            if candidate_ok:
                chosen = candidate
            else:
                # (2) 같은 연도·제목·요구자의 DB 레코드가 있으면 그 ID를 물려준다.
                key = (str(sheet).strip(), row_title, row_requester)
                for lid in db_by_content.get(key, []):
                    if lid not in used and lid not in tombstones:
                        chosen = lid
                        break
            # (3) 명시 ID가 전혀 없는 레거시 시트는 위치 기반 ID를 그대로 고정한다.
            if chosen is None and not sheet_has_explicit and candidate not in used:
                chosen = candidate
            # (4) 그 밖에는 해당 연도의 다음 번호를 새로 발급한다.
            if chosen is None:
                token = str(sheet).strip()
                n = max_num.get(token, 0) + 1
                while f"REQ-{token}-{n:03d}" in used:
                    n += 1
                chosen = f"REQ-{token}-{n:03d}"

            _remember(chosen)
            assignments.append((sheet, r, chosen))

    return assignments, sheets_without_id_column


def stamp_ledger_ids(excel_path, db_records=None, base_dir=None) -> int:
    """마스터 엑셀의 빈 `대장ID` 칸을 채워 ID를 고정한다. 기록한 칸 수를 돌려준다.

    파일이 잠겨 있거나 권한이 없으면 예외가 올라가므로 호출자가 처리한다.
    `base_dir`: 백업(`.ledger_backup`)을 남길 시스템 폴더. 웹 반영 경로와 같은 곳에 백업이
    모이도록 넘긴다. 생략하면 엑셀 파일이 있는 폴더를 쓴다. (감사 R4-13)
    """
    excel_path = Path(excel_path)
    wb_vals = openpyxl.load_workbook(str(excel_path), data_only=True)
    try:
        assignments, missing_cols = plan_ledger_ids(wb_vals, db_records)
    finally:
        try:
            wb_vals.close()
        except Exception:
            pass

    if not assignments:
        return 0

    from copy import copy as _copy

    is_xlsm = excel_path.suffix.lower() == ".xlsm"
    wb = openpyxl.load_workbook(str(excel_path), keep_vba=is_xlsm)
    try:
        id_cols = {}
        for sheet, row, ledger_id in assignments:
            ws = wb[sheet]
            if sheet not in id_cols:
                header_map = build_header_map(ws)
                col = find_id_column(header_map)
                if col is None:
                    col = ws.max_column + 1
                    header_cell = ws.cell(1, col, value="대장ID")
                    if col > 1:
                        src = ws.cell(1, col - 1)
                        header_cell.font = cast(Any, _copy(src.font))
                        header_cell.fill = cast(Any, _copy(src.fill))
                        header_cell.border = cast(Any, _copy(src.border))
                        header_cell.alignment = cast(Any, _copy(src.alignment))
                id_cols[sheet] = col
            ws.cell(row, id_cols[sheet], value=ledger_id)
        save_workbook_atomic(wb, excel_path, base_dir=base_dir)
    finally:
        try:
            wb.close()
        except Exception:
            pass

    if missing_cols:
        print(f"✓ 마스터 엑셀에 '대장ID' 열을 추가했습니다: {', '.join(sorted(missing_cols))}")
    print(f"✓ 관리대장 ID 고정(스탬핑) 완료: {len(assignments)}건")
    return len(assignments)


LEDGER_VALUE_FIELDS = (
    "year", "seq_no", "party", "requester", "aide", "title", "details", "request_date",
    "deadline", "submit_date", "department", "status", "note", "request_type",
)
