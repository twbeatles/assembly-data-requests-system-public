# -*- coding: utf-8 -*-
"""엑셀 시트 빌더(SRP: 5개 시트 생성)."""
from collections import Counter
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from excel_report.styles import (
    title_font,
    section_font,
    regular_font,
    link_font,
    sub_header_font,
    sub_header_fill,
    accent_fill,
    zebra_fill,
    cell_border,
    apply_table_header,
)
from excel_report.text import _excel_body, sanitize_text
from excel_report.links import resolve_excel_hyperlink_path


def build_dashboard_sheet(wb, records, qa_items, ledger_items=None):
    import generate_excel_db as ge  # 지연 import: 진입점 전역 패치를 그대로 본다.
    ws = wb.active
    assert ws is not None
    ws.title = "📊 대시보드"
    ws.views.sheetView[0].showGridLines = True

    dept_name = ge.CFG.get("department_name", "자료요구담당부서")
    sys_title = ge.CFG.get("system_title", "국회·대외기관 자료요구 스마트 관리 및 검색 시스템")
    ws["B2"] = f"{dept_name} {sys_title}"
    ws["B2"].font = title_font
    ledger_count = len(ledger_items) if ledger_items else 0
    ws["B3"] = f"총 관리 문서: {len(records)}건 | 세부 Q&A: {len(qa_items)}건 | 관리대장: {ledger_count}건 | 답변 전문(Full Text) 탑재 완료"
    ws["B3"].font = Font(name="맑은 고딕", size=11, color="555555", italic=True)

    kpis = [
        ("총 관리 문서", f"{len(records)} 건", "B", "C"),
        ("세부 질의(Q&A)", f"{len(qa_items)} 건", "E", "F"),
        ("국회 관리대장", f"{ledger_count} 건", "H", "I"),
        ("2026년 문서", f"{sum(1 for r in records if r.get('year')=='2026')} 건", "K", "L"),
        ("통계/표 포함 문서", f"{sum(1 for r in records if r.get('has_tables')=='Y')} 건", "N", "O")
    ]

    for label, val, start_col, end_col in kpis:
        ws.merge_cells(f"{start_col}5:{end_col}5")
        ws.merge_cells(f"{start_col}6:{end_col}6")
        c5 = ws[f"{start_col}5"]
        c5.value = label
        c5.font = Font(name="맑은 고딕", size=10, bold=True, color="555555")
        c5.alignment = Alignment(horizontal="center", vertical="center")
        c5.fill = accent_fill
        
        c6 = ws[f"{start_col}6"]
        c6.value = val
        c6.font = Font(name="맑은 고딕", size=16, bold=True, color="1F4E79")
        c6.alignment = Alignment(horizontal="center", vertical="center")
        c6.fill = zebra_fill

        for r in (5, 6):
            ws[f"{start_col}{r}"].border = cell_border
            ws[f"{end_col}{r}"].border = cell_border

    # Requester rankings
    ws["B9"] = "■ 주요 요구기관 및 국회의원별 현황"
    ws["B9"].font = section_font
    
    req_counter = Counter(r.get('requester', '') for r in records)
    headers_req = ["순위", "요구기관 / 요구자", "문서 건수", "비중(%)"]
    for c_idx, h in enumerate(headers_req, start=2):
        cell = ws.cell(row=10, column=c_idx, value=h)
        cell.font = sub_header_font
        cell.fill = sub_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border
    
    total_recs = len(records) or 1
    for idx, (req, cnt) in enumerate(req_counter.most_common(12), start=1):
        r_idx = 10 + idx
        ws.cell(row=r_idx, column=2, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=3, value=sanitize_text(req)).alignment = Alignment(horizontal="left")
        ws.cell(row=r_idx, column=4, value=cnt).alignment = Alignment(horizontal="right")
        ws.cell(row=r_idx, column=5, value=f"{cnt/total_recs*100:.1f}%").alignment = Alignment(horizontal="right")
        for c in range(2, 6):
            ws.cell(row=r_idx, column=c).border = cell_border
            ws.cell(row=r_idx, column=c).font = regular_font

    # Topic distributions
    ws["G9"] = "■ 핵심 질의 주제별 문서 분포"
    ws["G9"].font = section_font

    topic_counter = Counter()
    for r in records:
        if r.get('topic_tags'):
            for t in r['topic_tags'].split(','):
                topic_counter[t.strip()] += 1

    headers_top = ["순위", "핵심 주제(키워드)", "관련 문서수", "비중(%)"]
    for c_idx, h in enumerate(headers_top, start=7):
        cell = ws.cell(row=10, column=c_idx, value=h)
        cell.font = sub_header_font
        cell.fill = sub_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border

    for idx, (top, cnt) in enumerate(topic_counter.most_common(12), start=1):
        r_idx = 10 + idx
        ws.cell(row=r_idx, column=7, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=8, value=sanitize_text(top)).alignment = Alignment(horizontal="left")
        ws.cell(row=r_idx, column=9, value=cnt).alignment = Alignment(horizontal="right")
        ws.cell(row=r_idx, column=10, value=f"{cnt/total_recs*100:.1f}%").alignment = Alignment(horizontal="right")
        for c in range(7, 11):
            ws.cell(row=r_idx, column=c).border = cell_border
            ws.cell(row=r_idx, column=c).font = regular_font

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 4
    ws.column_dimensions["G"].width = 8
    ws.column_dimensions["H"].width = 22
    ws.column_dimensions["I"].width = 14
    ws.column_dimensions["J"].width = 12
    ws.column_dimensions["K"].width = 12
    ws.column_dimensions["L"].width = 12
    ws.column_dimensions["M"].width = 4
    ws.column_dimensions["N"].width = 12
    ws.column_dimensions["O"].width = 12
def build_master_list_sheet(wb, records):
    ws = wb.create_sheet(title="📋 자료요구_통합목록")
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "연번", "문서ID", "작성/요구일", "연도", "기관", "요구자", "관리번호",
        "자료/질의 제목", "세부 질의 목록", "답변 요약", "답변 전문 (Full Text)", "담당부서/팀장",
        "표포함", "표개수", "표요약", "주제태그", "버전", "연계 대장ID", "마크다운 열기", "원본파일 열기"
    ]
    apply_table_header(ws, 1, headers)

    for idx, r in enumerate(records, start=1):
        row = idx + 1
        is_even = (idx % 2 == 0)
        curr_fill = zebra_fill if is_even else None

        ws.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=2, value=sanitize_text(r.get("doc_id", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=3, value=sanitize_text(r.get("request_date", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=4, value=sanitize_text(r.get("year", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=5, value=sanitize_text(r.get("institution", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=6, value=sanitize_text(r.get("requester", ""))).alignment = Alignment(horizontal="left", vertical="top")
        ws.cell(row=row, column=7, value=sanitize_text(r.get("doc_number", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=8, value=sanitize_text(r.get("title", ""))).alignment = Alignment(horizontal="left", vertical="top")
        ws.cell(row=row, column=9, value=sanitize_text(r.get("question_list", ""))).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=row, column=10, value=sanitize_text(r.get("answer_summary", ""))).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        # Full text cell
        # 엑셀 셀 한도(3만 자)는 **내보낼 때** 적용한다. 예전에는 같은 값을 DB 컬럼
        # answer_full_text에도 저장해 본문을 두 벌 들고 있었다(제안서 D6).
        ws.cell(row=row, column=11, value=sanitize_text(_excel_body(r, "full_markdown", "answer_full_text"))).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=row, column=12, value=sanitize_text(r.get("contact_person") or r.get("department", ""))).alignment = Alignment(horizontal="left", vertical="top")
        ws.cell(row=row, column=13, value=sanitize_text(r.get("has_tables", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=14, value=r.get("table_count", 0)).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=row, column=15, value=sanitize_text(r.get("table_summary", ""))).alignment = Alignment(horizontal="left", vertical="top")
        ws.cell(row=row, column=16, value=sanitize_text(r.get("topic_tags", ""))).alignment = Alignment(horizontal="left", vertical="top")
        ws.cell(row=row, column=17, value=sanitize_text(r.get("version", ""))).alignment = Alignment(horizontal="center", vertical="top")

        linked_ledger = r.get("linked_ledger_id", "")
        ledger_cell = ws.cell(row=row, column=18)
        ledger_cell.alignment = Alignment(horizontal="center", vertical="top")
        if linked_ledger:
            ledger_cell.value = f'=HYPERLINK("#\'📋 국회_요구자료_관리대장\'!A1", "{sanitize_text(linked_ledger)}")'
            ledger_cell.font = link_font
        else:
            ledger_cell.value = "-"

        md_cell = ws.cell(row=row, column=19)
        md_path_norm = resolve_excel_hyperlink_path(r.get("parsed_md_path", ""))
        md_cell.value = f'=HYPERLINK("{md_path_norm}", "마크다운 보기")'
        md_cell.font = link_font
        md_cell.alignment = Alignment(horizontal="center", vertical="top")

        orig_cell = ws.cell(row=row, column=20)
        orig_path_norm = resolve_excel_hyperlink_path(r.get("original_path", ""))
        orig_cell.value = f'=HYPERLINK("{orig_path_norm}", "원본 열기")'
        orig_cell.font = link_font
        orig_cell.alignment = Alignment(horizontal="center", vertical="top")

        for col_idx in range(1, 21):
            cell = ws.cell(row=row, column=col_idx)
            cell.border = cell_border
            if col_idx not in (18, 19, 20):
                cell.font = regular_font
            elif col_idx == 18 and not linked_ledger:
                cell.font = regular_font
            if curr_fill:
                cell.fill = curr_fill

    col_widths = {
        1: 6, 2: 15, 3: 13, 4: 8, 5: 14, 6: 16, 7: 11,
        8: 32, 9: 35, 10: 35, 11: 55, 12: 28, 13: 8, 14: 8, 15: 25,
        16: 22, 17: 10, 18: 15, 19: 15, 20: 15
    }
    for c_idx, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = w

    ws.auto_filter.ref = f"A1:T{len(records)+1}"
    ws.freeze_panes = "A2"
def build_qa_detail_sheet(wb, qa_items):
    ws = wb.create_sheet(title="📑 Q&A_상세DB")
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "연번", "Q&A ID", "문서ID", "작성/요구일", "연도", "기관", "요구자",
        "관리번호", "질의번호", "개별 질의 제목", "해당 질의 답변 전문 (Full Text)", "표포함",
        "마크다운 열기", "원본파일 열기"
    ]
    apply_table_header(ws, 1, headers)

    for idx, q in enumerate(qa_items, start=1):
        row = idx + 1
        is_even = (idx % 2 == 0)
        curr_fill = zebra_fill if is_even else None

        ws.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=2, value=sanitize_text(q.get("qa_id", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=3, value=sanitize_text(q.get("doc_id", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=4, value=sanitize_text(q.get("request_date", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=5, value=sanitize_text(q.get("year", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=6, value=sanitize_text(q.get("institution", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=7, value=sanitize_text(q.get("requester", ""))).alignment = Alignment(horizontal="left", vertical="top")
        ws.cell(row=row, column=8, value=sanitize_text(q.get("doc_number", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=9, value=sanitize_text(q.get("q_num", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=10, value=sanitize_text(q.get("question_title", ""))).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=row, column=11, value=sanitize_text(_excel_body(q, "answer_markdown", "answer_full"))).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=row, column=12, value=sanitize_text(q.get("has_tables", ""))).alignment = Alignment(horizontal="center", vertical="top")

        md_cell = ws.cell(row=row, column=13)
        md_path_norm = resolve_excel_hyperlink_path(q.get("parsed_md_path", ""))
        md_cell.value = f'=HYPERLINK("{md_path_norm}", "마크다운 보기")'
        md_cell.font = link_font
        md_cell.alignment = Alignment(horizontal="center", vertical="top")

        orig_cell = ws.cell(row=row, column=14)
        orig_path_norm = resolve_excel_hyperlink_path(q.get("original_path", ""))
        orig_cell.value = f'=HYPERLINK("{orig_path_norm}", "원본 열기")'
        orig_cell.font = link_font
        orig_cell.alignment = Alignment(horizontal="center", vertical="top")

        for col_idx in range(1, 15):
            cell = ws.cell(row=row, column=col_idx)
            cell.border = cell_border
            if col_idx not in (13, 14):
                cell.font = regular_font
            if curr_fill:
                cell.fill = curr_fill

    col_widths = {
        1: 6, 2: 17, 3: 15, 4: 13, 5: 8, 6: 14, 7: 16,
        8: 11, 9: 9, 10: 40, 11: 60, 12: 8, 13: 15, 14: 15
    }
    for c_idx, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = w

    ws.auto_filter.ref = f"A1:N{len(qa_items)+1}"
    ws.freeze_panes = "A2"
def build_statistics_index_sheet(wb, records):
    ws = wb.create_sheet(title="📈 주요통계_색인표")
    ws.views.sheetView[0].showGridLines = True

    table_records = [r for r in records if r.get("has_tables") == "Y"]

    headers = [
        "연번", "문서ID", "작성/요구일", "요구기관/요구자", "관리번호",
        "자료 제목", "포함된 표 제목/내용", "표 개수", "주제 태그", "마크다운 열기", "원본 열기"
    ]
    apply_table_header(ws, 1, headers)

    for idx, r in enumerate(table_records, start=1):
        row = idx + 1
        is_even = (idx % 2 == 0)
        curr_fill = zebra_fill if is_even else None

        ws.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=2, value=sanitize_text(r.get("doc_id", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=3, value=sanitize_text(r.get("request_date", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=4, value=sanitize_text(f"{r.get('institution', '')} / {r.get('requester', '')}")).alignment = Alignment(horizontal="left", vertical="top")
        ws.cell(row=row, column=5, value=sanitize_text(r.get("doc_number", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=6, value=sanitize_text(r.get("title", ""))).alignment = Alignment(horizontal="left", vertical="top")
        ws.cell(row=row, column=7, value=sanitize_text(r.get("table_summary") or "상세 수치표 포함")).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=row, column=8, value=r.get("table_count", 0)).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=row, column=9, value=sanitize_text(r.get("topic_tags", ""))).alignment = Alignment(horizontal="left", vertical="top")

        md_cell = ws.cell(row=row, column=10)
        md_path_norm = resolve_excel_hyperlink_path(r.get("parsed_md_path", ""))
        md_cell.value = f'=HYPERLINK("{md_path_norm}", "마크다운 보기")'
        md_cell.font = link_font
        md_cell.alignment = Alignment(horizontal="center", vertical="top")

        orig_cell = ws.cell(row=row, column=11)
        orig_path_norm = resolve_excel_hyperlink_path(r.get("original_path", ""))
        orig_cell.value = f'=HYPERLINK("{orig_path_norm}", "원본 열기")'
        orig_cell.font = link_font
        orig_cell.alignment = Alignment(horizontal="center", vertical="top")

        for col_idx in range(1, 12):
            cell = ws.cell(row=row, column=col_idx)
            cell.border = cell_border
            if col_idx not in (10, 11):
                cell.font = regular_font
            if curr_fill:
                cell.fill = curr_fill

    col_widths = {
        1: 6, 2: 15, 3: 13, 4: 24, 5: 11,
        6: 35, 7: 45, 8: 9, 9: 24, 10: 15, 11: 15
    }
    for c_idx, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = w

    ws.auto_filter.ref = f"A1:K{len(table_records)+1}"
    ws.freeze_panes = "A2"
def build_ledger_sheet(wb, ledger_items):
    ws = wb.create_sheet(title="📋 국회_요구자료_관리대장")
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "No.", "대장ID", "연도", "연번", "소속", "의원", "보좌관", "요구자료",
        "세부내역", "요구일", "마감일", "제출일", "부서", "진행상태", "비고", "요청형태", "답변문서 연결"
    ]
    apply_table_header(ws, 1, headers)

    for idx, r in enumerate(ledger_items, start=1):
        row = idx + 1
        curr_fill = zebra_fill if idx % 2 == 0 else None

        ws.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=2, value=sanitize_text(r.get("ledger_id", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=3, value=sanitize_text(r.get("year", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=4, value=sanitize_text(r.get("seq_no", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=5, value=sanitize_text(r.get("party", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=6, value=sanitize_text(r.get("requester", ""))).alignment = Alignment(horizontal="left", vertical="top")
        ws.cell(row=row, column=7, value=sanitize_text(r.get("aide", ""))).alignment = Alignment(horizontal="left", vertical="top")
        ws.cell(row=row, column=8, value=sanitize_text(r.get("title", ""))).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=row, column=9, value=sanitize_text(r.get("details", ""))).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=row, column=10, value=sanitize_text(r.get("request_date", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=11, value=sanitize_text(r.get("deadline", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=12, value=sanitize_text(r.get("submit_date", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=13, value=sanitize_text(r.get("department", ""))).alignment = Alignment(horizontal="left", vertical="top")
        ws.cell(row=row, column=14, value=sanitize_text(r.get("status", ""))).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=15, value=sanitize_text(r.get("note", ""))).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=row, column=16, value=sanitize_text(r.get("request_type", ""))).alignment = Alignment(horizontal="center", vertical="top")
        
        linked_doc = r.get("linked_doc_id", "")
        doc_cell = ws.cell(row=row, column=17)
        doc_cell.alignment = Alignment(horizontal="center", vertical="top")
        if linked_doc:
            doc_cell.value = f'=HYPERLINK("#\'📋 자료요구_통합목록\'!A1", "{sanitize_text(linked_doc)}")'
            doc_cell.font = link_font
        else:
            doc_cell.value = "-"

        for col_idx in range(1, 18):
            cell = ws.cell(row=row, column=col_idx)
            cell.border = cell_border
            if col_idx != 17 or not linked_doc:
                cell.font = regular_font
            if curr_fill:
                cell.fill = curr_fill

    col_widths = {
        1: 6, 2: 15, 3: 8, 4: 8, 5: 14, 6: 12, 7: 10,
        8: 30, 9: 50, 10: 12, 11: 12, 12: 12, 13: 18, 14: 10,
        15: 20, 16: 10, 17: 15
    }
    for c_idx, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = w

    ws.auto_filter.ref = f"A1:Q{len(ledger_items)+1}"
    ws.freeze_panes = "A2"
