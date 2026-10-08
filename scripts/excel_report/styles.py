# -*- coding: utf-8 -*-
"""엑셀 테마 상수·헤더 스타일(SRP: 서식)."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


NAVY_HEADER = "1F4E79"
SUB_HEADER = "2F5597"
ACCENT_BLUE = "D9E1F2"
ZEBRA_FILL = "F9FAFC"
WHITE = "FFFFFF"
BORDER_GRAY = "D3D3D3"

thin_side = Side(style='thin', color=BORDER_GRAY)
cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

header_font = Font(name="맑은 고딕", size=11, bold=True, color=WHITE)
sub_header_font = Font(name="맑은 고딕", size=11, bold=True, color=WHITE)
title_font = Font(name="맑은 고딕", size=16, bold=True, color="1F4E79")
section_font = Font(name="맑은 고딕", size=13, bold=True, color="1F4E79")
regular_font = Font(name="맑은 고딕", size=10)
bold_font = Font(name="맑은 고딕", size=10, bold=True)
link_font = Font(name="맑은 고딕", size=10, color="0563C1", underline="single")

header_fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
sub_header_fill = PatternFill(start_color=SUB_HEADER, end_color=SUB_HEADER, fill_type="solid")
accent_fill = PatternFill(start_color=ACCENT_BLUE, end_color=ACCENT_BLUE, fill_type="solid")
zebra_fill = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")
def apply_table_header(ws, row, headers):
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws.row_dimensions[row].height = 28
