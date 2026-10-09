"""Inspect and, on request, rebuild the JSON cache from the SQLite source of truth.

Usage:
    python scripts/reconcile_storage.py           # read-only report
    python scripts/reconcile_storage.py --write   # atomically rebuild JSON cache
"""

import json
import sqlite3
import sys
from pathlib import Path

SCRIPTS_DIR = Path(__file__).resolve().parent
SYSTEM_DIR = SCRIPTS_DIR.parent
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from services.ledger_service import LedgerService


def db_counts(db_path: Path) -> dict:
    if not db_path.exists():
        return {"available": False, "error": "SQLite DB 파일이 없습니다."}
    conn = sqlite3.connect(str(db_path))
    try:
        tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        required = {"documents", "qa_items", "request_ledger"}
        if not required.issubset(tables):
            return {"available": False, "error": "SQLite 스키마가 완전하지 않습니다."}
        return {
            "available": True,
            "documents": conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0],
            "qa_items": conn.execute("SELECT COUNT(*) FROM qa_items").fetchone()[0],
            "request_ledger": conn.execute("SELECT COUNT(*) FROM request_ledger WHERE COALESCE(status, '') != '[삭제]'").fetchone()[0],
            "tombstones": conn.execute("SELECT COUNT(*) FROM request_ledger WHERE status = '[삭제]'").fetchone()[0],
        }
    finally:
        conn.close()


def json_counts(json_path: Path) -> dict:
    if not json_path.exists():
        return {"available": False, "error": "JSON 캐시 파일이 없습니다."}
    try:
        data = json.loads(json_path.read_text(encoding="utf-8"))
        return {
            "available": True,
            "documents": len(data.get("documents", [])),
            "qa_items": len(data.get("qa_items", [])),
            "request_ledger": len(data.get("request_ledger", [])),
        }
    except (OSError, ValueError, TypeError) as exc:
        return {"available": False, "error": f"JSON 읽기 실패: {exc}"}


def pending_report() -> dict:
    """보류 큐(.excel_pending_queue.json)에 남은 엑셀 반영 작업을 읽는다."""
    path = SYSTEM_DIR / ".excel_pending_queue.json"
    if not path.exists():
        return {"pending": 0, "ids": []}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        ops = data if isinstance(data, list) else []
        return {
            "pending": len(ops),
            "ids": [str((op.get("item") or {}).get("ledger_id") or "") for op in ops],
        }
    except (OSError, ValueError, TypeError) as exc:
        return {"pending": -1, "error": f"보류 큐 읽기 실패: {exc}"}


def excel_report(db_path: Path) -> dict:
    """마스터 엑셀과 SQLite 대장의 ID 집합을 비교한다 (읽기 전용)."""
    if not db_path.exists():
        return {"available": False, "error": "SQLite DB 파일이 없습니다."}
    try:
        from extractors.ledger_parser import load_request_ledger
        parsed = load_request_ledger(SYSTEM_DIR, [], merge_db_records=False)
    except Exception as exc:
        return {"available": False, "error": f"엑셀 대장 파싱 실패: {exc}"}

    excel_ids = {str(item.get("ledger_id")) for item in parsed if item.get("ledger_id")}
    conn = sqlite3.connect(str(db_path))
    try:
        rows = {
            r[0]: (str(r[1] or ""), str(r[2] or ""))
            for r in conn.execute("SELECT ledger_id, status, request_type FROM request_ledger")
        }
    finally:
        conn.close()

    active = {lid for lid, (status, _t) in rows.items() if status != "[삭제]"}
    tombstoned = {lid for lid, (status, _t) in rows.items() if status == "[삭제]"}
    return {
        "available": True,
        "excel_rows": len(excel_ids),
        "db_active": len(active),
        "only_in_excel": sorted(excel_ids - set(rows)),
        "only_in_db": sorted(lid for lid in (active - excel_ids) if rows[lid][1] != "시스템"),
        "excel_rows_tombstoned_in_db": sorted(excel_ids & tombstoned),
        "ghost_rows": ghost_excel_rows(),
    }


def ghost_excel_rows() -> list:
    """`대장ID`만 있고 본문 칸이 모두 빈 행을 찾는다.

    파서는 연번·제목·세부내역·의원이 전부 비면 데이터 행으로 보지 않아 이런 행을
    완전히 무시한다. 그래서 `only_in_excel`에도 잡히지 않고 워크북에 영구히 쌓인다.
    """
    try:
        import openpyxl
        from extractors.ledger_parser import (
            build_header_map, find_id_column, is_ledger_data_row, year_sheet_names,
        )
        from services.excel_sync_service import ExcelSyncService
    except Exception:
        return []

    master = ExcelSyncService(base_dir=SYSTEM_DIR).find_master_excel()
    if not master or not master.exists():
        return []

    ghosts = []
    try:
        wb = openpyxl.load_workbook(str(master), data_only=True)
    except Exception:
        return []
    try:
        for sheet in year_sheet_names(wb.sheetnames):
            ws = wb[sheet]
            header_map = build_header_map(ws)
            c_id = find_id_column(header_map)
            if not c_id:
                continue
            for r in range(2, ws.max_row + 1):
                lid = str(ws.cell(r, c_id).value or "").strip()
                if lid and not is_ledger_data_row(ws, r, header_map):
                    ghosts.append(f"{sheet}!{r} ({lid})")
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return ghosts


def main(argv=None) -> int:
    argv = argv or sys.argv[1:]
    if any(arg not in {"--write", "--excel", "--help", "-h"} for arg in argv):
        print("사용법: python scripts/reconcile_storage.py [--write] [--excel]")
        return 2
    if "--help" in argv or "-h" in argv:
        print(__doc__)
        return 0

    db_path = SYSTEM_DIR / "data_requests.db"
    json_path = SYSTEM_DIR / "data_requests.json"
    db = db_counts(db_path)
    cached = json_counts(json_path)
    print("SQLite:", db)
    print("JSON:  ", cached)

    queued = pending_report()
    if queued.get("pending"):
        print("보류 큐:", queued)

    if "--excel" in argv:
        excel = excel_report(db_path)
        print("Excel: ", excel)
        if excel.get("available"):
            if excel.get("ghost_rows"):
                print(
                    f"⚠ 대장ID만 있고 본문이 빈 유령 행 {len(excel['ghost_rows'])}건: "
                    f"{', '.join(excel['ghost_rows'][:10])}\n"
                    "  파서가 무시하는 행입니다. 엑셀에서 직접 삭제하세요."
                )
            if excel["only_in_excel"] or excel["only_in_db"] or excel["excel_rows_tombstoned_in_db"]:
                print(
                    "⚠ SQLite와 마스터 엑셀의 대장 항목이 일치하지 않습니다. "
                    "웹 화면에서 해당 항목을 확인하거나, 엑셀을 닫고 POST /api/sync/excel 로 재동기화하세요."
                )
            else:
                print("✓ SQLite와 마스터 엑셀의 대장 ID 집합이 일치합니다.")

    if not db.get("available"):
        print("✗ SQLite를 정본으로 확인할 수 없어 복구를 수행하지 않았습니다.")
        return 1

    mismatch = (not cached.get("available")) or any(
        db[key] != cached.get(key) for key in ("documents", "qa_items", "request_ledger")
    )
    if not mismatch:
        print("✓ SQLite와 JSON 캐시의 레코드 수가 일치합니다.")
        return 0
    if "--write" not in argv:
        print("⚠ 불일치 감지. JSON만 재생성하려면 --write를 사용하세요.")
        return 1

    LedgerService(db_path=db_path, json_path=json_path, base_dir=SYSTEM_DIR).sync_json_file(async_mode=False)
    repaired = json_counts(json_path)
    if repaired.get("available") and all(db[key] == repaired.get(key) for key in ("documents", "qa_items", "request_ledger")):
        print("✓ JSON 캐시를 SQLite 기준으로 원자적으로 재생성했습니다.")
        return 0
    print("✗ JSON 재생성 뒤에도 불일치가 남았습니다. DB/Excel을 확인하세요.")
    return 1


if __name__ == "__main__":
    sys.exit(main())
