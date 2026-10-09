from typing import Any, cast
import os
import sys
import json
import shutil
from pathlib import Path
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

cast(Any, sys.stdout).reconfigure(encoding='utf-8')




SCRIPTS_DIR = Path(__file__).resolve().parent
SYSTEM_DIR = SCRIPTS_DIR.parent
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))
import system_config
from excel_report.text import _excel_body, sanitize_text
from excel_report.styles import (
    NAVY_HEADER, SUB_HEADER, ACCENT_BLUE, ZEBRA_FILL, WHITE, BORDER_GRAY,
    thin_side, cell_border, header_font, sub_header_font, title_font,
    section_font, regular_font, bold_font, link_font, header_fill,
    sub_header_fill, accent_fill, zebra_fill, apply_table_header,
)
from excel_report.links import resolve_excel_hyperlink_path
from excel_report.sheets import (
    build_dashboard_sheet, build_master_list_sheet, build_qa_detail_sheet,
    build_statistics_index_sheet, build_ledger_sheet,
)

# 하이퍼링크 기준 폴더는 파싱·DB 구축과 같은 규칙으로 정한다. (배포본은 자기 폴더가 루트)
WORKSPACE_ROOT = system_config.find_workspace_root(SYSTEM_DIR)
ROOT_DIR = WORKSPACE_ROOT

CFG = system_config.get_config(SYSTEM_DIR)
JSON_PATH = SYSTEM_DIR / "data_requests.json"
EXCEL_PATH = SYSTEM_DIR / CFG.get("excel_filename", "국회_대외기관_자료요구_통합DB.xlsx")









def main():
    print("=" * 60)
    print("고도화 엑셀 통합 데이터베이스 생성 시작 (답변 전문 및 Q&A 상세DB, 관리대장)")
    print("=" * 60)

    if not JSON_PATH.exists():
        print(f"오류: {JSON_PATH} 파일이 없습니다.")
        return False

    with open(JSON_PATH, "r", encoding="utf-8") as fp:
        raw_data = json.load(fp)

    if isinstance(raw_data, dict):
        records = raw_data.get("documents", [])
        qa_items = raw_data.get("qa_items", [])
        ledger_items = raw_data.get("request_ledger", [])
    else:
        records = raw_data
        qa_items = []
        ledger_items = []

    print(f"로드된 문서 레코드: {len(records)}건")
    print(f"로드된 세부 Q&A 레코드: {len(qa_items)}건")
    print(f"로드된 관리대장 레코드: {len(ledger_items)}건")

    wb = openpyxl.Workbook()

    # 1. Dashboard
    build_dashboard_sheet(wb, records, qa_items, ledger_items)

    # 2. Master List with Full Text column & Ledger link
    build_master_list_sheet(wb, records)

    # 3. Q&A Detail DB (1:N)
    if qa_items:
        build_qa_detail_sheet(wb, qa_items)

    # 4. Request Ledger (933건 전체 마스터 대장)
    if ledger_items:
        build_ledger_sheet(wb, ledger_items)

    # 5. Statistics Index
    build_statistics_index_sheet(wb, records)

    ok = save_workbook_atomic(wb, EXCEL_PATH)
    try:
        wb.close()
    except Exception:
        pass
    return ok


def save_workbook_atomic(wb, excel_path: Path) -> bool:
    """xlsx를 tmp에 저장한 뒤 원자적으로 교체한다. 잠금 시 대기 파일을 남기고 False."""
    excel_path = Path(excel_path)
    tmp_path = excel_path.with_name(excel_path.stem + ".tmp.xlsx")
    try:
        wb.save(tmp_path)
        # 저장 결과가 실제로 열리는지 확인한 뒤에만 원본을 교체한다. (감사 5회차 ISSUE-002)
        probe = openpyxl.load_workbook(tmp_path, read_only=True)
        try:
            if not probe.sheetnames:
                raise ValueError("저장된 워크북에 시트가 없습니다.")
        finally:
            probe.close()
        try:
            os.replace(tmp_path, excel_path)
        except PermissionError:
            try:
                shutil.copy2(tmp_path, excel_path)
                tmp_path.unlink(missing_ok=True)
            except PermissionError:
                backup_path = excel_path.with_name(f"{excel_path.stem}_최신동기화_대기.xlsx")
                try:
                    os.replace(tmp_path, backup_path)
                except Exception:
                    wb.save(backup_path)
                    tmp_path.unlink(missing_ok=True)
                print("\n" + "!" * 65)
                print(f"⚠️ 경고: {excel_path.name} 파일이 엑셀에서 열려 있어 덮어쓰지 못했습니다.")
                print(f"최신본은 {backup_path.name} 에 저장했습니다. 엑셀을 닫은 뒤 03 배치로 열어 주세요.")
                print("!" * 65 + "\n")
                return False
        print("=" * 60)
        print(f"고도화 엑셀 통합 DB 생성 완료: {excel_path.name} ({excel_path.stat().st_size:,} bytes)")
        print("=" * 60)
        return True
    except Exception as e:
        print(f"엑셀 저장 실패: {e}")
        return False
    finally:
        # 실패 잔재(tmp)가 다음 실행의 대장으로 오인되지 않게 정리한다.
        try:
            if tmp_path.exists():
                tmp_path.unlink()
        except OSError:
            pass

if __name__ == "__main__":
    main()
