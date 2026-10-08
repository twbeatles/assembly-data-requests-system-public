# -*- coding: utf-8 -*-
"""kordoc 실패 시 대용량 엑셀 대체 추출(SRP: 폴백 파서)."""
import html
from pathlib import Path

import openpyxl


def handle_large_xlsx_fallback(file_path: Path, out_file: Path, max_rows: int = 2000, max_cols: int = 40):
    """kordoc 실패 시에만 쓰는 대용량 엑셀 대체 추출. 시트 전체와 최대 max_rows행을 보존한다."""
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        md_lines = [
            f"# {file_path.name}\n\n",
            f"**파일 크기**: {file_path.stat().st_size / (1024*1024):.2f} MB\n",
            f"**시트 목록**: {list(wb.sheetnames)}\n\n"
        ]
        for sname in wb.sheetnames:
            ws = wb[sname]
            md_lines.append(f"## 시트: {sname}\n\n")
            rows = []
            truncated = False
            for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
                if r_idx >= max_rows:
                    truncated = True
                    break
                rows.append([str(c or '') for c in row[:max_cols]])
            if rows:
                md_lines.append("<table>\n")
                for r in rows:
                    md_lines.append("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in r) + "</tr>\n")
                md_lines.append("</table>\n\n")
                if truncated:
                    md_lines.append(f"* (대체 파싱: {max_rows}행까지 수록, 이후 행 생략)\n\n")
        wb.close()

        with open(out_file, "w", encoding="utf-8") as fp:
            fp.write("".join(md_lines))
        return True, None
    except Exception as e:
        return False, str(e)
