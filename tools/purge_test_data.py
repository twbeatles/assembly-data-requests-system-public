from typing import Any, cast
# -*- coding: utf-8 -*-
"""운영 대장에 남은 테스트 데이터 정리 도구 (감사 3회차 R3-01 후속).

과거 테스트 실행이 운영 마스터 엑셀·SQLite·보류 큐에 남긴 테스트 페이로드를 제거한다.
테스트 코드의 페이로드와 **(요구자, 제목) 쌍이 정확히 일치하는 항목만** 대상으로 삼는다.

    python tools/purge_test_data.py            # 점검만 한다 (아무것도 바꾸지 않음)
    python tools/purge_test_data.py --apply    # 백업 → 엑셀 정리 → DB 정리 → 격리 파일 정리

실행 전 조건: 02번 웹관리서버·05번 프로그램·00번 파이프라인이 꺼져 있고, 마스터 엑셀이
어떤 Excel 창에도 열려 있지 않아야 한다.

엑셀은 openpyxl이 아니라 Excel 자체(COM, 매크로·이벤트 비활성)로 편집해 VBA·단추·서식을
보존한다. 정리가 끝나면 `00_새자료_추가_및_DB동기화.bat`을 실행해 엑셀 수정을 DB에 흡수하고
JSON·대시보드·통합 엑셀·배포본을 다시 만든다.
"""

import argparse
import datetime
import hashlib
import json
import shutil
import sqlite3
import sys
from pathlib import Path

SYSTEM_DIR = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = SYSTEM_DIR / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))
cast(Any, sys.stdout).reconfigure(encoding="utf-8")

import openpyxl  # noqa: E402
from extractors.ledger_parser import plan_ledger_ids, connect_readonly  # noqa: E402
from services.excel_sync_service import ExcelSyncService  # noqa: E402
from pipeline_guard import PipelineGuard  # noqa: E402

# 테스트 코드가 등록·수정에 쓰던 (요구자, 제목) 쌍. 운영 데이터와 겹칠 수 없는 조합만 둔다.
TEST_SIGNATURES = {
    ("테스트의원", "테스트 제목"),
    ("테스트의원", "수정 제목"),
    ("테스트의원", "딥페이크 테스트"),
    ("단위테스트의원", "단위테스트 요구자료 제목"),
    ("의원", "원자적 동기화 테스트"),
    ("홍길동 의원", "추가 자동 연계 테스트 자료"),
    ("홍길동의원", "테스트 요구자료 제목"),
    ("신규의원", "백업 확인용"),
}
CONSOLIDATED_SHEET = "통합목록(검색용)"


def is_test_pair(requester, title) -> bool:
    return (str(requester or "").strip(), str(title or "").strip()) in TEST_SIGNATURES


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fp:
        for chunk in iter(lambda: fp.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------------------------------------------------------------------------
# 점검
# ---------------------------------------------------------------------------
def plan_database(db_path: Path) -> dict:
    conn = connect_readonly(db_path)
    conn.row_factory = sqlite3.Row
    try:
        ledger = [dict(r) for r in conn.execute("SELECT * FROM request_ledger")]
        purge_ids = sorted(r["ledger_id"] for r in ledger if is_test_pair(r["requester"], r["title"]))
        history = []
        for h in conn.execute("SELECT history_id, ledger_id, action, snapshot FROM ledger_history"):
            if h["ledger_id"] in purge_ids:
                history.append(h["history_id"])
                continue
            if h["action"] == "EXCEL_APPLIED":
                try:
                    snap = json.loads(h["snapshot"] or "{}")
                except ValueError:
                    continue
                if is_test_pair(snap.get("requester"), snap.get("title")):
                    history.append(h["history_id"])
        docs = [r[0] for r in conn.execute(
            f"SELECT doc_id FROM documents WHERE linked_ledger_id IN ({','.join('?' * len(purge_ids))})", purge_ids
        )] if purge_ids else []
        remaining = [r for r in ledger if r["ledger_id"] not in purge_ids]
        return {"purge_ids": purge_ids, "history_ids": sorted(history), "unlink_docs": docs,
                "remaining_records": remaining, "ledger_rows": ledger}
    finally:
        conn.close()


def header_index(values_row):
    return {str(v).strip(): i for i, v in enumerate(values_row) if v is not None and str(v).strip()}


def plan_workbook(excel_path: Path, remaining_records) -> dict:
    wb = openpyxl.load_workbook(str(excel_path), read_only=True, data_only=True)
    try:
        plan = {"delete_rows": {}, "renumber": False}
        for name in wb.sheetnames:
            ws = wb[name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            idx = header_index(rows[0])
            req_col = idx.get("의원")
            title_col = idx.get("요구자료")
            if req_col is None or title_col is None:
                continue
            hits = [
                rn for rn, row in enumerate(rows[1:], start=2)
                if len(row) > max(req_col, title_col) and is_test_pair(row[req_col], row[title_col])
            ]
            if hits:
                plan["delete_rows"][name] = hits
            if name == CONSOLIDATED_SHEET and "No." in idx:
                nos = [row[idx["No."]] for row in rows[1:] if any(v not in (None, "") for v in row)]
                plan["renumber"] = nos == list(range(1, len(nos) + 1))
    finally:
        wb.close()

    # 테스트 행을 뺀 상태에서 비어 있는 대장ID에 줄 ID를 계산한다(스탬핑을 Excel로 직접 기록).
    wb_vals = openpyxl.load_workbook(str(excel_path), data_only=True)
    try:
        for name, hits in plan["delete_rows"].items():
            if name.isdigit():
                ws = wb_vals[name]
                for rn in sorted(hits, reverse=True):
                    ws.delete_rows(rn)
        assignments, missing_cols = plan_ledger_ids(wb_vals, remaining_records)
    finally:
        wb_vals.close()
    plan["id_assignments"] = assignments
    plan["sheets_without_id_column"] = sorted(missing_cols)
    return plan


def plan_quarantine(path: Path) -> dict:
    if not path.exists():
        return {"remove": 0, "keep": 0}
    ops = json.loads(path.read_text(encoding="utf-8"))
    remove = [op for op in ops if is_test_pair((op.get("item") or {}).get("requester"), (op.get("item") or {}).get("title"))]
    return {"remove": len(remove), "keep": len(ops) - len(remove)}


# ---------------------------------------------------------------------------
# 실행
# ---------------------------------------------------------------------------
def preconditions(excel_path: Path) -> list:
    problems = []
    if PipelineGuard(SYSTEM_DIR).is_locked():
        problems.append("파이프라인 락(.pipeline.lock)이 있습니다. 00번 파이프라인이 끝난 뒤 실행하세요.")
    if any(SYSTEM_DIR.glob("~$*.xls*")):
        problems.append("Excel 소유자 파일(~$...)이 있습니다. 마스터 엑셀을 닫은 뒤 실행하세요.")
    if ExcelSyncService(base_dir=SYSTEM_DIR).is_file_locked(excel_path):
        problems.append("마스터 엑셀이 다른 프로그램에서 열려 있습니다.")
    import socket
    with socket.socket() as s:
        s.settimeout(0.5)
        if s.connect_ex(("127.0.0.1", 8080)) == 0:
            problems.append("127.0.0.1:8080에서 서버가 실행 중입니다. 02번 웹관리서버를 종료하세요.")
    # 8080이 아닌 포트로 뜬 서버도 실행 락으로 잡는다. (감사 R4-05)
    web_lock = SYSTEM_DIR / ".web_server.lock"
    if web_lock.exists():
        from pipeline_guard import pid_is_alive
        try:
            pid = int(json.loads(web_lock.read_text(encoding="utf-8")).get("pid") or 0)
        except (OSError, ValueError, TypeError, AttributeError):
            pid = 0
        if pid and pid_is_alive(pid):
            problems.append(f"02번 웹관리서버가 실행 중입니다(PID {pid}). 서버 창을 닫은 뒤 실행하세요.")
    return problems


def make_backup(excel_path: Path, db_path: Path) -> Path:
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = SYSTEM_DIR / ".maintenance_backup" / f"{stamp}_test_data_purge"
    dest.mkdir(parents=True)
    manifest = {}
    for name in [excel_path.name, "data_requests.json", ".excel_pending_queue.json", ".excel_pending_failed.json",
                 "config.json"]:
        src = SYSTEM_DIR / name
        if src.exists():
            shutil.copy2(src, dest / name)
            if sha256(src) != sha256(dest / name):
                raise RuntimeError(f"백업 검증 실패: {name}")
            manifest[name] = sha256(src)
    src = connect_readonly(db_path)
    dst = sqlite3.connect(str(dest / "data_requests.db"))
    try:
        src.backup(dst)
    finally:
        dst.close()
        src.close()
    chk = sqlite3.connect(str(dest / "data_requests.db"))
    try:
        if chk.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise RuntimeError("DB 백업 무결성 검사 실패")
    finally:
        chk.close()
    manifest["data_requests.db"] = "sqlite backup API, integrity ok"
    (dest / "MANIFEST.json").write_text(json.dumps({"created_at": stamp, "files": manifest}, ensure_ascii=False, indent=2),
                                        encoding="utf-8")
    return dest


def apply_workbook(excel_path: Path, plan: dict):
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    wb = None
    try:
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AutomationSecurity = 3  # msoAutomationSecurityForceDisable: 열 때 매크로 실행 안 함
        excel.EnableEvents = False
        wb = excel.Workbooks.Open(str(excel_path), 0, False)
        for name, hits in plan["delete_rows"].items():
            ws = wb.Worksheets(name)
            last_col = ws.UsedRange.Columns.Count + ws.UsedRange.Column - 1
            header = [ws.Cells(1, c).Value for c in range(1, last_col + 1)]
            idx = header_index(header)
            for rn in sorted(hits, reverse=True):
                requester = ws.Cells(rn, idx["의원"] + 1).Value
                title = ws.Cells(rn, idx["요구자료"] + 1).Value
                if not is_test_pair(requester, title):
                    raise RuntimeError(f"{name}!{rn} 행이 점검 때와 달라 중단합니다: {requester} / {title}")
                ws.Rows(rn).Delete()
            if name == CONSOLIDATED_SHEET and plan.get("renumber") and "No." in idx:
                col = idx["No."] + 1
                last_row = ws.UsedRange.Rows.Count + ws.UsedRange.Row - 1
                n = 0
                for r in range(2, last_row + 1):
                    if any(ws.Cells(r, c).Value not in (None, "") for c in range(2, min(last_col, 8) + 1)):
                        n += 1
                        ws.Cells(r, col).Value = n
        for name, row, ledger_id in plan["id_assignments"]:
            ws = wb.Worksheets(name)
            last_col = ws.UsedRange.Columns.Count + ws.UsedRange.Column - 1
            idx = header_index([ws.Cells(1, c).Value for c in range(1, last_col + 1)])
            if "대장ID" not in idx:
                raise RuntimeError(f"{name} 시트에 대장ID 열이 없습니다. 00번 파이프라인이 열을 추가하게 두세요.")
            ws.Cells(row, idx["대장ID"] + 1).Value = ledger_id
        wb.Save()
    finally:
        if wb is not None:
            wb.Close(False)
        excel.Quit()
        pythoncom.CoUninitialize()


def apply_database(db_path: Path, plan: dict, backup_dir: Path):
    ids = plan["purge_ids"]
    conn = sqlite3.connect(str(db_path), timeout=10)
    try:
        conn.execute("PRAGMA busy_timeout = 10000")
        conn.execute("BEGIN IMMEDIATE")
        q = ",".join("?" * len(ids))
        if plan["history_ids"]:
            conn.execute(f"DELETE FROM ledger_history WHERE history_id IN ({','.join('?' * len(plan['history_ids']))})",
                         plan["history_ids"])
        if ids:
            conn.execute(f"UPDATE documents SET linked_ledger_id = '' WHERE linked_ledger_id IN ({q})", ids)
            conn.execute(f"DELETE FROM request_ledger WHERE ledger_id IN ({q})", ids)
        conn.execute(
            "INSERT INTO ledger_history (ledger_id, action, changed_at, snapshot) VALUES (?, ?, ?, ?)",
            ("MAINTENANCE", "TEST_DATA_PURGE", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
             json.dumps({"reason": "감사 R3-01 테스트 데이터 정리", "backup": str(backup_dir.relative_to(SYSTEM_DIR)),
                         "purged_ledger_ids": ids, "purged_history_rows": len(plan["history_ids"]),
                         "unlinked_documents": plan["unlink_docs"]}, ensure_ascii=False)),
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def apply_quarantine(path: Path, backup_dir: Path):
    if not path.exists():
        return
    ops = json.loads(path.read_text(encoding="utf-8"))
    keep = [op for op in ops if not is_test_pair((op.get("item") or {}).get("requester"), (op.get("item") or {}).get("title"))]
    if keep:
        tmp = path.with_suffix(".json.tmp")
        tmp.write_text(json.dumps(keep, ensure_ascii=False, indent=2), encoding="utf-8")
        tmp.replace(path)
    else:
        path.unlink()


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--apply", action="store_true", help="점검 결과대로 실제 정리를 수행한다")
    args = parser.parse_args()

    sync = ExcelSyncService(base_dir=SYSTEM_DIR)
    excel_path = sync.find_master_excel()
    db_path = SYSTEM_DIR / "data_requests.db"
    if excel_path is None or not db_path.exists():
        print("마스터 엑셀 또는 data_requests.db를 찾을 수 없습니다.")
        return 1

    db_plan = plan_database(db_path)
    wb_plan = plan_workbook(excel_path, db_plan["remaining_records"])
    q_plan = plan_quarantine(SYSTEM_DIR / ".excel_pending_failed.json")
    pending = json.loads((SYSTEM_DIR / ".excel_pending_queue.json").read_text(encoding="utf-8")) \
        if (SYSTEM_DIR / ".excel_pending_queue.json").exists() else []
    pending_test = [op for op in pending if is_test_pair((op.get("item") or {}).get("requester"), (op.get("item") or {}).get("title"))]

    print("=" * 70)
    print(f"마스터 엑셀: {excel_path.name}")
    print(f"[DB] 삭제할 테스트 대장 행 {len(db_plan['purge_ids'])}건: {', '.join(db_plan['purge_ids'])}")
    print(f"[DB] 삭제할 테스트 이력 {len(db_plan['history_ids'])}건, 연결 해제할 문서 {db_plan['unlink_docs']}")
    for name, hits in wb_plan["delete_rows"].items():
        print(f"[엑셀] {name}: 테스트 행 {len(hits)}건 삭제 (행 {hits[0]}~{hits[-1]})")
    if wb_plan["renumber"]:
        print(f"[엑셀] {CONSOLIDATED_SHEET}: No. 열 번호를 다시 매깁니다")
    for name, row, lid in wb_plan["id_assignments"]:
        print(f"[엑셀] {name}!{row}: 비어 있는 대장ID에 {lid} 기록")
    if wb_plan["sheets_without_id_column"]:
        print(f"[엑셀] 대장ID 열이 없는 시트: {wb_plan['sheets_without_id_column']} (00번 파이프라인이 추가)")
    print(f"[격리 파일] 테스트 작업 {q_plan['remove']}건 삭제, 나머지 {q_plan['keep']}건 유지")
    if pending_test:
        print(f"⚠️ 보류 큐에 테스트 작업 {len(pending_test)}건이 남아 있습니다. 먼저 큐를 확인하세요.")
    print("=" * 70)

    if not args.apply:
        print("점검만 했습니다. 실제로 정리하려면 --apply를 붙여 실행하세요.")
        return 0
    if pending_test:
        print("보류 큐에 테스트 작업이 있어 중단합니다.")
        return 1
    problems = preconditions(excel_path)
    if problems:
        print("실행 조건을 만족하지 않아 중단합니다:")
        for p in problems:
            print(f"  - {p}")
        return 1

    backup_dir = make_backup(excel_path, db_path)
    print(f"✓ 백업 완료: {backup_dir}")
    # 엑셀을 먼저 정리한다. DB만 먼저 지우면, 엑셀 정리가 실패했을 때 다음 동기화가
    # 엑셀의 테스트 행을 새 항목으로 다시 넣는다.
    apply_workbook(excel_path, wb_plan)
    print("✓ 마스터 엑셀 정리 완료 (Excel로 저장, 매크로·서식 보존)")
    apply_database(db_path, db_plan, backup_dir)
    print("✓ DB 정리 완료")
    apply_quarantine(SYSTEM_DIR / ".excel_pending_failed.json", backup_dir)
    print("✓ 격리 파일 정리 완료")

    after_db = plan_database(db_path)
    after_wb = plan_workbook(excel_path, after_db["remaining_records"])
    ok = not after_db["purge_ids"] and not after_wb["delete_rows"] and not after_wb["id_assignments"]
    print("✓ 사후 점검: 남은 테스트 데이터 없음" if ok else f"⚠️ 사후 점검에서 남은 항목이 있습니다: {after_db['purge_ids']} {after_wb}")
    print("다음 단계: 00_새자료_추가_및_DB동기화.bat을 실행해 엑셀 수정을 DB에 반영하고 산출물·배포본을 다시 만드세요.")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
