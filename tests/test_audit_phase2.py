from typing import Any
import os
import sys
import unittest
import tempfile
import shutil
import sqlite3
import re
from pathlib import Path

# Add scripts directory to sys.path
TESTS_DIR = Path(__file__).resolve().parent
PROJECT_DIR = TESTS_DIR.parent
SCRIPTS_DIR = PROJECT_DIR / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import parse_all
import extract_and_build_db
import generate_excel_db
import launcher_gui

class TestAuditPhase2(unittest.TestCase):
    """
    Phase 2 Audit Improvements Verification Tests:
    1. Exclusion of staging folders from markdown parser
    2. Pruning of orphan markdowns when source file deleted/moved
    3. GUI DB connection release before external operations and on window close
    4. openpyxl control character sanitization
    5. Excel sheet name accuracy and bidirectional hyperlinks
    6. Web dashboard deep linking for #ledger= and #req=
    7. QA extraction pattern precision (no false positives for □ bullets, supports □ ... ?)
    8. PII masking in request ledger (aide, details, note)
    """

    def test_1_staging_folders_excluded(self):
        """Verify DEFAULT_EXCLUDE_DIRS contains staging/intake directories"""
        self.assertIn("새자료_투입폴더", parse_all.DEFAULT_EXCLUDE_DIRS)
        self.assertIn("새자료_넣는곳", parse_all.DEFAULT_EXCLUDE_DIRS)

    def test_2_orphan_markdown_pruning(self):
        """Verify prune_orphan_markdowns deletes .md files when original files no longer exist"""
        with tempfile.TemporaryDirectory() as tmp_dir:
            root_p = Path(tmp_dir)
            md_dir = root_p / "_parsed_markdown"
            year_dir = root_p / "2026"
            md_year_dir = md_dir / "2026"
            year_dir.mkdir(parents=True)
            md_year_dir.mkdir(parents=True)

            # Valid doc: source and md exist
            source1 = year_dir / "doc1.hwpx"
            source1.write_text("dummy", encoding="utf-8")
            md1 = md_year_dir / "doc1.hwpx.md"
            md1.write_text("# Doc 1 Content", encoding="utf-8")

            # Orphan doc: source was moved/deleted, md still exists
            md_orphan = md_year_dir / "ghost.hwpx.md"
            md_orphan.write_text("# Ghost Content", encoding="utf-8")

            # Staging folder leftover in md_dir
            md_staging = md_dir / "새자료_투입폴더"
            md_staging.mkdir(parents=True)
            staging_file = md_staging / "staging.hwpx.md"
            staging_file.write_text("# Staging Content", encoding="utf-8")

            pruned = parse_all.prune_orphan_markdowns(root_p, md_dir)
            
            # ghost and staging md must be pruned
            self.assertTrue(md1.exists(), "Valid markdown must be preserved")
            self.assertFalse(md_orphan.exists(), "Orphan markdown must be removed")
            self.assertFalse(staging_file.exists(), "Staging markdown must be removed")
            self.assertEqual(pruned, 2)

    def test_3_launcher_gui_connection_lifecycle(self):
        """Verify launcher_gui provides close_db_conn to prevent Windows SQLite file locks"""
        self.assertTrue(hasattr(launcher_gui.LauncherApp, "close_db_conn"))
        # Verify close_db_conn safely handles null or closed connection
        gui_dummy: Any = type("DummyGUI", (), {})()
        gui_dummy._db_conn = sqlite3.connect(":memory:")
        # Attach method
        launcher_gui.LauncherApp.close_db_conn(gui_dummy)
        self.assertIsNone(gui_dummy._db_conn)
        # Calling again should not raise error
        launcher_gui.LauncherApp.close_db_conn(gui_dummy)
        self.assertIsNone(gui_dummy._db_conn)

    def test_4_excel_control_character_sanitization(self):
        """Verify sanitize_text removes illegal XML/openpyxl control characters"""
        dirty_str = "Clean text\x00with null\x08and backspace\x1fcontrol\nnewline preserved\tTab preserved"
        clean = generate_excel_db.sanitize_text(dirty_str)
        self.assertNotIn("\x00", clean)
        self.assertNotIn("\x08", clean)
        self.assertNotIn("\x1f", clean)
        self.assertIn("\n", clean)
        self.assertIn("\t", clean)
        self.assertIn("Clean text", clean)
        # Non-string input should pass untouched
        self.assertEqual(generate_excel_db.sanitize_text(12345), 12345)
        self.assertIsNone(generate_excel_db.sanitize_text(None))

    def test_5_excel_sheet_names_and_bidirectional_links(self):
        """Verify Excel sheet names match formulas and bidirectional links exist"""
        import openpyxl
        excel_file = PROJECT_DIR / "국회_대외기관_자료요구_통합DB.xlsx"
        if not excel_file.exists():
            self.skipTest("생성 Excel 산출물은 소스 체크아웃에 포함되지 않습니다.")
        
        wb = openpyxl.load_workbook(excel_file, read_only=False, data_only=False)
        expected_sheets = ["📊 대시보드", "📋 자료요구_통합목록", "📑 Q&A_상세DB", "📋 국회_요구자료_관리대장", "📈 주요통계_색인표"]
        for s in expected_sheets:
            self.assertIn(s, wb.sheetnames)

        # Check Master sheet column 18 is '연계 대장ID'
        ws_master = wb["📋 자료요구_통합목록"]
        self.assertEqual(ws_master.cell(1, 18).value, "연계 대장ID")
        self.assertEqual(ws_master.cell(1, 19).value, "마크다운 열기")
        self.assertEqual(ws_master.cell(1, 20).value, "원본파일 열기")

        # Check Ledger sheet column 17 references '📋 자료요구_통합목록'
        ws_ledger = wb["📋 국회_요구자료_관리대장"]
        self.assertEqual(ws_ledger.cell(1, 17).value, "답변문서 연결")
        
        # Check linked formula references the correct sheet
        found_link = False
        for r in range(2, min(100, ws_ledger.max_row + 1)):
            v = ws_ledger.cell(r, 17).value
            if v and str(v).startswith("=HYPERLINK"):
                self.assertIn("📋 자료요구_통합목록", str(v))
                self.assertNotIn("전체_자료요구_목록", str(v))
                found_link = True
                break
        self.assertTrue(found_link, "Expected at least one linked doc hyperlink in ledger sheet")

    def test_6_web_dashboard_hash_deep_links(self):
        """Verify web dashboard JavaScript template contains #ledger= and #req= hash routing"""
        tpl_file = SCRIPTS_DIR / "templates" / "dashboard_template.html"
        gen_script = tpl_file.read_text(encoding="utf-8") if tpl_file.exists() else (SCRIPTS_DIR / "generate_web_dashboard.py").read_text(encoding="utf-8")
        self.assertIn("applyDeepLink", gen_script)
        self.assertIn("hashchange", gen_script)
        self.assertIn("startsWith('ledger=')", gen_script)
        self.assertIn("startsWith('req=')", gen_script)
        self.assertIn("openLedgerModal(found)", gen_script)
        self.assertIn("switchMode('ledger')", gen_script)

    def test_7_qa_extraction_pattern_and_interrogative_detection(self):
        """Verify Q_PATTERN eliminates false positives for statement bullets and detects interrogative □ ... ?"""
        # Test statement bullets that should NOT be detected as questions
        statement_md = """
# 제목
□ 추진경과
○ 2024년 5월 회의 개최
□ 주요 현황
○ 관련 통계 보고
□ 향후 조치계획
○ 2026년 하반기 추진
"""
        details = extract_and_build_db.extract_content_details(statement_md)
        # Questions string should NOT contain statement headers like 추진경과, 주요 현황
        self.assertEqual(details["questions"], "")

        # Test valid question starting with □ but ending with ?
        question_md = """
# 제목
□ 대외기관의 딥페이크 성범죄 영상물 대응 현황 및 향후 계획은?
○ 이에 대해 신속한 모니터링을 진행하고 있습니다.
"""
        meta = {"year": "2026", "requester": "의원실", "title": "딥페이크 질의", "institution": "국회"}
        qa_items = extract_and_build_db.split_into_qa_items("DOC-TEST-001", meta, question_md)
        self.assertTrue(any("대응 현황 및 향후 계획은?" in q["question_title"] for q in qa_items))

    def test_8_pii_masking_in_request_ledger(self):
        """Verify mask_pii masks phone numbers and emails in ledger fields"""
        test_text = "보좌관 연락처 010-9876-5432, 이메일 aide_kim@sample.go.kr 문의 요망"
        masked = extract_and_build_db.mask_pii(test_text)
        self.assertNotIn("010-9876-5432", masked)
        self.assertIn("010-****-5432", masked)
        self.assertNotIn("aide_kim@sample.go.kr", masked)
        self.assertIn("ai***@sample.go.kr", masked)

    def test_9_recursive_drop_folder_scanning(self):
        """Verify that add_documents_smart recursively scans nested subfolders and identifies year from parent folder"""
        import add_documents_smart

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_root = Path(tmp_dir)
            drop_dir = tmp_root / "새자료_투입폴더"
            nested_dir = drop_dir / "2024_하반기" / "국회자료"
            nested_dir.mkdir(parents=True)

            test_file = nested_dir / "홍길동의원_질의서.hwp"
            test_file.write_text("dummy hwp", encoding="utf-8")

            # Mock roots
            orig_root = add_documents_smart.ROOT_DIR
            orig_drop = add_documents_smart.DROP_DIRS
            try:
                add_documents_smart.ROOT_DIR = tmp_root
                add_documents_smart.DROP_DIRS = [drop_dir]

                # Run scanning logic
                target_folder = add_documents_smart.determine_target_folder(test_file)
                self.assertEqual(target_folder.name, "2024")

                # Test file moving
                moved = add_documents_smart.copy_or_move_files([test_file], move=True)
                self.assertEqual(len(moved), 1)
                self.assertTrue(moved[0].exists())
                self.assertEqual(moved[0].parent.name, "2024")
                self.assertFalse(test_file.exists())
            finally:
                add_documents_smart.ROOT_DIR = orig_root
                add_documents_smart.DROP_DIRS = orig_drop

if __name__ == "__main__":
    unittest.main()
