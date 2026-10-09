# -*- coding: utf-8 -*-
"""
Unit Test Suite for Excel Bidirectional Synchronization.
Tests both directions:
1. Web -> DB -> Master Excel (real-time append, update, soft-delete, VBA preservation)
2. Master Excel manual edits -> DB (startup sync, runtime upsert, version tracking)
3. Windows file locking and deferred pending queue
"""

import os
import sys
import time
import json
import shutil
import tempfile
import sqlite3
import unittest
from pathlib import Path

SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import openpyxl
from db.database_manager import DatabaseManager
from services.ledger_service import LedgerService
from services.excel_sync_service import ExcelSyncService
import web_server


class TestExcelBidirectionalSync(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp(prefix="test_excel_sync_")
        self.base_dir = Path(self.temp_dir)
        self.db_path = self.base_dir / "data_requests.db"
        self.json_path = self.base_dir / "data_requests.json"

        # Initialize test SQLite DB
        conn = sqlite3.connect(str(self.db_path))
        DatabaseManager.init_schema(conn)
        conn.close()

        # Create dummy master Excel ledger with 2026 sheet
        self.excel_path = self.base_dir / "260904 국회 요구자료 목록(테스트)_스마트검색.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "2026"
        headers = [
            "연번", "소속", "의원", "보좌관", "요구자료", "세부내역",
            "요구일", "마감일", "제출일", "부서", "제출", "비고", "요청형태"
        ]
        ws.append(headers)
        ws.append([
            "1", "과학기술정보상급감독기관", "김의원", "이보좌관", "기존 엑셀 요구자료",
            "세부내역 1", "2026-09-01", "2026-09-03", "2026-09-03", "기획예산팀", "완료", "", "시스템"
        ])
        wb.save(str(self.excel_path))
        wb.close()

        # Seed initial record in DB corresponding to Excel row 1
        conn = sqlite3.connect(str(self.db_path))
        conn.execute("""
            INSERT INTO request_ledger (
                ledger_id, year, seq_no, party, requester, aide,
                title, details, request_date, deadline, submit_date,
                department, status, note, request_type, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            "REQ-2026-001", "2026", "1", "과학기술정보상급감독기관", "김의원", "이보좌관",
            "기존 엑셀 요구자료", "세부내역 1", "2026-09-01", "2026-09-03", "2026-09-03",
            "기획예산팀", "완료", "", "시스템", "2026-09-01 10:00:00", "2026-09-01 10:00:00"
        ))
        conn.commit()
        conn.close()

        self.service = LedgerService(db_path=self.db_path, json_path=self.json_path, base_dir=self.base_dir)
        self.sync_service = ExcelSyncService.get_instance(base_dir=self.base_dir)
        self.sync_service.db_path = self.db_path
        self.sync_service.json_path = self.json_path

    def tearDown(self):
        self.sync_service.stop_file_watcher()
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_01_find_master_excel(self):
        found = self.sync_service.find_master_excel()
        self.assertIsNotNone(found)
        self.assertEqual(found.resolve(), self.excel_path.resolve())

    def test_02_web_insert_syncs_to_excel(self):
        # Insert a new ledger item via service (simulating POST /api/ledger)
        new_item_data = {
            "year": "2026",
            "seq_no": "2",
            "party": "과방위",
            "requester": "박의원",
            "aide": "최보좌관",
            "title": "웹에서 추가한 신규 요구자료",
            "details": "웹 상세 요구내역",
            "request_date": "2026-09-09",
            "deadline": "2026-09-12",
            "submit_date": "",
            "department": "기획예산팀",
            "status": "작성중",
            "note": "웹 등록 테스트",
            "request_type": "일반"
        }

        res = self.service.insert_ledger_item(new_item_data)
        self.assertTrue(res["success"])
        self.assertTrue(res.get("excel_synced"))

        # Verify in Excel file
        wb = openpyxl.load_workbook(str(self.excel_path), data_only=True)
        ws = wb["2026"]
        self.assertEqual(ws.max_row, 3) # Header + 2 data rows
        last_row_vals = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
        wb.close()

        # Check title, requester, party in Excel
        self.assertIn("웹에서 추가한 신규 요구자료", last_row_vals)
        self.assertIn("박의원", last_row_vals)
        self.assertIn("과방위", last_row_vals)

    def test_03_web_update_syncs_to_excel(self):
        # Update existing record REQ-2026-001 (row 1)
        update_data = {
            "title": "기존 엑셀 요구자료(웹에서 수정됨)",
            "status": "제출완료",
            "note": "비고 수정"
        }
        res = self.service.update_ledger_item("REQ-2026-001", update_data)
        self.assertTrue(res["success"])
        self.assertTrue(res.get("excel_synced"))

        # Verify Excel row was updated
        wb = openpyxl.load_workbook(str(self.excel_path), data_only=True)
        ws = wb["2026"]
        row2_vals = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        wb.close()

        self.assertIn("기존 엑셀 요구자료(웹에서 수정됨)", row2_vals)
        self.assertIn("제출완료", row2_vals)
        self.assertIn("비고 수정", row2_vals)

    def test_04_web_delete_syncs_to_excel(self):
        res = self.service.delete_ledger_item("REQ-2026-001")
        self.assertTrue(res["success"])
        self.assertTrue(res.get("excel_synced"))

        wb = openpyxl.load_workbook(str(self.excel_path), data_only=True)
        ws = wb["2026"]
        row2_vals = [str(ws.cell(2, c).value or "") for c in range(1, ws.max_column + 1)]
        wb.close()

        self.assertTrue(any("[삭제]" in v for v in row2_vals))

    def test_05_manual_excel_edits_sync_to_db(self):
        # Simulate user opening Excel manually and adding a 2nd row
        wb = openpyxl.load_workbook(str(self.excel_path))
        ws = wb["2026"]
        ws.append([
            "2", "과방위", "이의원", "정보좌관", "엑셀 수기 추가 자료",
            "수기 상세내역", "2026-09-05", "2026-09-08", "", "기획예산팀", "작성중", "수기작성", "공문"
        ])
        wb.save(str(self.excel_path))
        wb.close()

        # Run sync_excel_to_db
        res = self.sync_service.sync_excel_to_db()
        self.assertTrue(res["success"])
        self.assertGreaterEqual(res["inserted"], 1)

        # Check DB has the new record
        conn = sqlite3.connect(str(self.db_path))
        conn.row_factory = sqlite3.Row
        row = conn.execute("SELECT * FROM request_ledger WHERE title = '엑셀 수기 추가 자료'").fetchone()
        conn.close()

        self.assertIsNotNone(row)
        self.assertEqual(row["requester"], "이의원")
        self.assertEqual(row["party"], "과방위")

    def test_06_manual_excel_cell_update_syncs_to_db(self):
        # User manually edits an existing row's title and status in Excel
        wb = openpyxl.load_workbook(str(self.excel_path))
        ws = wb["2026"]
        # Column 5 is '요구자료' (title)
        ws.cell(2, 5, value="기존 요구자료 (엑셀에서 수기 수정)")
        ws.cell(2, 11, value="보완요청")
        wb.save(str(self.excel_path))
        wb.close()

        # Trigger sync
        res = self.sync_service.sync_excel_to_db()
        self.assertTrue(res["success"])
        self.assertGreaterEqual(res["updated"], 1)

        # Verify DB reflects manual edit
        conn = sqlite3.connect(str(self.db_path))
        conn.row_factory = sqlite3.Row
        row = conn.execute("SELECT * FROM request_ledger WHERE ledger_id = 'REQ-2026-001'").fetchone()
        conn.close()

        self.assertEqual(row["title"], "기존 요구자료 (엑셀에서 수기 수정)")
        self.assertEqual(row["status"], "보완요청")

    def test_07_file_lock_deferred_queue(self):
        # Simulate an exclusive lock by keeping file open in write mode
        with open(self.excel_path, "r+b") as locked_fp:
            # While file is locked, web inserts a new item
            new_item = {
                "year": "2026",
                "seq_no": "99",
                "party": "행안위",
                "requester": "정의원",
                "title": "잠금 중 등록된 자료",
                "details": "잠금 테스트"
            }
            res = self.service.insert_ledger_item(new_item)
            self.assertTrue(res["success"])
            self.assertFalse(res.get("excel_synced"))
            self.assertTrue(res.get("excel_pending"))

            # Check item was queued
            with self.sync_service._pending_lock:
                self.assertGreaterEqual(len(self.sync_service._pending_queue), 1)

        # After locked_fp closes (Excel program closed), flush pending queue
        flushed = self.sync_service.flush_pending_queue()
        self.assertGreaterEqual(flushed, 1)

        # Verify row was written to Excel after unlock
        wb = openpyxl.load_workbook(str(self.excel_path), data_only=True)
        ws = wb["2026"]
        titles = [ws.cell(r, 5).value for r in range(2, ws.max_row + 1)]
        wb.close()
        self.assertIn("잠금 중 등록된 자료", titles)

    def test_08_startup_sync(self):
        # Advance Excel mtime into the future
        future_time = time.time() + 100
        os.utime(self.excel_path, (future_time, future_time))

        # Startup sync check
        self.sync_service.check_and_sync_startup()
        self.assertGreaterEqual(self.sync_service.get_version(), 0)


if __name__ == '__main__':
    unittest.main()
