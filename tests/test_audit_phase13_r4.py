# -*- coding: utf-8 -*-
"""감사 4회차(PROJECT_AUDIT.md R4-xx) 회귀 테스트.

모든 테스트는 임시 디렉터리만 쓴다. 운영 폴더 무변경·운영 DB 미접근은
test_zzz_production_isolation이 확인한다.
"""

from typing import Any, cast
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from db.database_manager import DatabaseManager
from extractors.ledger.discovery import discover_master_excel
from extractors.ledger.dates import LEDGER_COLUMN_ALIASES
from extractors.ledger_parser import load_request_ledger, stamp_ledger_ids
from services.excel_sync_service import ExcelSyncService, QUARANTINE_CONFLICT, QUARANTINE_RETRY_LIMIT
from services.ledger_service import LedgerService


def sleeper():
    """살아 있는 다른 프로세스(PID 락 보유자 흉내)."""
    return subprocess.Popen(
        [sys.executable, "-c", "import time; time.sleep(60)"],
        stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
    )


class Fixture(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="audit_phase13_"))
        self.sys_dir = self.tmp
        self.db_path = self.tmp / "data_requests.db"
        self.json_path = self.tmp / "data_requests.json"
        self.init_db(self.db_path)
        ExcelSyncService._instance = None

    def tearDown(self):
        svc = ExcelSyncService._instance
        if svc is not None:
            svc.stop_file_watcher()
        ExcelSyncService._instance = None
        LedgerService.wait_last_sync(5)
        shutil.rmtree(self.tmp, ignore_errors=True)

    @staticmethod
    def init_db(path):
        path.parent.mkdir(parents=True, exist_ok=True)
        conn = sqlite3.connect(str(path))
        DatabaseManager.init_schema(conn)
        conn.close()

    @staticmethod
    def make_excel(path, headers, rows, sheet="2026"):
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = sheet
        ws.append(headers)
        for row in rows:
            ws.append(row)
        wb.save(path)
        wb.close()
        return path

    @staticmethod
    def sheet_rows(path, sheet="2026"):
        wb = openpyxl.load_workbook(path, data_only=True)
        try:
            ws = wb[sheet]
            header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            return [
                {str(header[c - 1]): ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
                for r in range(2, ws.max_row + 1)
            ]
        finally:
            wb.close()

    def db_rows(self, db_path=None):
        conn = sqlite3.connect(str(db_path or self.db_path))
        conn.row_factory = sqlite3.Row
        try:
            return {r["ledger_id"]: dict(r) for r in conn.execute("SELECT * FROM request_ledger")}
        finally:
            conn.close()


# ---------------------------------------------------------------------------
# R4-03 마스터 엑셀 탐색 단일화
# ---------------------------------------------------------------------------
class TestMasterExcelDiscovery(Fixture):
    HEADERS = ["대장ID", "연번", "의원", "요구자료", "제출"]

    def test_excel_in_workspace_root_is_parsed_by_excel_to_db(self):
        """대장이 상위 워크스페이스에 있어도 쓰기와 읽기가 같은 파일을 쓴다."""
        ws_root = self.tmp / "ws"
        (ws_root / "2026").mkdir(parents=True)
        system = ws_root / "system"
        db = system / "data_requests.db"
        self.init_db(db)
        excel = self.make_excel(ws_root / "260904 국회 요구자료 목록(워크스페이스).xlsx", self.HEADERS,
                                [["REQ-2026-001", "1", "김의원", "워크스페이스 대장 항목", "미제출"]])
        svc = ExcelSyncService(base_dir=system, db_path=db, json_path=system / "data_requests.json")
        self.assertEqual(svc.find_master_excel().resolve(), excel.resolve())
        res = svc.sync_excel_to_db(write_json=False)
        self.assertTrue(res["success"], res)
        self.assertEqual(res["inserted"], 1, "예전에는 파서가 상위 폴더를 보지 않아 0건이었다")
        self.assertIn("REQ-2026-001", self.db_rows(db))

    def test_loose_patterns_are_not_used_in_workspace_root(self):
        """답변 문서가 섞인 원문 폴더에서는 느슨한 패턴으로 대장을 고르지 않는다."""
        ws_root = self.tmp / "ws"
        (ws_root / "2026").mkdir(parents=True)
        system = ws_root / "system"
        system.mkdir()
        self.make_excel(ws_root / "2026_요구자료_답변_통계.xlsx", self.HEADERS, [])
        chosen, _candidates, _pinned = discover_master_excel(system)
        self.assertIsNone(chosen)

    def test_pinned_master_excel_wins(self):
        a = self.make_excel(self.tmp / "국회 요구자료 목록 A.xlsx", self.HEADERS, [])
        self.make_excel(self.tmp / "국회 요구자료 목록 B.xlsx", self.HEADERS, [])
        (self.tmp / "config.json").write_text(json.dumps({"master_excel": a.name}, ensure_ascii=False), encoding="utf-8")
        chosen, candidates, pinned = discover_master_excel(self.tmp)
        self.assertTrue(pinned)
        assert chosen is not None
        self.assertEqual(chosen.resolve(), a.resolve())

    def test_explicit_excel_path_is_honored(self):
        self.make_excel(self.tmp / "국회 요구자료 목록 최신.xlsx", self.HEADERS,
                        [["REQ-2026-001", "1", "갑", "자동 탐색 파일", "미제출"]])
        other = self.make_excel(self.tmp / "sub" / "다른파일.xlsx", self.HEADERS,
                                [["REQ-2026-009", "9", "을", "명시한 파일", "미제출"]])
        items = load_request_ledger(self.tmp, [], merge_db_records=False, db_path=self.db_path,
                                    system_dir=self.tmp, excel_path=other)
        self.assertEqual([i["title"] for i in items], ["명시한 파일"])


# ---------------------------------------------------------------------------
# R4-02 경로 해석이 전역 모듈에 끌려가지 않는다
# ---------------------------------------------------------------------------
class TestLedgerPathResolution(Fixture):
    def test_root_dir_wins_over_imported_pipeline_module(self):
        import extract_and_build_db  # noqa: F401  운영 경로 전역을 가진 모듈이 import된 상태
        from extractors.ledger.paths import _get_active_db_path, _get_active_system_dir
        self.assertEqual(_get_active_db_path(self.tmp), self.tmp / "data_requests.db")
        self.assertEqual(_get_active_system_dir(self.tmp), self.tmp)


# ---------------------------------------------------------------------------
# R4-04 헤더 별칭 SSOT
# ---------------------------------------------------------------------------
class TestHeaderAliasesRoundTrip(Fixture):
    def test_apply_uses_parser_alias_table(self):
        self.assertIs(ExcelSyncService.COLUMN_ALIASES, LEDGER_COLUMN_ALIASES)

    def test_update_writes_alias_headers(self):
        headers = ["대장ID", "연번", "상임위", "요구기관", "질의제목", "제출"]
        excel = self.make_excel(self.tmp / "260904 국회 요구자료 목록(별칭).xlsx", headers,
                                [["REQ-2026-001", "1", "과방위", "A기관", "원래 제목", "미제출"]])
        svc = ExcelSyncService.get_instance(self.tmp)
        self.assertEqual(svc.sync_excel_to_db(write_json=False)["inserted"], 1)
        row = self.db_rows()["REQ-2026-001"]
        self.assertEqual((row["requester"], row["title"], row["party"]), ("A기관", "원래 제목", "과방위"))

        ledger = LedgerService(self.db_path, self.json_path, self.tmp)
        res = ledger.update_ledger_item("REQ-2026-001", {"requester": "B기관", "title": "바뀐 제목", "party": "법사위"})
        self.assertTrue(res["success"], res)
        self.assertTrue(res["excel_synced"], res)
        sheet = self.sheet_rows(excel)[0]
        self.assertEqual((sheet["요구기관"], sheet["질의제목"], sheet["상임위"]), ("B기관", "바뀐 제목", "법사위"))

        # 다음 엑셀→DB 동기화가 웹 수정을 옛값으로 되돌리지 않는다.
        time.sleep(0.05)
        os.utime(excel, None)
        svc.sync_excel_to_db(write_json=False)
        self.assertEqual(self.db_rows()["REQ-2026-001"]["requester"], "B기관")


# ---------------------------------------------------------------------------
# R4-07 0건 파싱 신호
# ---------------------------------------------------------------------------
class TestParsedZeroIsReported(Fixture):
    HEADERS = ["대장ID", "연번", "의원", "요구자료", "제출"]

    def test_zero_rows_with_excel_backed_db_is_not_success(self):
        excel = self.make_excel(self.tmp / "260904 국회 요구자료 목록(0건).xlsx", self.HEADERS,
                                [["REQ-2026-001", "1", "김의원", "엑셀 항목", "미제출"]])
        svc = ExcelSyncService.get_instance(self.tmp)
        self.assertEqual(svc.sync_excel_to_db(write_json=False)["inserted"], 1)

        # 헤더 인식이 깨진 파일(시트가 비었음)로 바뀐다.
        self.make_excel(excel, ["알수없는열"], [])
        res = svc.sync_excel_to_db(write_json=False)
        self.assertFalse(res["success"])
        self.assertEqual(res["reason"], "parsed_zero")
        self.assertEqual(svc.last_excel_sync()["reason"], "parsed_zero")
        self.assertNotEqual(self.db_rows()["REQ-2026-001"]["status"], "[삭제]")

    def test_empty_new_ledger_is_still_success(self):
        self.make_excel(self.tmp / "260904 국회 요구자료 목록(새 대장).xlsx", self.HEADERS, [])
        res = ExcelSyncService.get_instance(self.tmp).sync_excel_to_db(write_json=False)
        self.assertTrue(res["success"], res)


# ---------------------------------------------------------------------------
# R4-11 보류 큐 손상 보존
# ---------------------------------------------------------------------------
class TestCorruptPendingQueue(Fixture):
    def test_corrupt_queue_is_preserved_and_reported(self):
        queue = self.tmp / ".excel_pending_queue.json"
        queue.write_text("[{broken", encoding="utf-8")
        svc = ExcelSyncService(base_dir=self.tmp)
        health = svc.queue_health()
        self.assertIsNotNone(health)
        backups = list(self.tmp.glob(".excel_pending_queue.corrupt.*.json"))
        self.assertEqual(len(backups), 1)
        self.assertEqual(backups[0].read_text(encoding="utf-8"), "[{broken")
        self.assertEqual(svc.pending_count(), 0)

    def test_healthy_queue_reports_none(self):
        (self.tmp / ".excel_pending_queue.json").write_text("[]", encoding="utf-8")
        self.assertIsNone(ExcelSyncService(base_dir=self.tmp).queue_health())


# ---------------------------------------------------------------------------
# R4-13 스탬핑 백업 위치
# ---------------------------------------------------------------------------
class TestStampBackupLocation(Fixture):
    def test_stamp_backup_goes_to_base_dir(self):
        excel = self.make_excel(self.tmp / "elsewhere" / "대장.xlsx", ["연번", "의원", "요구자료"],
                                [["1", "김의원", "ID 없는 행"]])
        system = self.tmp / "system"
        system.mkdir()
        written = stamp_ledger_ids(excel, [], base_dir=system)
        self.assertEqual(written, 1)
        self.assertTrue(list((system / ".ledger_backup").glob("대장_*.xlsx")))
        self.assertFalse((excel.parent / ".ledger_backup").exists())


# ---------------------------------------------------------------------------
# R4-05 웹 서버 단일 인스턴스
# ---------------------------------------------------------------------------
class TestServerInstanceLock(Fixture):
    def test_live_holder_blocks_second_server(self):
        from web.instance_lock import ServerInstanceLock, LOCK_NAME
        holder = sleeper()
        try:
            (self.tmp / LOCK_NAME).write_text(json.dumps({"pid": holder.pid, "port": 8123}), encoding="utf-8")
            ok, info = ServerInstanceLock(self.tmp).acquire()
            self.assertFalse(ok)
            assert info is not None
            self.assertEqual(info["port"], 8123)
        finally:
            holder.kill()
            holder.wait()

    def test_dead_holder_lock_is_reclaimed_and_released(self):
        from web.instance_lock import ServerInstanceLock, LOCK_NAME
        dead = sleeper()
        dead.kill()
        dead.wait()
        (self.tmp / LOCK_NAME).write_text(json.dumps({"pid": dead.pid, "port": 8080}), encoding="utf-8")
        lock = ServerInstanceLock(self.tmp)
        ok, _ = lock.acquire()
        self.assertTrue(ok)
        lock.set_port(8099)
        self.assertEqual(lock.read(), {"pid": os.getpid(), "port": 8099})
        lock.release()
        self.assertFalse((self.tmp / LOCK_NAME).exists())

    def test_run_server_exits_when_already_running(self):
        import web_server
        from web.instance_lock import LOCK_NAME
        holder = sleeper()
        try:
            (self.tmp / LOCK_NAME).write_text(json.dumps({"pid": holder.pid, "port": 8123}), encoding="utf-8")
            with mock.patch.object(web_server, "BASE_DIR", self.tmp), \
                 mock.patch.object(web_server, "init_db") as init_db, \
                 mock.patch("webbrowser.open"):
                with self.assertRaises(SystemExit) as cm:
                    web_server.run_server(port=8123)
            self.assertEqual(cm.exception.code, web_server.ALREADY_RUNNING_EXIT_CODE)
            init_db.assert_not_called()
        finally:
            holder.kill()
            holder.wait()


# ---------------------------------------------------------------------------
# R4-08 extract_and_build_db 직접 실행도 파이프라인 락을 존중한다
# ---------------------------------------------------------------------------
class TestBuildRespectsPipelineLock(Fixture):
    def test_direct_build_refuses_while_other_process_holds_lock(self):
        import extract_and_build_db as builder
        holder = sleeper()
        try:
            (self.tmp / ".pipeline.lock").write_text(str(holder.pid), encoding="utf-8")
            with mock.patch.object(builder, "DB_PATH", self.db_path), \
                 mock.patch.object(builder, "_main_locked") as body:
                result = builder.main()
            self.assertFalse(result["success"])
            body.assert_not_called()
        finally:
            holder.kill()
            holder.wait()

    def test_direct_build_takes_and_releases_lock(self):
        import extract_and_build_db as builder
        seen = {}

        def fake_body():
            seen["locked"] = (self.tmp / ".pipeline.lock").exists()
            return {"success": True}

        with mock.patch.object(builder, "DB_PATH", self.db_path), \
             mock.patch.object(builder, "_main_locked", side_effect=fake_body):
            self.assertTrue(builder.main()["success"])
        self.assertTrue(seen["locked"])
        self.assertFalse((self.tmp / ".pipeline.lock").exists())


# ---------------------------------------------------------------------------
# R4-10 락 획득 순서: self.lock → _failed_lock → _pending_lock
# ---------------------------------------------------------------------------
class OrderedLock:
    RANK = {"lock": 0, "failed": 1, "pending": 2}

    def __init__(self, name, registry):
        self.name = name
        self.registry = registry
        self._lock = threading.Lock()

    def acquire(self, *args, **kwargs):
        held = self.registry["held"].setdefault(threading.get_ident(), [])
        for other in held:
            if self.RANK[other] > self.RANK[self.name]:
                self.registry["violations"].append(f"{other} -> {self.name}")
        ok = self._lock.acquire(*args, **kwargs)
        if ok:
            held.append(self.name)
        return ok

    def release(self):
        self.registry["held"][threading.get_ident()].remove(self.name)
        self._lock.release()

    __enter__ = acquire

    def __exit__(self, *exc):
        self.release()


class TestLockOrder(Fixture):
    HEADERS = ["대장ID", "연번", "의원", "요구자료", "제출"]

    def test_no_path_acquires_locks_in_reverse_order(self):
        excel = self.make_excel(self.tmp / "260904 국회 요구자료 목록(락).xlsx", self.HEADERS,
                                [["REQ-2026-001", "1", "김의원", "락 순서", "미제출"]])
        svc = ExcelSyncService.get_instance(self.tmp)
        registry = {"held": {}, "violations": []}
        svc.lock = cast(Any, OrderedLock("lock", registry))
        svc._failed_lock = cast(Any, OrderedLock("failed", registry))
        svc._pending_lock = cast(Any, OrderedLock("pending", registry))

        svc.sync_excel_to_db(write_json=False)
        item = dict(self.db_rows()["REQ-2026-001"])
        # 반영 실패(일반 예외) → self.lock 안에서 큐 적재
        with mock.patch.object(svc, "_apply_to_excel", side_effect=RuntimeError("disk")):
            svc.sync_item_to_excel(dict(item, title="큐 적재"), action="update")
        self.assertEqual(svc.pending_count(), 1)
        svc.flush_pending_queue()
        # 충돌 → self.lock 안에서 격리
        conflict_item = dict(item, ledger_id="REQ-2026-001", title="전혀 다른 항목", requester="다른의원",
                             _expected={"title": "또 다른", "requester": "누군가"})
        svc.sync_item_to_excel(conflict_item, action="update")
        quarantined = svc.list_quarantined()
        self.assertTrue(quarantined)
        svc.retry_quarantined(quarantined[0]["operation_id"])
        svc._quarantine_pending_operations([svc._new_operation(item, "update")], QUARANTINE_RETRY_LIMIT)
        for op in svc.list_quarantined():
            svc.dismiss_quarantined(op["operation_id"])
        self.assertTrue(excel.exists())
        self.assertEqual(registry["violations"], [])


# ---------------------------------------------------------------------------
# 격리 항목 관리 (4.1-3)
# ---------------------------------------------------------------------------
class TestQuarantineManagement(Fixture):
    def write_failed(self, ops):
        (self.tmp / ".excel_pending_failed.json").write_text(json.dumps(ops, ensure_ascii=False), encoding="utf-8")

    def op(self, op_id, reason, lid):
        return {"operation_id": op_id, "action": "update", "attempts": 5, "timestamp": 0,
                "item": {"ledger_id": lid, "title": f"{lid} 제목", "_changed_fields": ["title"]},
                "quarantine_reason": reason, "quarantined_at": "2026-09-17 10:00:00"}

    def test_list_retry_dismiss(self):
        self.write_failed([self.op("a1", QUARANTINE_RETRY_LIMIT, "REQ-2026-001"),
                           self.op("b2", QUARANTINE_CONFLICT, "REQ-2026-002")])
        svc = ExcelSyncService(base_dir=self.tmp)
        listed = svc.list_quarantined()
        self.assertEqual([x["operation_id"] for x in listed], ["a1", "b2"])
        self.assertEqual(svc.protected_ledger_ids(), {"REQ-2026-001"})

        res = svc.retry_quarantined("a1")
        self.assertTrue(res["success"], res)
        self.assertEqual(svc.pending_count(), 1)
        self.assertEqual(svc.failed_count(), 1)
        queued = json.loads((self.tmp / ".excel_pending_queue.json").read_text(encoding="utf-8"))[0]
        self.assertEqual(queued["attempts"], 0)
        self.assertNotIn("quarantine_reason", queued)
        # 큐로 옮겨도 보호는 유지된다.
        self.assertEqual(svc.protected_ledger_ids(), {"REQ-2026-001"})

        res = svc.dismiss_quarantined("b2")
        self.assertTrue(res["success"], res)
        self.assertFalse((self.tmp / ".excel_pending_failed.json").exists())
        conn = sqlite3.connect(str(self.db_path))
        try:
            actions = [r[0] for r in conn.execute("SELECT action FROM ledger_history WHERE ledger_id='REQ-2026-002'")]
        finally:
            conn.close()
        self.assertIn("EXCEL_QUARANTINE_DISMISSED", actions)

    def test_unknown_operation(self):
        svc = ExcelSyncService(base_dir=self.tmp)
        self.assertFalse(svc.retry_quarantined("nope")["success"])
        self.assertFalse(svc.dismiss_quarantined("nope")["success"])


# ---------------------------------------------------------------------------
# 낙관적 잠금 (4.1-1)
# ---------------------------------------------------------------------------
class TestOptimisticLocking(Fixture):
    def test_stale_expected_updated_at_is_rejected(self):
        ledger = LedgerService(self.db_path, self.json_path, self.tmp)
        created = ledger.insert_ledger_item({"title": "잠금 시험", "requester": "김의원", "year": "2026"})
        self.assertTrue(created["success"], created)
        lid = created["item"]["ledger_id"]
        original = self.db_rows()[lid]["updated_at"]

        stale = ledger.update_ledger_item(lid, {"title": "남의 수정 위에 덮기", "expected_updated_at": "1999-01-01 00:00:00"})
        self.assertFalse(stale["success"])
        self.assertTrue(stale["conflict"])
        self.assertEqual(self.db_rows()[lid]["title"], "잠금 시험")

        time.sleep(1.1)  # updated_at은 초 단위
        ok = ledger.update_ledger_item(lid, {"title": "정상 수정", "expected_updated_at": original})
        self.assertTrue(ok["success"], ok)
        self.assertEqual(self.db_rows()[lid]["title"], "정상 수정")
        self.assertEqual(ok["updated_at"], self.db_rows()[lid]["updated_at"])

        # 옛 시각으로 다시 보내면 거절된다.
        again = ledger.update_ledger_item(lid, {"title": "두 번째", "expected_updated_at": original})
        self.assertTrue(again.get("conflict"))

    def test_without_expected_value_behaves_as_before(self):
        ledger = LedgerService(self.db_path, self.json_path, self.tmp)
        lid = ledger.insert_ledger_item({"title": "기존 호출", "requester": "이의원", "year": "2026"})["item"]["ledger_id"]
        self.assertTrue(ledger.update_ledger_item(lid, {"note": "메모"})["success"])

    def test_same_second_writes_conflict(self):
        """같은 초 안의 두 번째 쓰기는 충돌로 거절되어야 한다. (ISSUE-001)

        updated_at이 초 단위면 두 쓰기의 버전 문자열이 같아져 낙관적 잠금이
        동작하지 않고 조용히 덮어쓴다.
        """
        ledger = LedgerService(self.db_path, self.json_path, self.tmp)
        created = ledger.insert_ledger_item({"title": "동시 수정", "requester": "김의원", "year": "2026"})
        self.assertTrue(created["success"], created)
        lid = created["item"]["ledger_id"]
        original = self.db_rows()[lid]["updated_at"]

        first = ledger.update_ledger_item(lid, {"title": "첫 수정", "expected_updated_at": original})
        self.assertTrue(first["success"], first)
        self.assertNotEqual(self.db_rows()[lid]["updated_at"], original)

        second = ledger.update_ledger_item(lid, {"title": "두 번째", "expected_updated_at": original})
        self.assertFalse(second["success"])
        self.assertTrue(second.get("conflict"))
        self.assertEqual(self.db_rows()[lid]["title"], "첫 수정")


class TestWriteContentionRetry(Fixture):
    def test_begin_immediate_retries_then_succeeds(self):
        """BEGIN 경합 1회는 재시도 후 성공한다."""
        cur = mock.MagicMock()
        cur.execute.side_effect = [sqlite3.OperationalError("database is locked"), None]
        LedgerService._begin_immediate(cur)
        self.assertEqual(cur.execute.call_count, 2)

    def test_begin_immediate_gives_up(self):
        """계속 잠겨 있으면 OperationalError를 그대로 올린다."""
        cur = mock.MagicMock()
        cur.execute.side_effect = sqlite3.OperationalError("database is locked")
        with self.assertRaises(sqlite3.OperationalError):
            LedgerService._begin_immediate(cur, attempts=2)
        self.assertEqual(cur.execute.call_count, 2)

    def test_update_succeeds_under_concurrent_writer(self):
        """RESERVED 락 보유 중의 update는 대기 후 성공한다."""
        ledger = LedgerService(self.db_path, self.json_path, self.tmp)
        lid = ledger.insert_ledger_item(
            {"title": "경합 시험", "requester": "김의원", "year": "2026"})["item"]["ledger_id"]

        release = threading.Event()

        def hold_lock():
            blocker = sqlite3.connect(str(self.db_path), timeout=30.0)
            try:
                blocker.execute("BEGIN IMMEDIATE")
                blocker.execute("UPDATE request_ledger SET note = ? WHERE ledger_id = ?",
                                ("차단자", lid))
                release.wait(timeout=10)
                blocker.commit()
            finally:
                blocker.close()

        worker = threading.Thread(target=hold_lock, daemon=True)
        worker.start()
        try:
            time.sleep(0.3)  # 차단자가 RESERVED를 잡을 시간을 준다.
            res = ledger.update_ledger_item(lid, {"title": "경합 후 수정"})
            self.assertTrue(res.get("success"), res)
        finally:
            release.set()
            worker.join(timeout=15)
        self.assertEqual(self.db_rows()[lid]["title"], "경합 후 수정")


# ---------------------------------------------------------------------------
# R4-09 검색 폴백
# ---------------------------------------------------------------------------
class TestFallbackSearch(Fixture):
    def setUp(self):
        super().setUp()
        conn = sqlite3.connect(str(self.db_path))
        rows = [
            ("D1", "2026", "딥페이크 대응 현황", "김의원", "본문 하나", "# 전문 하나"),
            ("D2", "2025", "텔레그램 시정요구", "이의원", "본문 둘", "# 전문 둘"),
        ]
        conn.executemany(
            "INSERT INTO documents (doc_id, year, title, requester, answer_full_text, full_markdown) VALUES (?, ?, ?, ?, ?, ?)",
            rows,
        )
        conn.commit()
        conn.close()
        self.ledger = LedgerService(self.db_path, self.json_path, self.tmp)

    def test_choseong_query_scans_only_short_columns(self):
        statements = []
        original = self.ledger.get_conn

        def traced(*args, **kwargs):
            conn = original(*args, **kwargs)
            conn.set_trace_callback(statements.append)
            return conn

        with mock.patch.object(self.ledger, "get_conn", side_effect=traced):
            docs = self.ledger.query_documents(kw="ㄷㅍㅇㅋ")
        self.assertEqual([d["doc_id"] for d in docs], ["D1"])
        self.assertEqual(docs[0]["full_markdown"], "# 전문 하나")
        scans = [s for s in statements if "_rid" in s and "IN (" not in s]
        self.assertTrue(scans)
        for s in scans:
            self.assertNotIn("full_markdown", s)
            self.assertNotIn("answer_full_text", s, "초성 질의는 본문 열을 읽지 않는다")
            self.assertNotIn("l.*", s)

    def test_fallback_keeps_order_and_slim(self):
        docs = self.ledger.query_documents(kw="ㅅ", slim=True)  # 한 글자 초성 → 폴백
        self.assertEqual([d["doc_id"] for d in docs], ["D2"])
        self.assertNotIn("full_markdown", docs[0])


# ---------------------------------------------------------------------------
# ledger_history 보존 정리 도구 (4.1-2)
# ---------------------------------------------------------------------------
def load_archive_tool():
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "archive_ledger_history", SCRIPTS_DIR.parent / "tools" / "archive_ledger_history.py")
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class TestLedgerHistoryArchive(Fixture):
    def seed(self, rows):
        conn = sqlite3.connect(str(self.db_path))
        conn.executemany(
            "INSERT INTO ledger_history (ledger_id, action, changed_at, snapshot) VALUES (?, ?, ?, '{}')", rows)
        conn.commit()
        conn.close()

    def remaining(self):
        conn = sqlite3.connect(str(self.db_path))
        try:
            return [tuple(r) for r in conn.execute("SELECT ledger_id, action, changed_at FROM ledger_history ORDER BY history_id")]
        finally:
            conn.close()

    def test_keeps_identity_markers_and_recent_rows(self):
        tool = load_archive_tool()
        self.seed([
            ("REQ-2024-001", "INSERT", "2024-01-01 09:00:00"),
            ("REQ-2024-001", "UPDATE", "2024-01-02 09:00:00"),
            ("REQ-2024-001", "EXCEL_APPLIED", "2024-01-02 09:00:01"),
            ("REQ-2024-001", "UPDATE", "2024-02-01 09:00:00"),
            ("REQ-2024-002", "EXCEL_SYNC_INSERT", "2024-03-01 09:00:00"),
            ("REQ-2024-002", "EXCEL_SYNC_UPDATE", "2024-03-02 09:00:00"),
            ("REQ-2026-001", "UPDATE", "2026-09-01 09:00:00"),
            ("REQ-2024-003", "UPDATE", "알 수 없음"),
        ])
        now = __import__("datetime").datetime(2026, 9, 17)
        dry = tool.archive_history(self.db_path, self.tmp / "backup", days=365, apply=False, now=now)
        self.assertEqual((dry["total"], dry["archive"]), (8, 3))
        self.assertEqual(len(self.remaining()), 8, "점검 모드는 DB를 바꾸지 않는다")

        res = tool.archive_history(self.db_path, self.tmp / "backup", days=365, apply=True, now=now)
        self.assertTrue(res["applied"])
        self.assertEqual(self.remaining(), [
            ("REQ-2024-001", "INSERT", "2024-01-01 09:00:00"),
            ("REQ-2024-001", "EXCEL_APPLIED", "2024-01-02 09:00:01"),
            ("REQ-2024-002", "EXCEL_SYNC_UPDATE", "2024-03-02 09:00:00"),
            ("REQ-2026-001", "UPDATE", "2026-09-01 09:00:00"),
            ("REQ-2024-003", "UPDATE", "알 수 없음"),
        ])
        lines = Path(res["archive_file"]).read_text(encoding="utf-8").strip().splitlines()
        self.assertEqual(len(lines), 3)
        # 웹 전용 판정이 바뀌지 않는다.
        conn = sqlite3.connect(str(self.db_path))
        try:
            self.assertEqual(ExcelSyncService._web_only_ledger_ids(conn.cursor()), set())
        finally:
            conn.close()


# ---------------------------------------------------------------------------
# R4-14 정적 게이트: 미정의 이름 (분할 잔재가 NameError로 숨어 있지 않게)
# ---------------------------------------------------------------------------
try:
    from pyflakes import api as _pyflakes_api, messages as _pyflakes_messages
    from pyflakes.reporter import Reporter as _PyflakesReporter
except ImportError:  # pragma: no cover
    _pyflakes_api = None
    _pyflakes_messages = None
    _PyflakesReporter = object


@unittest.skipUnless(_pyflakes_api, "pyflakes가 없어 정적 검사를 건너뜁니다")
class TestNoUndefinedNames(unittest.TestCase):
    def test_scripts_and_tools_have_no_undefined_names(self):
        assert _pyflakes_api is not None
        assert _pyflakes_messages is not None
        found = []

        class Collect(cast(Any, _PyflakesReporter)):
            def __init__(self):
                pass

            def flake(self, message):
                assert _pyflakes_messages is not None
                if isinstance(message, (_pyflakes_messages.UndefinedName, _pyflakes_messages.UndefinedLocal)):
                    found.append(str(message))

            def unexpectedError(self, filename, msg):
                found.append(f"{filename}: {msg}")

            def syntaxError(self, filename, msg, lineno, offset, text):
                found.append(f"{filename}:{lineno}: {msg}")

        root = SCRIPTS_DIR.parent
        for folder in ("scripts", "tools"):
            for path in sorted((root / folder).rglob("*.py")):
                if "__pycache__" in path.parts:
                    continue
                _pyflakes_api.check(path.read_text(encoding="utf-8"), str(path), Collect())
        self.assertEqual(found, [])


# ---------------------------------------------------------------------------
# 4.3 추정 항목 후속 (2026-09-17 운영 대장 점검 후 보강)
# ---------------------------------------------------------------------------
class TestEstimateFollowUps(Fixture):
    def test_invalid_calendar_dates_are_not_fabricated(self):
        from extractors.ledger.dates import normalize_ledger_date as n
        self.assertEqual(n("0528", "2018"), "2018-05-28")
        self.assertEqual(n("9. 7.", "2021"), "2021-09-07")
        self.assertEqual(n("250115", "2025"), "2025-01-15")
        self.assertEqual(n("20240229", "2024"), "2024-02-29")
        self.assertEqual(n("0231", "2026"), "0231", "달력에 없는 날짜는 원래 값 유지")
        self.assertEqual(n("251345", "2025"), "251345")
        self.assertEqual(n("991231", "2026"), "991231", "2099년으로 읽지 않는다")
        self.assertEqual(n("20250230", "2025"), "20250230")

    def test_daily_snapshot_survives_rolling_backups(self):
        from extractors.ledger import workbook
        excel = self.make_excel(self.tmp / "대장.xlsx", ["연번"], [["1"]])
        with mock.patch.dict(workbook._LAST_BACKUP_AT, clear=True):
            for _ in range(7):
                workbook.backup_master_excel(excel, self.tmp, force=True)
        rolling = list((self.tmp / ".ledger_backup").glob("대장_*.xlsx"))
        daily = list((self.tmp / ".ledger_backup" / "daily").glob("대장_*.xlsx"))
        self.assertEqual(len(rolling), workbook.BACKUP_KEEP)
        self.assertEqual(len(daily), 1, "그날 첫 덮어쓰기 직전 본은 롤링과 별도로 남는다")

    def test_multiple_candidates_are_reported(self):
        headers = ["대장ID", "연번", "의원", "요구자료", "제출"]
        self.make_excel(self.tmp / "국회 요구자료 목록 원본.xlsx", headers, [])
        svc = ExcelSyncService(base_dir=self.tmp)
        self.assertEqual(svc.master_excel_ambiguity(), [])
        time.sleep(0.05)
        self.make_excel(self.tmp / "국회 요구자료 목록 사본.xlsx", headers, [])
        self.assertEqual(len(svc.master_excel_ambiguity()), 2)
        (self.tmp / "config.json").write_text(json.dumps({"master_excel": "국회 요구자료 목록 원본.xlsx"},
                                                         ensure_ascii=False), encoding="utf-8")
        self.assertEqual(svc.master_excel_ambiguity(), [], "고정하면 경고하지 않는다")


# ---------------------------------------------------------------------------
# 투입 파일이 사용 중일 때 (2026-09-17 운영 반영 중 발견)
# ---------------------------------------------------------------------------
class TestLockedInboxFile(Fixture):
    def setUp(self):
        super().setUp()
        import add_documents_smart
        self.mod = add_documents_smart
        self.workspace = self.tmp / "ws"
        (self.workspace / "2026").mkdir(parents=True)
        self.inbox = self.tmp / "inbox"
        self.inbox.mkdir()

    def test_in_use_file_is_skipped_not_crashing(self):
        busy = self.inbox / "260917 김의원 요구자료(1)_사용중.hwp"
        free = self.inbox / "260917 이의원 요구자료(2).hwp"
        busy.write_bytes(b"busy")
        free.write_bytes(b"free")
        skipped = []
        with mock.patch.object(self.mod, "ROOT_DIR", self.workspace), \
             mock.patch.object(self.mod, "_is_in_use", side_effect=lambda p: p.name == busy.name):
            moved = self.mod.copy_or_move_files([str(busy), str(free)], move=True, skipped=skipped)
        self.assertEqual([p.name for p in moved], [free.name])
        self.assertEqual([name for name, _ in skipped], [busy.name])
        self.assertTrue(busy.exists())
        self.assertFalse((self.workspace / "2026" / busy.name).exists())

    def test_partial_move_copy_is_rolled_back(self):
        src = self.inbox / "260917 박의원 요구자료(3).hwp"
        src.write_bytes(b"payload")

        def copy_then_fail(s, d):
            shutil.copy2(s, d)
            raise PermissionError("원본 삭제 실패(사용 중)")

        skipped = []
        with mock.patch.object(self.mod, "ROOT_DIR", self.workspace), \
             mock.patch.object(self.mod, "_is_in_use", return_value=False), \
             mock.patch.object(self.mod.shutil, "move", side_effect=copy_then_fail):
            moved = self.mod.copy_or_move_files([str(src)], move=True, skipped=skipped)
        self.assertEqual(moved, [])
        self.assertEqual(len(skipped), 1)
        self.assertTrue(src.exists())
        self.assertFalse((self.workspace / "2026" / src.name).exists(), "실패한 이동이 만든 사본은 되돌린다")


# ---------------------------------------------------------------------------
# R4-15d 정적 파일 허용 목록
# ---------------------------------------------------------------------------
class TestPublicStaticAllowlist(unittest.TestCase):
    def test_arbitrary_html_is_not_served(self):
        from web_server import RequestLedgerHandler
        self.assertFalse(RequestLedgerHandler.is_public_static("/내보낸_보고서.html"))
        self.assertTrue(RequestLedgerHandler.is_public_static("/자료요구_통합검색_대시보드.html"))


# ---------------------------------------------------------------------------
# 대시보드 JS (Node 실행)
# ---------------------------------------------------------------------------
try:  # pytest(패키지) / unittest discover -s tests(최상위 모듈) 둘 다 지원
    from tests.test_audit_phase10_dashboard import NODE, TEMPLATE, extract_js  # noqa: E402
except ImportError:  # pragma: no cover
    from test_audit_phase10_dashboard import NODE, TEMPLATE, extract_js  # noqa: E402


def run_js(names, body, prelude=""):
    src = TEMPLATE.read_text(encoding="utf-8")
    code = prelude + "\n" + "\n".join(extract_js(src, n) for n in names) + "\n" + body
    with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False, encoding="utf-8") as fp:
        fp.write(code)
        path = fp.name
    try:
        assert NODE is not None
        assert NODE is not None
        proc = subprocess.run([NODE, path], capture_output=True, text=True, encoding="utf-8",
                              timeout=60, stdin=subprocess.DEVNULL)
    finally:
        Path(path).unlink(missing_ok=True)
    if proc.returncode != 0:
        raise AssertionError(f"node 실행 실패:\n{proc.stderr}")
    return json.loads(proc.stdout.strip().splitlines()[-1])


@unittest.skipUnless(NODE, "node가 없어 대시보드 JS 실행 테스트를 건너뜁니다")
class TestDashboardR4(unittest.TestCase):
    def test_highlight_marks_every_text_node(self):
        """R4-12: g 플래그 lastIndex 때문에 번갈아 빠지던 강조."""
        prelude = """
const NodeFilter = { SHOW_TEXT: 4 };
const texts = ['텔레그램 대응', '다시 텔레그램', '텔레그램과 텔레그램', '관련 없음'].map(v => ({ nodeValue: v }));
texts.forEach(t => { t.parentNode = { nodeName: 'P', replaced: null, replaceChild(frag) { this.replaced = frag; } }; });
const document = {
  createTreeWalker() { let i = 0; return { nextNode: () => texts[i++] || null }; },
  createDocumentFragment() { return { kids: [], appendChild(n) { this.kids.push(n); } }; },
  createElement(tag) { return { tag, className: '', textContent: '' }; },
  createTextNode(text) { return { text }; },
};
"""
        out = run_js(["isChoseongOnly", "applyDomHighlight"], """
applyDomHighlight({}, ['텔레그램']);
console.log(JSON.stringify(texts.map(t => {
  const f = t.parentNode.replaced;
  return f ? { marks: f.kids.filter(k => k.tag === 'mark').map(k => k.textContent), plain: f.kids.filter(k => k.text !== undefined).map(k => k.text) } : null;
})));
""", prelude)
        self.assertEqual(out[0]["marks"], ["텔레그램"])
        self.assertEqual(out[1]["marks"], ["텔레그램"])
        self.assertEqual(out[2]["marks"], ["텔레그램", "텔레그램"])
        self.assertEqual(out[2]["plain"], ["과 "])
        self.assertIsNone(out[3])

    def test_offline_write_block_and_local_merge(self):
        """R4-06 / R4-15a."""
        prelude = """
const window = { location: { protocol: 'file:' } };
const store = { kocsc_local_ledger: JSON.stringify([{ ledger_id: 'LOCAL-2026-b' }, { ledger_id: 'LOCAL-2026-a' }, { ledger_id: 'REQ-2026-001' }]) };
const localStorage = { getItem: k => store[k] || null, setItem: (k, v) => { store[k] = v; } };
"""
        out = run_js(["isLocalOnlyLedgerId", "ledgerWriteBlockReason", "mergeLocalLedger"], """
const offlineServer = ledgerWriteBlockReason({ ledger_id: 'REQ-2026-001' });
const offlineLocal = ledgerWriteBlockReason({ ledger_id: 'LOCAL-2026-a' });
window.location.protocol = 'http:';
const online = ledgerWriteBlockReason({ ledger_id: 'REQ-2026-001' });
const merged = mergeLocalLedger([{ ledger_id: 'REQ-2026-001', title: 'server' }]).map(x => x.ledger_id);
console.log(JSON.stringify({ offlineServer, offlineLocal, online, merged }));
""", prelude)
        self.assertTrue(out["offlineServer"])
        self.assertEqual(out["offlineLocal"], "")
        self.assertEqual(out["online"], "")
        self.assertEqual(out["merged"], ["LOCAL-2026-b", "LOCAL-2026-a", "REQ-2026-001"])

    def test_status_banner_reports_silent_failures(self):
        """R4-07 / R4-11 / R4-15e."""
        out = run_js(["excelStatusMessage"], """
console.log(JSON.stringify({
  ok: excelStatusMessage({ excel_pending: 0, excel_failed: 0, excel_last_sync: { success: true } }),
  zero: excelStatusMessage({ excel_last_sync: { success: false, reason: 'parsed_zero', message: '읽지 못함' } }),
  corrupt: excelStatusMessage({ excel_queue_corrupt: { backup: 'q.corrupt.json' } }),
  startup: excelStatusMessage({ excel_startup_error: '엑셀 시작 동기화 실패: x' }),
  failed: excelStatusMessage({ excel_failed: 2 }),
  ambiguous: excelStatusMessage({ excel_master_candidates: ['사본.xlsm', '원본.xlsm'] }),
}));
""")
        self.assertIsNone(out["ok"])
        self.assertEqual(out["zero"]["level"], "error")
        self.assertIn("읽지 못함", out["zero"]["text"])
        self.assertIn("q.corrupt.json", out["corrupt"]["text"])
        self.assertIn("시작 동기화", out["startup"]["text"])
        self.assertTrue(out["failed"]["canManageQuarantine"])
        self.assertFalse(out["zero"]["canManageQuarantine"])
        self.assertEqual(out["ambiguous"]["level"], "warn")
        self.assertIn("사본.xlsm", out["ambiguous"]["text"])

    def test_search_cache_invalidates_same_length_change(self):
        """R4-15b."""
        out = run_js(["searchCacheSignature", "searchCache"], """
const item = {};
let calls = 0;
const lower = (x) => { calls += 1; return x.toLowerCase(); };
const a = searchCache(item, 'k', 'ABCD', lower);
const b = searchCache(item, 'k', 'ABCD', lower);
const c = searchCache(item, 'k', 'WXYZ', lower);
const long1 = 'x'.repeat(300) + 'A' + 'y'.repeat(300);
const long2 = 'x'.repeat(300) + 'B' + 'y'.repeat(300);
searchCache(item, 'L', long1, lower);
const d = searchCache(item, 'L', long2, lower);
console.log(JSON.stringify({ a, b, c, d: d.includes('b'), calls }));
""")
        self.assertEqual((out["a"], out["b"], out["c"]), ("abcd", "abcd", "wxyz"))
        self.assertTrue(out["d"])
        self.assertEqual(out["calls"], 4)


if __name__ == "__main__":
    unittest.main()
