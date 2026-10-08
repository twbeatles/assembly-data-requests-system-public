# -*- coding: utf-8 -*-
"""감사 5회차(PROJECT_AUDIT.md 5회차 ISSUE-001~003) 회귀 테스트.

모든 테스트는 임시 디렉터리만 쓴다. 운영 폴더 무변경·운영 DB 미접근은
test_zzz_production_isolation이 확인한다.
"""

from typing import Any, cast
import shutil
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from db.database_manager import DatabaseManager


HISTORY_COLS_6 = ("ledger_id", "action", "changed_at", "snapshot",
                  "changed_fields", "actor")


def _make_src_db(path, rows, old_schema=False):
    """rows: (ledger_id, action, changed_at, snapshot[, changed_fields, actor])."""
    conn = sqlite3.connect(str(path))
    try:
        if old_schema:
            conn.execute(
                "CREATE TABLE ledger_history (history_id INTEGER PRIMARY KEY AUTOINCREMENT,"
                " ledger_id TEXT, action TEXT, changed_at TEXT, snapshot TEXT)"
            )
            conn.executemany(
                "INSERT INTO ledger_history (ledger_id, action, changed_at, snapshot)"
                " VALUES (?, ?, ?, ?)",
                [r[:4] for r in rows],
            )
        else:
            DatabaseManager.init_schema(conn)
            conn.executemany(
                "INSERT INTO ledger_history (ledger_id, action, changed_at, snapshot,"
                " changed_fields, actor) VALUES (?, ?, ?, ?, ?, ?)",
                rows,
            )
        conn.commit()
    finally:
        conn.close()


def _read_dest_rows(conn):
    return conn.execute(
        "SELECT ledger_id, action, changed_at, snapshot, changed_fields, actor"
        " FROM ledger_history ORDER BY history_id"
    ).fetchall()


class TestCopyLedgerHistoryKeepsNewColumns(unittest.TestCase):
    """ISSUE-001: changed_fields·actor가 이관돼야 한다."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="audit_phase14_"))
        self.src = self.tmp / "src.db"
        self.dest_conn = sqlite3.connect(str(self.tmp / "dest.db"))
        DatabaseManager.init_schema(self.dest_conn)

    def tearDown(self):
        self.dest_conn.close()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_six_columns_are_copied(self):
        rows = [
            ("REQ-2026-0001", "UPDATE", "2026-09-20 10:00:00", '{"a": 1}',
             '["deadline"]', "web:admin"),
            ("REQ-2026-0002", "DELETE", "2026-09-20 11:00:00", '{"b": 2}',
             "", "excel:watcher"),
        ]
        _make_src_db(self.src, rows)
        n = DatabaseManager.copy_ledger_history(self.src, self.dest_conn)
        self.assertEqual(n, 2)
        got = _read_dest_rows(self.dest_conn)
        self.assertEqual([tuple(r) for r in got], rows)

    def test_old_schema_falls_back_to_four_columns(self):
        _make_src_db(self.src, [("REQ-2026-0009", "INSERT", "2026-01-01 00:00:00", "{}")],
                     old_schema=True)
        n = DatabaseManager.copy_ledger_history(self.src, self.dest_conn)
        self.assertEqual(n, 1)
        got = _read_dest_rows(self.dest_conn)
        self.assertEqual(len(got), 1)
        self.assertEqual(got[0][:4], ("REQ-2026-0009", "INSERT", "2026-01-01 00:00:00", "{}"))
        self.assertIsNone(got[0][4])
        self.assertIsNone(got[0][5])

    def test_missing_source_returns_zero(self):
        self.assertEqual(
            DatabaseManager.copy_ledger_history(self.tmp / "nope.db", self.dest_conn), 0)

    def test_corrupt_source_returns_zero_without_raising(self):
        self.src.write_bytes(b"not a sqlite database at all" * 100)
        self.assertEqual(
            DatabaseManager.copy_ledger_history(self.src, self.dest_conn), 0)
        self.assertEqual(_read_dest_rows(self.dest_conn), [])


class TestDerivedWorkbookAtomicSave(unittest.TestCase):
    """ISSUE-002: 파생본 저장 실패 시 원본 보존 + tmp 잔재 없음."""

    def setUp(self):
        import generate_excel_db
        self.mod = generate_excel_db
        self.tmp = Path(tempfile.mkdtemp(prefix="audit_phase14_"))

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _workbook(self):
        wb = openpyxl.Workbook()
        cast(Any, wb.active)["A1"] = "값"
        return wb

    def test_success_saves_and_leaves_no_tmp(self):
        target = self.tmp / "통합.xlsx"
        self.assertTrue(self.mod.save_workbook_atomic(self._workbook(), target))
        self.assertTrue(target.exists())
        self.assertFalse((self.tmp / "통합.tmp.xlsx").exists())
        probe = openpyxl.load_workbook(target, read_only=True)
        try:
            self.assertIn("값", [c.value for row in cast(Any, probe.active).iter_rows() for c in row])
        finally:
            probe.close()

    def test_save_failure_keeps_target_and_cleans_tmp(self):
        target = self.tmp / "통합.xlsx"
        wb = self._workbook()
        wb.save(target)
        before = target.read_bytes()
        broken = self._workbook()
        with mock.patch.object(broken, "save", side_effect=RuntimeError("디스크 부족")):
            self.assertFalse(self.mod.save_workbook_atomic(broken, target))
        self.assertEqual(target.read_bytes(), before)
        self.assertFalse((self.tmp / "통합.tmp.xlsx").exists())

    def test_probe_failure_keeps_target(self):
        target = self.tmp / "통합.xlsx"
        self._workbook().save(target)
        before = target.read_bytes()
        with mock.patch("openpyxl.load_workbook", side_effect=ValueError("깨진 저장")):
            self.assertFalse(self.mod.save_workbook_atomic(self._workbook(), target))
        self.assertEqual(target.read_bytes(), before)
        self.assertFalse((self.tmp / "통합.tmp.xlsx").exists())


class TestCopyPathWarnAndSkip(unittest.TestCase):
    """ISSUE-003: copy 실패는 경고+건너뛰기, 배치는 계속된다."""

    def setUp(self):
        import add_documents_smart
        self.mod = add_documents_smart
        self.tmp = Path(tempfile.mkdtemp(prefix="audit_phase14_"))
        self.workspace = self.tmp / "ws"
        (self.workspace / "2026").mkdir(parents=True)
        self.inbox = self.tmp / "inbox"
        self.inbox.mkdir()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _run(self, files, **kwargs):
        with mock.patch.object(self.mod, "ROOT_DIR", self.workspace):
            return self.mod.copy_or_move_files(files, **kwargs)

    def test_copy_failure_is_skipped_not_crashing(self):
        busy = self.inbox / "260920 김의원 요구자료(1).hwp"
        free = self.inbox / "260920 이의원 요구자료(2).hwp"
        busy.write_bytes(b"busy")
        free.write_bytes(b"free")
        real_copy2 = shutil.copy2
        skipped = []

        def copy2_or_fail(s, d):
            if Path(s).name == busy.name:
                raise PermissionError("대상 잠김")
            return real_copy2(s, d)

        with mock.patch.object(self.mod.shutil, "copy2", side_effect=copy2_or_fail):
            moved = self._run([str(busy), str(free)], move=False, skipped=skipped)
        self.assertEqual([p.name for p in moved], [free.name])
        self.assertEqual([name for name, _ in skipped], [busy.name])
        self.assertTrue((self.workspace / "2026" / free.name).exists())
        self.assertFalse((self.workspace / "2026" / busy.name).exists())

    def test_copy_failure_without_skipped_list_still_skips(self):
        busy = self.inbox / "260920 박의원 요구자료(3).hwp"
        busy.write_bytes(b"busy")
        with mock.patch.object(self.mod.shutil, "copy2",
                               side_effect=PermissionError("대상 잠김")):
            moved = self._run([str(busy)], move=False)
        self.assertEqual(moved, [])
        self.assertFalse((self.workspace / "2026" / busy.name).exists())


if __name__ == "__main__":
    unittest.main()
