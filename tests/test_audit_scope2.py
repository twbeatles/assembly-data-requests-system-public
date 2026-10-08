# -*- coding: utf-8 -*-
"""SCOPE2 감사 개선 회귀 테스트."""

from typing import Any
import json
import sqlite3
import subprocess
import sys
import tempfile
import time
import unittest
import shutil
from pathlib import Path
from unittest.mock import patch

SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "scripts"
SYSTEM_DIR = SCRIPTS_DIR.parent
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from db.database_manager import DatabaseManager
from services.ledger_service import LedgerService, auto_link_ledger_to_doc
from services.sync_service import SyncService
from services.excel_sync_service import ExcelSyncService
from extractors.ledger_parser import LedgerParser
from template_renderer import DashboardRenderer
import extract_and_build_db
import parse_all
import web_server
import system_config
import create_ledger_template


class TestAuditScope2(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.base = Path(self.temp_dir)
        self.db_path = self.base / "data_requests.db"
        self.json_path = self.base / "data_requests.json"
        conn = sqlite3.connect(str(self.db_path))
        try:
            DatabaseManager.init_schema(conn)
        finally:
            conn.close()
        self.service = LedgerService(db_path=self.db_path, json_path=self.json_path, base_dir=self.base)

    def tearDown(self):
        LedgerService.wait_last_sync(timeout=2.0)
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_copy_ledger_history_between_databases(self):
        """파이프라인 재적재 시 ledger_history가 이관되어야 한다."""
        src = self.base / "src.db"
        dst = self.base / "dst.db"
        src_conn = sqlite3.connect(str(src))
        try:
            DatabaseManager.init_schema(src_conn)
            src_conn.execute(
                "INSERT INTO ledger_history (ledger_id, action, changed_at, snapshot) VALUES (?,?,?,?)",
                ("REQ-2026-001", "INSERT", "2026-09-10 10:00:00", "{}")
            )
            src_conn.commit()
        finally:
            src_conn.close()

        dst_conn = sqlite3.connect(str(dst))
        try:
            DatabaseManager.init_schema(dst_conn)
            copied = DatabaseManager.copy_ledger_history(src, dst_conn)
            dst_conn.commit()
            count = dst_conn.execute("SELECT COUNT(*) FROM ledger_history").fetchone()[0]
        finally:
            dst_conn.close()
        self.assertEqual(copied, 1)
        self.assertEqual(count, 1)

    def test_extract_pipeline_copies_history_then_writes_json_atomically(self):
        """extract_and_build_db는 스왑 이후 JSON을 원자적으로 쓰고 history를 복사해야 한다."""
        code = (SCRIPTS_DIR / "extract_and_build_db.py").read_text(encoding="utf-8")
        self.assertIn("copy_ledger_history", code)
        self.assertIn("write_json_atomic", code)
        self.assertLess(code.index("atomic_swap"), code.index("write_json_atomic(JSON_PATH"))

    def test_auto_link_rejects_partial_seq_match(self):
        """연번 1이 문서번호 10/11/100에 붙으면 안 된다."""
        conn = sqlite3.connect(str(self.db_path))
        conn.row_factory = sqlite3.Row
        try:
            DatabaseManager.ensure_schema(conn)
            conn.execute(
                """INSERT INTO documents (
                    doc_id, year, request_date, institution, requester, doc_number, version, title,
                    department, contact_person, question_list, answer_summary, answer_full_text,
                    has_tables, table_count, table_summary, topic_tags, parsed_md_path, original_path,
                    full_markdown, linked_ledger_id
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                ("DOC-2026-010", "2026", "2026-09-01", "국회", "김의원", "10", "v1", "열번째 자료",
                 "팀", "", "", "", "", "N", 0, "", "", "a.md", "2026/요구자료(10).hwp", "", "")
            )
            conn.execute(
                """INSERT INTO documents (
                    doc_id, year, request_date, institution, requester, doc_number, version, title,
                    department, contact_person, question_list, answer_summary, answer_full_text,
                    has_tables, table_count, table_summary, topic_tags, parsed_md_path, original_path,
                    full_markdown, linked_ledger_id
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                ("DOC-2026-001", "2026", "2026-09-01", "국회", "김의원", "1", "v1", "첫번째 자료",
                 "팀", "", "", "", "", "N", 0, "", "", "b.md", "2026/요구자료(1).hwp", "", "")
            )
            conn.commit()
            item = {"year": "2026", "seq_no": "1", "ledger_id": "REQ-2026-099", "requester": "김의원", "request_date": ""}
            linked = auto_link_ledger_to_doc(conn, item)
            self.assertEqual(linked, "DOC-2026-001")
            item2 = {"year": "2026", "seq_no": "1", "ledger_id": "REQ-2026-100", "requester": "김의원", "request_date": ""}
            # 연번 10 문서는 연번 1과 연결되면 안 됨
            conn.execute("DELETE FROM documents WHERE doc_id = 'DOC-2026-001'")
            conn.commit()
            linked2 = auto_link_ledger_to_doc(conn, dict(item2))
            self.assertEqual(linked2, "")
        finally:
            conn.close()

    def test_auto_link_skips_when_multiple_token_matches(self):
        conn = sqlite3.connect(str(self.db_path))
        conn.row_factory = sqlite3.Row
        try:
            DatabaseManager.ensure_schema(conn)
            for i, did in enumerate(("DOC-A", "DOC-B"), start=1):
                conn.execute(
                    """INSERT INTO documents (
                        doc_id, year, request_date, institution, requester, doc_number, version, title,
                        department, contact_person, question_list, answer_summary, answer_full_text,
                        has_tables, table_count, table_summary, topic_tags, parsed_md_path, original_path,
                        full_markdown, linked_ledger_id
                    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (did, "2026", "2026-09-01", "국회", "김의원", "15", "v1", f"자료{i}",
                     "팀", "", "", "", "", "N", 0, "", "", f"{i}.md", f"2026/요구자료(15)_{i}.hwp", "", "")
                )
            conn.commit()
            linked = auto_link_ledger_to_doc(conn, {
                "year": "2026", "seq_no": "15", "ledger_id": "REQ-X", "requester": "김의원", "request_date": ""
            })
            self.assertEqual(linked, "")
        finally:
            conn.close()

    def test_api_file_rejects_db_and_scripts(self):
        handler: Any = web_server.RequestLedgerHandler.__new__(web_server.RequestLedgerHandler)
        ok, _, err = handler.resolve_download_target("data_requests.db")
        self.assertFalse(ok)
        self.assertIn("거부", err)
        ok2, _, err2 = handler.resolve_download_target("scripts/web_server.py")
        self.assertFalse(ok2)

    def test_static_handler_rejects_sqlite(self):
        handler: Any = web_server.RequestLedgerHandler.__new__(web_server.RequestLedgerHandler)
        self.assertFalse(handler.is_public_static("/data_requests.db"))
        self.assertFalse(handler.is_public_static("/config.json"))
        self.assertFalse(handler.is_public_static("/scripts/web_server.py"))
        self.assertTrue(handler.is_public_static("/자료요구_통합검색_대시보드.html"))

    def test_insert_validation_uses_http_400(self):
        handler: Any = web_server.RequestLedgerHandler.__new__(web_server.RequestLedgerHandler)
        captured = []

        def mock_send(data, status_code=200):
            captured.append((status_code, data))

        handler.send_json_response = mock_send
        handler.headers = {"Content-Length": "2"}
        handler.rfile = type("R", (), {"read": lambda self, n: b"{}"})()
        handler.path = "/api/ledger"
        handler.read_json_body = lambda: {"year": "2026", "requester": "의원"}
        old_db, old_json, old_base = web_server.DB_PATH, web_server.JSON_PATH, web_server.BASE_DIR
        try:
            web_server.DB_PATH = self.db_path
            web_server.JSON_PATH = self.json_path
            web_server.BASE_DIR = self.base
            handler.do_POST()
        finally:
            web_server.DB_PATH = old_db
            web_server.JSON_PATH = old_json
            web_server.BASE_DIR = old_base
        self.assertTrue(captured)
        self.assertEqual(captured[0][0], 400)
        self.assertFalse(captured[0][1].get("success"))

    def test_insert_rejects_invalid_dates(self):
        """웹 등록의 날짜 3종은 실재하는 날짜(또는 빈 값)만 받는다. (직전 회차 ISSUE-004)

        엑셀 파서가 읽는 표기(`2026. 3. 5`, `2026/03/05`)는 `YYYY-MM-DD`로 정규화해 받는다.
        웹만 더 좁게 받으면 엑셀에서 들어온 값을 화면에서 다시 저장할 수 없다. (7회차 ISSUE-003)
        """
        base = {"title": "기한 시험", "requester": "김의원", "year": "2026"}
        for bad in ("2026-13-40", "내일까지", "2026-02-30"):
            payload = dict(base, deadline=bad)
            res = self.service.insert_ledger_item(payload)
            self.assertFalse(res.get("success"), bad)
        for i, text in enumerate(("2026. 3. 5", "2026/03/05")):
            res = self.service.insert_ledger_item(dict(base, title=f"정규화 {i}", deadline=text))
            self.assertTrue(res.get("success"), res)
            self.assertEqual(res["item"]["deadline"], "2026-03-05")
        ok = self.service.insert_ledger_item(dict(base, deadline="2026-09-30"))
        self.assertTrue(ok.get("success"), ok)
        empty = self.service.insert_ledger_item(dict(base, title="기한 없음", deadline=""))
        self.assertTrue(empty.get("success"), empty)

    def test_update_rejects_invalid_dates(self):
        """웹 수정의 날짜 3종도 같은 규칙을 따른다. (ISSUE-004)"""
        lid = self.service.insert_ledger_item(
            {"title": "수정 시험", "requester": "김의원", "year": "2026"})["item"]["ledger_id"]
        bad = self.service.update_ledger_item(lid, {"deadline": "2026-02-30"})
        self.assertFalse(bad.get("success"))
        row = self.service.get_ledger_item(lid)
        assert row is not None
        self.assertEqual(row.get("deadline") or "", "")
        good = self.service.update_ledger_item(lid, {"deadline": "2026-09-30"})
        self.assertTrue(good.get("success"), good)

    def test_delete_missing_id_is_not_success(self):
        res = self.service.delete_ledger_item("REQ-2099-999")
        self.assertFalse(res.get("success"))

    def test_get_stats_failure_flag(self):
        bad_path = self.base / "not_a_db_dir"
        bad_path.mkdir()
        broken = LedgerService(db_path=bad_path, json_path=self.json_path, base_dir=self.base)
        stats = broken.get_stats()
        self.assertFalse(stats.get("success"))

    def test_config_post_rejects_unknown_keys(self):
        handler: Any = web_server.RequestLedgerHandler.__new__(web_server.RequestLedgerHandler)
        captured = []
        handler.send_json_response = lambda data, status_code=200: captured.append((status_code, data))
        handler.send_error_response = lambda code, msg: captured.append((code, {"success": False, "error": msg}))
        handler.read_json_body = lambda: {"ledger_patterns": ["*.xls*"], "web_port": 1}
        handler.path = "/api/config"
        handler.do_POST()
        self.assertEqual(captured[0][0], 400)

    def test_oversized_body_reports_413(self):
        """5MB 초과 본문은 413과 용량 안내를 반환해야 한다. (ISSUE-003)"""
        handler: Any = web_server.RequestLedgerHandler.__new__(web_server.RequestLedgerHandler)
        captured = []
        handler.send_json_response = lambda data, status_code=200: captured.append((status_code, data))
        handler.headers = {"Content-Length": str(5 * 1024 * 1024 + 1)}
        handler.rfile = type("R", (), {"read": lambda self, n: b"{}"})()
        handler.path = "/api/ledger"
        handler.do_POST()
        self.assertTrue(captured)
        self.assertEqual(captured[0][0], 413)
        self.assertIn("5MB", captured[0][1].get("error", ""))

    def test_malformed_json_body_reports_detail(self):
        """깨진 JSON 본문은 400과 JSON 형식 안내를 반환해야 한다. (ISSUE-003)"""
        handler: Any = web_server.RequestLedgerHandler.__new__(web_server.RequestLedgerHandler)
        captured = []
        handler.send_json_response = lambda data, status_code=200: captured.append((status_code, data))
        body = b"{bad json"
        handler.headers = {"Content-Length": str(len(body))}
        handler.rfile = type("R", (), {"read": lambda self, n: body})()
        handler.path = "/api/ledger"
        handler.do_POST()
        self.assertTrue(captured)
        self.assertEqual(captured[0][0], 400)
        self.assertIn("JSON", captured[0][1].get("error", ""))

    def test_download_streams_file_bytes(self):
        """원문 다운로드는 파일 바이트와 Content-Length를 그대로 전달한다. (Gap-3)"""
        ws_root = system_config.find_workspace_root(self.base)
        doc_dir = ws_root / "2026"
        doc_dir.mkdir(parents=True, exist_ok=True)
        payload = ("본문 내용\n" * 4096).encode("utf-8")
        (doc_dir / "테스트_질의서.md").write_bytes(payload)

        handler: Any = web_server.RequestLedgerHandler.__new__(web_server.RequestLedgerHandler)
        sent = {}
        chunks = []
        handler.send_response = lambda code: sent.setdefault("code", code)
        handler.send_header = lambda k, v: sent.setdefault(k, v)
        handler.end_headers = lambda: None
        handler.wfile = type("W", (), {"write": lambda self, b: chunks.append(bytes(b))})()
        handler.headers = {}
        from urllib.parse import quote
        handler.path = "/api/download?path=" + quote("2026/테스트_질의서.md")
        old_base = web_server.BASE_DIR
        try:
            web_server.BASE_DIR = self.base
            handler.do_GET()
        finally:
            web_server.BASE_DIR = old_base
        self.assertEqual(sent.get("code"), 200)
        self.assertEqual(sent.get("Content-Length"), str(len(payload)))
        self.assertEqual(b"".join(chunks), payload)

    def test_slim_payload_keeps_full_markdown_for_normal_docs(self):
        payload = DashboardRenderer.build_slim_payload({
            "documents": [{"doc_id": "D1", "title": "t", "full_markdown": "# 전문 본문입니다", "answer_full_text": "잘린본문"}],
            "qa_items": [{"qa_id": "Q1", "doc_id": "D1", "question_title": "질", "answer_markdown": "장문 답변 내용"}],
            "request_ledger": []
        })
        self.assertEqual(payload["documents"][0]["full_markdown"], "# 전문 본문입니다")
        self.assertFalse(payload["documents"][0].get("body_deferred"))
        self.assertEqual(payload["qa_items"][0]["answer_markdown"], "장문 답변 내용")
        self.assertEqual(payload["qa_items"][0]["question_title"], "질")

    def test_slim_payload_never_drops_full_markdown(self):
        """용량이 커도 대시보드 페이로드에서 문서·Q&A 전문을 빼지 않는다."""
        huge = "가" * 250_000
        qa_long = "답" * 80_000
        payload = DashboardRenderer.build_slim_payload({
            "documents": [{
                "doc_id": "D2", "title": "big", "full_markdown": huge,
                "parsed_md_path": "_parsed_markdown/x.md"
            }],
            "qa_items": [{
                "qa_id": "Q2", "doc_id": "D2", "question_title": "질",
                "answer_markdown": qa_long
            }],
            "request_ledger": []
        })
        self.assertEqual(payload["documents"][0]["full_markdown"], huge)
        self.assertFalse(payload["documents"][0].get("body_deferred"))
        self.assertEqual(payload["qa_items"][0]["answer_markdown"], qa_long)

    def test_modal_prefers_full_markdown_over_truncated_answer(self):
        tmpl = (SCRIPTS_DIR / "templates" / "dashboard_template.html").read_text(encoding="utf-8")
        self.assertIn("doc.full_markdown || doc.answer_full_text", tmpl)
        self.assertNotIn("innerText = doc.answer_full_text || doc.full_markdown", tmpl)
        self.assertIn("body_incomplete", tmpl)
        self.assertIn("/api/file", tmpl)

    def test_save_atomic_writes_via_tempfile(self):
        out = self.base / "dash.html"
        ok = DashboardRenderer.save_atomic(out, "<html>ok</html>")
        self.assertTrue(ok)
        self.assertTrue(out.exists())
        self.assertFalse(out.with_suffix(".tmp.html").exists())
        src = (SCRIPTS_DIR / "template_renderer.py").read_text(encoding="utf-8")
        block = src.split("def save_atomic")[1].split("return saved")[0]
        self.assertIn("tmp_out.write_text", block)
        self.assertNotIn("out_path.write_text(content", block)

    def test_xlsx_fallback_escapes_html(self):
        import openpyxl
        xlsx = self.base / "big.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["<svg/onload=alert(1)>", "정상"])
        wb.save(str(xlsx))
        wb.close()
        out = self.base / "out.md"
        ok, err = parse_all.handle_large_xlsx_fallback(xlsx, out)
        self.assertTrue(ok)
        text = out.read_text(encoding="utf-8")
        self.assertNotIn("<svg/onload=alert(1)>", text)
        self.assertIn("&lt;svg", text)

    def test_parse_all_uses_hash_sidecar(self):
        src = (SCRIPTS_DIR / "parse_all.py").read_text(encoding="utf-8")
        self.assertIn("sha256", src.lower())
        self.assertIn(".sha256", src)
        self.assertIn("--ocr", src)
        self.assertIn("--keep-empty-cols", src)
        self.assertIn("enrich_image_only_markdown", src)
        self.assertNotIn("xlsx' and file_path.stat().st_size > 10 * 1024 * 1024", src)

    def test_chart_and_href_use_escapeHtml(self):
        tmpl = (SCRIPTS_DIR / "templates" / "dashboard_template.html").read_text(encoding="utf-8")
        self.assertIn("escapeHtml(req)", tmpl)
        self.assertIn("escapeHtml(top)", tmpl)
        self.assertIn("escapeHtml(doc.parsed_md_path", tmpl)
        self.assertIn("svg|math", tmpl)
        self.assertIn("ensureDocBody", tmpl)

    def test_ledger_parser_class_default_dir(self):
        parser = LedgerParser()
        self.assertTrue(parser.root_dir.exists() or parser.root_dir.name)

    def test_ledger_template_has_id_and_dept_columns(self):
        out = self.base / "tpl.xlsx"
        create_ledger_template.create_template(out)
        import openpyxl
        wb = openpyxl.load_workbook(out)
        names = list(wb.sheetnames)
        ws = wb["2026"] if "2026" in names else wb[names[0]]
        headers = [c.value for c in ws[1]]
        wb.close()
        self.assertIn("대장ID", headers)
        self.assertIn("담당부서", headers)
        self.assertGreaterEqual(len(names), 2)

    def test_sync_service_timeout_clears_flag(self):
        svc = SyncService(scripts_dir=self.base, timeout_seconds=1)
        with patch("subprocess.run", side_effect=subprocess.TimeoutExpired(cmd="py", timeout=1)):
            started, _ = svc.trigger_async_sync()
            self.assertTrue(started)
            for _ in range(20):
                time.sleep(0.05)
                if not svc.get_status()["is_syncing"]:
                    break
            st = svc.get_status()
            self.assertFalse(st["is_syncing"])
            self.assertTrue(st["last_error"])

    def test_pending_queue_persisted_to_disk(self):
        svc = ExcelSyncService(base_dir=self.base, db_path=self.db_path, json_path=self.json_path)
        with svc._pending_lock:
            svc._pending_queue.append({"action": "insert", "item": {"ledger_id": "REQ-1"}, "timestamp": 1})
            svc._save_pending_queue()
        qfile = self.base / ".excel_pending_queue.json"
        self.assertTrue(qfile.exists())
        svc2 = ExcelSyncService(base_dir=self.base, db_path=self.db_path, json_path=self.json_path)
        self.assertEqual(len(svc2._pending_queue), 1)

    def test_default_ledger_patterns_are_specific(self):
        pats = system_config.DEFAULT_CONFIG["ledger_patterns"]
        self.assertNotIn("*요구자료*.xls*", pats)
        self.assertNotIn("*자료요구*.xls*", pats)

    def test_universal_config_patterns_are_specific(self):
        univ = SYSTEM_DIR.parent / "국회자료요구_스마트시스템_범용배포용" / "config.json"
        if univ.exists():
            cfg = json.loads(univ.read_text(encoding="utf-8"))
            self.assertNotIn("*요구자료*.xls*", cfg.get("ledger_patterns", []))

    def test_write_json_atomic_helper(self):
        target = self.base / "data_requests.json"
        extract_and_build_db.write_json_atomic(target, {"documents": []})
        self.assertTrue(target.exists())
        self.assertFalse(target.with_name("data_requests.json.tmp").exists())
        data = json.loads(target.read_text(encoding="utf-8"))
        self.assertEqual(data["documents"], [])


if __name__ == "__main__":
    unittest.main()
