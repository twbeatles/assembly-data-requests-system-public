# -*- coding: utf-8 -*-
"""
PROJECT_AUDIT.md Phase 1~2 개선에 대한 회귀 테스트.

감사에서 재현된 경로를 그대로 고정한다.
- ISSUE-001: 재시작 시 보류 큐를 Excel→DB 동기화보다 먼저 반영
- ISSUE-002: 웹 insert/update가 다른 ID의 Excel 행을 덮어쓰지 않음
- ISSUE-003: `대장ID` 고정(스탬핑) 후 행 삭제·정렬에도 항목이 소실되지 않음
- ISSUE-004: Excel 열림 중 수기 저장이 flush에 가려지지 않음
- ISSUE-005: 연도 시트가 없으면 새로 만들어 연도가 바뀌지 않음
- ISSUE-006: 대시보드 저장 실패를 성공으로 보고하지 않음
- ISSUE-007: 파이프라인 실행 중에는 Excel 동기화를 건너뜀
- GAP-003/004: 본문 타입·연도 검증, Host/Origin 검증
"""

from typing import Any
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import openpyxl

import system_config  # noqa: F401  (config 기본값 경로 확인용)
from db.database_manager import DatabaseManager
from extractors.ledger_parser import load_request_ledger, stamp_ledger_ids
from services.excel_sync_service import ExcelSyncService
from services.ledger_service import LedgerService

HEADERS_WITH_ID = ["대장ID", "연번", "의원", "요구자료", "제출", "요청형태", "비고"]
HEADERS_NO_ID = ["연번", "의원", "요구자료", "제출", "요청형태", "비고"]


class LedgerSyncFixtureMixin:
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="audit_phase8_"))
        self.db_path = self.tmp / "data_requests.db"
        self.json_path = self.tmp / "data_requests.json"
        conn = sqlite3.connect(str(self.db_path))
        DatabaseManager.init_schema(conn)
        conn.close()
        ExcelSyncService._instance = None

    def tearDown(self):
        svc = ExcelSyncService._instance
        if svc is not None:
            svc.stop_file_watcher()
        ExcelSyncService._instance = None
        shutil.rmtree(self.tmp, ignore_errors=True)

    def make_excel(self, rows, with_id=True, sheets=("2026",)):
        path = self.tmp / "260904 국회 요구자료 목록(테스트).xlsx"
        header = HEADERS_WITH_ID if with_id else HEADERS_NO_ID
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = sheets[0]
        ws.append(header)
        for row in rows:
            ws.append(row)
        for extra in sheets[1:]:
            wb.create_sheet(extra).append(header)
        wb.save(path)
        wb.close()
        return path

    def sync_service(self):
        return ExcelSyncService.get_instance(self.tmp)

    def ledger_service(self):
        return LedgerService(self.db_path, self.json_path, self.tmp)

    def db_rows(self):
        conn = sqlite3.connect(str(self.db_path))
        conn.row_factory = sqlite3.Row
        try:
            return {r["ledger_id"]: dict(r) for r in conn.execute("SELECT * FROM request_ledger")}
        finally:
            conn.close()

    def excel_rows(self, path, sheet="2026"):
        wb = openpyxl.load_workbook(path, data_only=True)
        try:
            ws = wb[sheet]
            header = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
            out = []
            for r in range(2, ws.max_row + 1):
                out.append({
                    header[c - 1]: ws.cell(r, c).value
                    for c in range(1, ws.max_column + 1)
                })
            return out
        finally:
            wb.close()

    def touch_future(self, path, seconds=5):
        future = time.time() + seconds
        os.utime(path, (future, future))


class TestStartupAndWatcherOrdering(LedgerSyncFixtureMixin, unittest.TestCase):
    """ISSUE-001 / ISSUE-004: 보류 큐와 Excel→DB 동기화의 순서."""

    def test_restart_keeps_web_insert_and_update_and_absorbs_manual_edit(self):
        excel = self.make_excel([
            ["REQ-2026-001", "1", "김의원", "자료A", "작성중", "일반", ""],
            ["REQ-2026-002", "2", "이의원", "자료B", "작성중", "일반", ""],
            ["REQ-2026-003", "3", "박의원", "자료C", "작성중", "일반", ""],
        ])
        self.sync_service().sync_excel_to_db()
        ledger = self.ledger_service()

        # 엑셀이 열려 있는 동안의 웹 등록/수정은 보류 큐로 간다.
        with mock.patch.object(ExcelSyncService, "is_file_locked", return_value=True):
            inserted = ledger.insert_ledger_item({
                "year": "2026", "requester": "최의원",
                "title": "자료D(웹등록)", "request_type": "메일",
            })
            updated = ledger.update_ledger_item("REQ-2026-001", {"status": "제출"})
        LedgerService.wait_last_sync(3)
        self.assertTrue(inserted["success"])
        self.assertTrue(inserted["excel_pending"])
        self.assertTrue(updated["excel_pending"])
        new_id = inserted["item"]["ledger_id"]

        # 서버가 꺼진 뒤 사용자가 엑셀의 다른 행을 수기 수정하고 저장한다.
        wb = openpyxl.load_workbook(excel)
        wb["2026"].cell(4, 7, "수기 메모")
        wb.save(excel)
        wb.close()
        self.touch_future(excel)

        # 재시작: init_db() → check_and_sync_startup()
        ExcelSyncService._instance = None
        restarted = ExcelSyncService.get_instance(self.tmp)
        self.assertEqual(restarted.pending_count(), 2)
        restarted.check_and_sync_startup()

        rows = self.db_rows()
        self.assertIn(new_id, rows)
        self.assertNotEqual(rows[new_id]["status"], "[삭제]", "웹 등록 항목이 삭제 처리되면 안 된다")
        self.assertEqual(rows["REQ-2026-001"]["status"], "제출", "웹 수정이 엑셀 구값으로 되돌아가면 안 된다")
        self.assertEqual(rows["REQ-2026-003"]["note"], "수기 메모", "엑셀 수기 수정은 DB에 반영돼야 한다")

        excel_by_id = {str(r.get("대장ID") or ""): r for r in self.excel_rows(excel)}
        self.assertIn(new_id, excel_by_id, "보류 큐의 등록이 엑셀에도 기록돼야 한다")
        self.assertEqual(excel_by_id["REQ-2026-001"]["제출"], "제출")
        self.assertEqual(restarted.pending_count(), 0)

    def test_watcher_syncs_manual_edit_before_flushing_queue(self):
        excel = self.make_excel([
            ["REQ-2026-001", "1", "김의원", "자료A", "작성중", "일반", ""],
            ["REQ-2026-002", "2", "이의원", "자료B", "작성중", "일반", ""],
        ])
        svc = self.sync_service()
        svc.sync_excel_to_db()
        ledger = self.ledger_service()

        with mock.patch.object(ExcelSyncService, "is_file_locked", return_value=True):
            ledger.update_ledger_item("REQ-2026-001", {"status": "제출"})
        LedgerService.wait_last_sync(3)

        # 엑셀을 열어 둔 채 다른 행을 수정하고 저장한 뒤 엑셀을 닫은 상황.
        wb = openpyxl.load_workbook(excel)
        wb["2026"].cell(3, 5, "검토중")
        wb.save(excel)
        wb.close()
        self.touch_future(excel, seconds=2)

        svc.start_file_watcher(poll_interval=0.3)
        deadline = time.time() + 8
        while time.time() < deadline:
            if svc.pending_count() == 0 and self.db_rows()["REQ-2026-002"]["status"] == "검토중":
                break
            time.sleep(0.3)
        svc.stop_file_watcher()

        rows = self.db_rows()
        self.assertEqual(rows["REQ-2026-002"]["status"], "검토중", "수기 수정이 flush에 가려지면 안 된다")
        self.assertEqual(rows["REQ-2026-001"]["status"], "제출")
        excel_by_id = {str(r.get("대장ID") or ""): r for r in self.excel_rows(excel)}
        self.assertEqual(excel_by_id["REQ-2026-001"]["제출"], "제출")

    def test_pipeline_lock_defers_excel_sync(self):
        excel = self.make_excel([["REQ-2026-001", "1", "김의원", "자료A", "작성중", "일반", ""]])
        svc = self.sync_service()
        svc.sync_excel_to_db()

        wb = openpyxl.load_workbook(excel)
        wb["2026"].cell(2, 5, "검토중")
        wb.save(excel)
        wb.close()
        self.touch_future(excel)

        # 다른 살아 있는 프로세스가 파이프라인 락을 보유한 상황
        holder = subprocess.Popen(
            [sys.executable, "-c", "import time; time.sleep(60)"],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        try:
            (self.tmp / ".pipeline.lock").write_text(str(holder.pid), encoding="utf-8")
            before = svc._last_known_excel_mtime
            res = svc.sync_excel_to_db()
            self.assertFalse(res.get("success"))
            self.assertTrue(res.get("skipped"))
            self.assertEqual(svc._last_known_excel_mtime, before, "건너뛴 경우 mtime을 갱신하면 안 된다")
            self.assertEqual(self.db_rows()["REQ-2026-001"]["status"], "작성중")
        finally:
            holder.kill()
            holder.wait()
            (self.tmp / ".pipeline.lock").unlink(missing_ok=True)

        # 락이 풀리면 정상 반영
        self.assertTrue(svc.sync_excel_to_db().get("success"))
        self.assertEqual(self.db_rows()["REQ-2026-001"]["status"], "검토중")


class TestExcelRowMatching(LedgerSyncFixtureMixin, unittest.TestCase):
    """ISSUE-002: 행 매칭은 ledger_id를 우선한다."""

    def test_insert_with_same_requester_and_title_appends_new_row(self):
        excel = self.make_excel([
            ["REQ-2026-001", "1", "김의원", "딥페이크 시정요구 현황", "제출", "일반", "1분기"],
            ["REQ-2026-002", "2", "이의원", "자료B", "작성중", "일반", ""],
        ])
        self.sync_service().sync_excel_to_db()
        ledger = self.ledger_service()

        res = ledger.insert_ledger_item({
            "year": "2026", "requester": "김의원",
            "title": "딥페이크 시정요구 현황", "note": "2분기", "request_type": "메일",
        })
        LedgerService.wait_last_sync(3)
        self.assertTrue(res["success"])
        new_id = res["item"]["ledger_id"]

        rows = self.excel_rows(excel)
        self.assertEqual(len(rows), 3, "기존 행을 덮어쓰지 않고 새 행을 추가해야 한다")
        by_id = {str(r.get("대장ID") or ""): r for r in rows}
        self.assertEqual(by_id["REQ-2026-001"]["비고"], "1분기", "기존 행 내용이 유지돼야 한다")
        self.assertEqual(by_id[new_id]["비고"], "2분기")

        self.sync_service().sync_excel_to_db()
        db = self.db_rows()
        self.assertNotEqual(db["REQ-2026-001"]["status"], "[삭제]")
        self.assertNotEqual(db[new_id]["status"], "[삭제]")

    def test_update_matches_id_even_when_another_row_shares_seq_no(self):
        excel = self.make_excel([
            ["REQ-2026-001", "5", "김의원", "자료A", "작성중", "일반", "건드리지 말 것"],
            ["REQ-2026-002", "2", "이의원", "자료B", "작성중", "일반", ""],
        ])
        self.sync_service().sync_excel_to_db()
        ledger = self.ledger_service()

        res = ledger.update_ledger_item("REQ-2026-002", {"seq_no": "5", "status": "제출"})
        LedgerService.wait_last_sync(3)
        self.assertTrue(res["success"])

        by_id = {str(r.get("대장ID") or ""): r for r in self.excel_rows(excel)}
        self.assertEqual(by_id["REQ-2026-001"]["비고"], "건드리지 말 것")
        self.assertEqual(by_id["REQ-2026-001"]["제출"], "작성중")
        self.assertEqual(by_id["REQ-2026-002"]["제출"], "제출")

    def test_update_only_writes_changed_fields(self):
        excel = self.make_excel([
            ["REQ-2026-001", "1", "김의원", "자료A", "작성중", "일반", "원래 비고"],
        ])
        self.sync_service().sync_excel_to_db()
        ledger = self.ledger_service()

        with mock.patch.object(ExcelSyncService, "is_file_locked", return_value=True):
            ledger.update_ledger_item("REQ-2026-001", {"status": "제출"})
        LedgerService.wait_last_sync(3)

        # 보류 중 사용자가 같은 행의 다른 칸을 수기 수정한다.
        wb = openpyxl.load_workbook(excel)
        wb["2026"].cell(2, 7, "사용자 수기 비고")
        wb.save(excel)
        wb.close()

        self.sync_service().flush_pending_queue()
        row = self.excel_rows(excel)[0]
        self.assertEqual(row["제출"], "제출")
        self.assertEqual(row["비고"], "사용자 수기 비고", "바뀌지 않은 필드를 옛 스냅숏으로 덮어써선 안 된다")


class TestLedgerIdStamping(LedgerSyncFixtureMixin, unittest.TestCase):
    """ISSUE-003: 위치 기반 ID를 고정해 행 이동에 안전하게 만든다."""

    def test_row_delete_without_id_column_keeps_active_items(self):
        excel = self.make_excel([
            ["1", "김의원", "자료A", "작성중", "일반", ""],
            ["2", "이의원", "자료B", "작성중", "일반", ""],
            ["3", "박의원", "자료C", "작성중", "일반", ""],
            ["4", "최의원", "자료D", "작성중", "일반", ""],
        ], with_id=False)
        svc = self.sync_service()
        svc.sync_excel_to_db()

        seeded = {lid: r["title"] for lid, r in self.db_rows().items()}
        self.assertEqual(len(seeded), 4)
        header = [str(c.value or "") for c in openpyxl.load_workbook(excel)["2026"][1]]
        self.assertIn("대장ID", header, "대장ID 열이 자동 추가돼야 한다")

        ledger = self.ledger_service()
        ledger.delete_ledger_item("REQ-2026-002")  # 자료B 웹 삭제
        LedgerService.wait_last_sync(3)

        wb = openpyxl.load_workbook(excel)
        wb["2026"].delete_rows(2)  # 사용자가 자료A 행 삭제
        wb.save(excel)
        wb.close()
        self.touch_future(excel)

        svc.sync_excel_to_db()
        active = {r["title"] for r in self.db_rows().values() if r["status"] != "[삭제]"}
        self.assertEqual(active, {"자료C", "자료D"}, "행 삭제로 다른 항목이 사라지면 안 된다")

    def test_stamping_reassigns_duplicated_ids(self):
        excel = self.make_excel([
            ["REQ-2026-001", "1", "김의원", "자료A", "작성중", "일반", ""],
            ["REQ-2026-001", "2", "이의원", "복사된 행", "작성중", "일반", ""],
        ])
        written = stamp_ledger_ids(excel, [{"ledger_id": "REQ-2026-001", "year": "2026",
                                            "title": "자료A", "requester": "김의원"}])
        self.assertEqual(written, 1)
        ids = [str(r.get("대장ID") or "") for r in self.excel_rows(excel)]
        self.assertEqual(len(set(ids)), 2, "중복 ID는 새 ID로 분리돼야 한다")
        self.assertIn("REQ-2026-001", ids)

    def test_parser_skips_blank_id_row_in_stamped_sheet(self):
        excel = self.make_excel([
            ["REQ-2026-005", "1", "김의원", "자료A", "작성중", "일반", ""],
            ["", "2", "이의원", "ID 없는 신규 행", "작성중", "일반", ""],
        ])
        items = load_request_ledger(self.tmp, [], merge_db_records=False)
        ids = {item["ledger_id"] for item in items}
        self.assertEqual(ids, {"REQ-2026-005"},
                         "스탬핑 전 빈 ID 행에 위치 기반 ID를 매겨 기존 ID와 충돌시키면 안 된다")

        # 스탬핑하면 새 ID를 받아 정상 반영된다.
        stamp_ledger_ids(excel, [{"ledger_id": "REQ-2026-005", "year": "2026",
                                  "title": "자료A", "requester": "김의원"}])
        items2 = load_request_ledger(self.tmp, [], merge_db_records=False)
        self.assertEqual(len(items2), 2)
        self.assertEqual(len({i["ledger_id"] for i in items2}), 2)


class TestYearSheetCreation(LedgerSyncFixtureMixin, unittest.TestCase):
    """ISSUE-005: 연도 시트가 없으면 새로 만든다."""

    def test_insert_for_missing_year_creates_sheet_and_keeps_year(self):
        excel = self.make_excel(
            [["REQ-2025-001", "1", "김의원", "자료A", "제출", "일반", ""]],
            sheets=("2025", "2026"),
        )
        self.sync_service().sync_excel_to_db()
        ledger = self.ledger_service()

        res = ledger.insert_ledger_item({
            "year": "2027", "requester": "이의원", "title": "2027 신규", "request_type": "메일",
        })
        LedgerService.wait_last_sync(3)
        self.assertTrue(res["success"])
        new_id = res["item"]["ledger_id"]

        wb = openpyxl.load_workbook(excel)
        sheetnames = wb.sheetnames
        wb.close()
        self.assertIn("2027", sheetnames)

        rows_2025 = [str(r.get("대장ID") or "") for r in self.excel_rows(excel, "2025")]
        self.assertNotIn(new_id, rows_2025, "다른 연도 시트에 섞여 들어가면 안 된다")
        rows_2027 = [str(r.get("대장ID") or "") for r in self.excel_rows(excel, "2027")]
        self.assertIn(new_id, rows_2027)

        years = {i["ledger_id"]: i["year"] for i in load_request_ledger(self.tmp, [], merge_db_records=False)}
        self.assertEqual(years[new_id], "2027", "파이프라인 재빌드에서도 연도가 유지돼야 한다")


class TestValidationAndGuards(LedgerSyncFixtureMixin, unittest.TestCase):
    """GAP-003 / GAP-004 및 ISSUE-006 보고 경로."""

    def test_non_dict_body_and_bad_year_are_rejected(self):
        ledger = self.ledger_service()
        self.assertFalse(ledger.insert_ledger_item([1, 2])["success"])
        self.assertFalse(ledger.update_ledger_item("REQ-2026-001", "문자열")["success"])
        bad_year = ledger.insert_ledger_item({"year": "26년", "requester": "김의원", "title": "연도 오류"})
        self.assertFalse(bad_year["success"])
        self.assertIn("연도", bad_year["error"])

    def test_handler_rejects_foreign_host_and_non_dict_json(self):
        import web_server

        handler: Any = web_server.RequestLedgerHandler.__new__(web_server.RequestLedgerHandler)
        sent = []
        handler.send_json_response = lambda data, status_code=200: sent.append((status_code, data))

        handler.headers = {"Host": "attacker.example.com"}
        handler.path = "/api/sync"
        handler.do_GET()
        self.assertEqual(sent[-1][0], 403)

        handler.headers = {"Host": "127.0.0.1:8080"}
        handler.path = "/api/ledger"
        handler.read_json_body = lambda: [1, 2, 3]
        handler.do_POST()
        self.assertEqual(sent[-1][0], 400)

    def test_trusted_write_origin_requires_same_port(self):
        import web_server

        handler: Any = web_server.RequestLedgerHandler.__new__(web_server.RequestLedgerHandler)
        handler.server = mock.Mock()
        handler.server.server_address = ("127.0.0.1", 8080)

        handler.headers = {"Origin": "http://127.0.0.1:8080"}
        self.assertTrue(handler.has_trusted_write_origin())
        handler.headers = {"Origin": "http://localhost:3000"}
        self.assertFalse(handler.has_trusted_write_origin())
        handler.headers = {"Origin": "null"}
        self.assertFalse(handler.has_trusted_write_origin())

    def test_dashboard_generator_reports_save_failure(self):
        import generate_web_dashboard as gwd

        payload = {"documents": [], "qa_items": [], "request_ledger": []}
        self.json_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        with mock.patch.object(gwd, "JSON_PATH", self.json_path), \
             mock.patch.object(gwd._renderer, "save_atomic", return_value=False):
            result = gwd.main()
        self.assertIsInstance(result, dict)
        self.assertFalse(result["success"], "HTML 저장 실패를 성공으로 보고하면 안 된다")

    def test_gui_reports_pipeline_failure(self):
        code = (SCRIPTS_DIR / "launcher_gui.py").read_text(encoding="utf-8")
        self.assertIn('result_data.get("success", True)', code,
                      "GUI가 파이프라인 success 값을 확인해야 한다")
        # 문자열만 확인하면 그 코드가 실행 불가능해도 통과한다. 실제로 ISSUE-005
        # (nonlocal 누락 → UnboundLocalError)가 이 테스트 아래에서 오래 가려져 있었다.
        # 실행 기반 검증은 tests/test_audit_phase9_fixes.py
        # TestLauncherGuiCompletionHandler 에 있다.
        self.assertIn("nonlocal err_msg", code,
                      "err_msg에 대입하려면 nonlocal 선언이 있어야 UnboundLocalError를 피한다")

    def test_reconcile_reports_db_excel_divergence(self):
        import reconcile_storage

        excel = self.make_excel([["REQ-2026-001", "1", "김의원", "자료A", "작성중", "일반", ""]])
        self.assertTrue(excel.exists())
        with mock.patch.object(reconcile_storage, "SYSTEM_DIR", self.tmp):
            self.sync_service().sync_excel_to_db()
            ledger = self.ledger_service()
            with mock.patch.object(ExcelSyncService, "find_master_excel", return_value=None):
                # 엑셀에 기록되지 않은 웹 등록 1건을 만든다.
                res = ledger.insert_ledger_item({
                    "year": "2026", "requester": "이의원", "title": "엑셀 미반영", "request_type": "메일",
                })
            LedgerService.wait_last_sync(3)
            report = reconcile_storage.excel_report(self.db_path)

        self.assertTrue(report["available"])
        self.assertIn(res["item"]["ledger_id"], report["only_in_db"])
        self.assertEqual(report["only_in_excel"], [])


if __name__ == "__main__":
    unittest.main()
