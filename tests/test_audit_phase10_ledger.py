# -*- coding: utf-8 -*-
"""감사 3회차(PROJECT_AUDIT.md R3-xx) 관리대장 동기화 회귀 테스트.

모든 테스트는 임시 디렉터리만 쓴다. 운영 폴더 무변경은 test_zzz_production_isolation이 확인한다.
"""

import json
import os
import shutil
import sqlite3
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import write_guard
from db.database_manager import DatabaseManager
from extractors import ledger_parser
from extractors.ledger_parser import load_request_ledger, plan_ledger_ids, save_workbook_atomic
from services import ledger_service as ledger_service_module
from services.excel_sync_service import (
    ExcelApplyConflict,
    ExcelSyncService,
    MAX_PENDING_ATTEMPTS,
    QUARANTINE_CONFLICT,
)
from services.ledger_service import LedgerService

HEADERS = ["대장ID", "연번", "의원", "보좌관", "요구자료", "제출", "요청형태", "비고"]
HEADERS_NO_ID = HEADERS[1:]


class Fixture(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="audit_phase10_"))
        self.db_path = self.tmp / "data_requests.db"
        self.json_path = self.tmp / "data_requests.json"
        conn = sqlite3.connect(str(self.db_path))
        DatabaseManager.init_schema(conn)
        conn.close()
        ExcelSyncService._instance = None

    def tearDown(self):
        svc = ExcelSyncService._instance
        if svc is not None:
            svc.stop_file_watcher()
        ExcelSyncService._instance = None
        LedgerService.wait_last_sync(5)
        shutil.rmtree(self.tmp, ignore_errors=True)

    # -- helpers ---------------------------------------------------------
    def make_excel(self, sheets, with_id=True, name="260904 국회 요구자료 목록(테스트).xlsx"):
        path = self.tmp / name
        wb = openpyxl.Workbook()
        _default_sheet = wb.active
        assert _default_sheet is not None
        wb.remove(_default_sheet)
        for sheet, rows in sheets.items():
            ws = wb.create_sheet(sheet)
            ws.append(HEADERS if with_id else HEADERS_NO_ID)
            for row in rows:
                ws.append(row)
        wb.save(path)
        wb.close()
        return path

    def rows(self, path, sheet):
        wb = openpyxl.load_workbook(path, data_only=True)
        try:
            ws = wb[sheet]
            header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            out = []
            for r in range(2, ws.max_row + 1):
                row = {str(header[c - 1]): ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
                if any(v not in (None, "") for v in row.values()):
                    out.append(row)
            return out
        finally:
            wb.close()

    def svc(self):
        return ExcelSyncService.get_instance(self.tmp)

    def ledger(self):
        return LedgerService(self.db_path, self.json_path, self.tmp)

    def db_rows(self):
        conn = sqlite3.connect(str(self.db_path))
        conn.row_factory = sqlite3.Row
        try:
            return {r["ledger_id"]: dict(r) for r in conn.execute("SELECT * FROM request_ledger")}
        finally:
            conn.close()

    def touch_future(self, path, seconds=5):
        future = time.time() + seconds
        os.utime(path, (future, future))


# ---------------------------------------------------------------------------
# R3-01 테스트 격리
# ---------------------------------------------------------------------------
class TestProductionIsolation(Fixture):
    def test_base_dir_defaults_to_db_folder_not_production(self):
        """[T-ISO-2] base_dir를 생략하면 DB 폴더를 쓴다. 운영 폴더로 흐르면 안 된다."""
        svc = LedgerService(self.db_path, self.json_path)
        self.assertEqual(svc.base_dir.resolve(), self.tmp.resolve())
        sync = ExcelSyncService(db_path=self.db_path, json_path=self.json_path)
        self.assertEqual(sync.base_dir.resolve(), self.tmp.resolve())

    def test_guard_blocks_writes_into_production_folder_during_tests(self):
        """[T-ISO-3] 테스트 실행 중 운영 폴더의 마스터 엑셀·큐 쓰기는 거부된다."""
        probe = write_guard.SYSTEM_DIR / "__guard_probe__.xlsx"
        wb = openpyxl.Workbook()
        with self.assertRaises(write_guard.ProductionWriteBlocked):
            save_workbook_atomic(wb, probe)
        wb.close()
        self.assertFalse(probe.exists(), "가드는 파일을 만들기 전에 막아야 한다")
        # 임시 폴더는 막지 않는다.
        write_guard.ensure_writable(self.tmp / "data_requests.db")
        # 소스·테스트 폴더는 운영 데이터가 아니다.
        self.assertFalse(write_guard.is_protected_path(SCRIPTS_DIR / "x.py"))

    def test_guard_error_is_not_mistaken_for_excel_lock(self):
        """가드 예외가 PermissionError로 취급되면 운영 보류 큐에 다시 쓰게 된다."""
        self.assertFalse(issubclass(write_guard.ProductionWriteBlocked, PermissionError))


# ---------------------------------------------------------------------------
# R3-02 보류 큐 재생의 행 충돌 검증
# ---------------------------------------------------------------------------
class TestQueueReplayConflicts(Fixture):
    def test_insert_does_not_overwrite_a_different_row_with_the_same_id(self):
        excel = self.make_excel({"2026": [["REQ-2026-001", "1", "김의원", "", "실제 요구자료", "제출", "시스템", ""]]})
        svc = self.svc()
        with self.assertRaises(ExcelApplyConflict):
            svc._apply_to_excel(excel, {"ledger_id": "REQ-2026-001", "year": "2026",
                                        "requester": "테스트의원", "title": "테스트 제목"}, "insert")
        self.assertEqual(self.rows(excel, "2026")[0]["요구자료"], "실제 요구자료")

    def test_unsynced_web_insert_reports_conflict_instead_of_overwriting(self):
        """DB가 엑셀 ID를 모르는 상태에서 발급된 같은 ID가 기존 행을 덮어쓰지 않는다."""
        excel = self.make_excel({"2026": [["REQ-2026-001", "1", "김의원", "", "자료A", "작성중", "일반", ""]]})
        res = self.ledger().insert_ledger_item({"year": "2026", "requester": "신규의원", "title": "새 자료"})
        self.assertTrue(res["success"])
        self.assertFalse(res["excel_synced"])
        self.assertIn("덮어쓰지 않았습니다", res["excel_message"])
        self.assertEqual(self.rows(excel, "2026")[0]["요구자료"], "자료A")

    def test_polluted_legacy_queue_leaves_real_ledger_intact(self):
        """[T-02c] 운영에서 발견된 오염 큐 패턴(검증 정보 없는 과거 작업)을 재생해도 실제 행이 보존된다."""
        # 열: 연번, 의원, 보좌관, 요구자료, 제출, 요청형태, 비고 (운영 대장처럼 대장ID 열 없음)
        excel = self.make_excel({
            "2026": [["1", "", "김보좌", "업무보고 자료 관련 문의", "제출", "시스템", ""],
                     ["2", "", "이보좌", "처리 대기 건수 현행화 요청", "제출", "시스템", ""]],
            "2025": [["133", "", "박보좌", "", "제출", "시스템", ""]],
        }, with_id=False)
        test_item = {"seq_no": "001", "requester": "테스트의원", "title": "테스트 제목", "status": "작성중"}
        legacy_ops = [
            {"action": "insert", "item": dict(test_item, ledger_id="REQ-2025-001", year="2025")},
            {"action": "update", "item": dict(test_item, ledger_id="REQ-2025-001", year="2026")},
            {"action": "insert", "item": {"ledger_id": "REQ-2026-001", "year": "2026", "seq_no": "001",
                                          "requester": "의원", "title": "원자적 동기화 테스트"}},
            {"action": "insert", "item": {"ledger_id": "REQ-2026-002", "year": "2026", "seq_no": "0517",
                                          "requester": "홍길동 의원", "title": "추가 자동 연계 테스트 자료"}},
            {"action": "delete", "item": {"ledger_id": "REQ-2026-002", "requester": "단위테스트의원",
                                          "title": "단위테스트 요구자료 제목", "seq_no": "TEST-01"}},
        ]
        for i, op in enumerate(legacy_ops):
            op.update({"operation_id": f"legacy-{i}", "timestamp": 0, "attempts": 0})
        (self.tmp / ".excel_pending_queue.json").write_text(json.dumps(legacy_ops, ensure_ascii=False), encoding="utf-8")

        svc = self.svc()
        svc.flush_pending_queue()

        rows_2026 = self.rows(excel, "2026")
        rows_2025 = self.rows(excel, "2025")
        self.assertEqual([r["요구자료"] for r in rows_2026],
                         ["업무보고 자료 관련 문의", "처리 대기 건수 현행화 요청"])
        self.assertTrue(all(r["제출"] == "제출" for r in rows_2026), "실제 행이 [삭제]되면 안 된다")
        self.assertEqual([r["보좌관"] for r in rows_2025], ["박보좌"], "2025 실제 행이 옮겨지거나 지워지면 안 된다")
        self.assertEqual(svc.pending_count(), 0)
        self.assertGreaterEqual(svc.failed_count(), 3)
        self.assertTrue(list((self.tmp / ".ledger_backup").glob("*.xlsx")), "재생 직전 강제 백업이 있어야 한다")

    def test_web_update_with_fingerprint_still_applies(self):
        excel = self.make_excel({"2026": [["REQ-2026-001", "1", "김의원", "", "원래 제목", "작성중", "일반", ""]]})
        self.svc().sync_excel_to_db()
        res = self.ledger().update_ledger_item("REQ-2026-001", {"title": "고친 제목", "status": "제출"})
        self.assertTrue(res["excel_synced"], res)
        row = self.rows(excel, "2026")[0]
        self.assertEqual((row["요구자료"], row["제출"]), ("고친 제목", "제출"))

    def test_conflict_quarantine_is_not_protected_but_retry_limit_is(self):
        svc = self.svc()
        svc._quarantine_pending_operations([svc._new_operation({"ledger_id": "REQ-2026-010"}, "insert")],
                                           QUARANTINE_CONFLICT, "different row")
        svc._quarantine_pending_operations([svc._new_operation({"ledger_id": "REQ-2026-011"}, "update")])
        protected = svc.protected_ledger_ids()
        self.assertNotIn("REQ-2026-010", protected)
        self.assertIn("REQ-2026-011", protected)
        self.assertEqual(svc.failed_count(), 2)


# ---------------------------------------------------------------------------
# R3-05 ID 스탬핑 재사용
# ---------------------------------------------------------------------------
class TestStampingIdReuse(Fixture):
    def test_tombstoned_id_is_not_given_to_a_new_manual_row(self):
        excel = self.make_excel({"2026": [
            ["REQ-2026-002", "2", "의원B", "", "제목B", "제출", "", ""],
            ["REQ-2026-003", "3", "의원C", "", "제목C", "제출", "", ""],
            [None, "4", "새의원", "", "새로 추가한 수기 행", "작성중", "", ""],
        ]})
        wb = openpyxl.load_workbook(excel, data_only=True)
        records = [{"ledger_id": "REQ-2026-001", "year": "2026", "title": "제목A", "requester": "의원A", "status": "[삭제]"},
                   {"ledger_id": "REQ-2026-002", "year": "2026", "title": "제목B", "requester": "의원B"},
                   {"ledger_id": "REQ-2026-003", "year": "2026", "title": "제목C", "requester": "의원C"}]
        assignments, _ = plan_ledger_ids(wb, records)
        wb.close()
        self.assertEqual(assignments, [("2026", 4, "REQ-2026-004")])

    def test_id_of_item_moved_to_another_year_sheet_is_not_reissued(self):
        excel = self.make_excel({
            "2026": [[None, "5", "새의원", "", "새 수기 행", "작성중", "", ""],
                     ["REQ-2026-002", "6", "의원X", "", "기존", "제출", "", ""]],
            "2025": [["REQ-2026-001", "1", "의원M", "", "2026→2025로 옮긴 항목", "제출", "", ""]],
        })
        wb = openpyxl.load_workbook(excel, data_only=True)
        records = [{"ledger_id": "REQ-2026-001", "year": "2025", "title": "2026→2025로 옮긴 항목", "requester": "의원M"},
                   {"ledger_id": "REQ-2026-002", "year": "2026", "title": "기존", "requester": "의원X"}]
        assignments, _ = plan_ledger_ids(wb, records)
        wb.close()
        self.assertEqual(len(assignments), 1)
        self.assertEqual(assignments[0][:2], ("2026", 2))
        self.assertNotIn(assignments[0][2], {"REQ-2026-001", "REQ-2026-002"})

    def test_legacy_sheet_keeps_positional_ids(self):
        """명시 ID가 전혀 없는 레거시 시트는 기존 위치 기반 ID를 그대로 고정한다."""
        excel = self.make_excel({"2026": [
            [None, "1", "의원A", "", "제목A(엑셀에서 고침)", "제출", "", ""],
            [None, "2", "의원B", "", "제목B", "[삭제]", "", ""],
        ]})
        wb = openpyxl.load_workbook(excel, data_only=True)
        records = [{"ledger_id": "REQ-2026-001", "year": "2026", "title": "제목A", "requester": "의원A"},
                   {"ledger_id": "REQ-2026-002", "year": "2026", "title": "제목B", "requester": "의원B", "status": "[삭제]"}]
        assignments, _ = plan_ledger_ids(wb, records)
        wb.close()
        self.assertEqual([a[2] for a in assignments], ["REQ-2026-001", "REQ-2026-002"])


# ---------------------------------------------------------------------------
# R3-06 파이프라인 병합 규칙
# ---------------------------------------------------------------------------
class TestPipelineMergeRule(Fixture):
    def _seed_db(self, **overrides):
        base = {"ledger_id": "REQ-2026-001", "year": "2026", "seq_no": "1", "party": "", "requester": "김의원",
                "aide": "", "title": "자료A", "details": "", "request_date": "", "deadline": "", "submit_date": "",
                "department": "", "status": "작성중", "note": "웹 메모", "request_type": "일반", "linked_doc_id": "",
                "created_at": "2026-09-01 09:00:00", "updated_at": "2026-09-05 09:00:00"}
        base.update(overrides)
        conn = sqlite3.connect(str(self.db_path))
        conn.execute(f"INSERT INTO request_ledger ({', '.join(base)}) VALUES ({', '.join('?' * len(base))})",
                     list(base.values()))
        conn.commit()
        conn.close()

    def test_excel_value_wins_even_if_db_row_was_edited(self):
        self.make_excel({"2026": [["REQ-2026-001", "1", "김의원", "", "자료A", "제출", "일반", "웹 메모"]]})
        self._seed_db()
        item = next(i for i in load_request_ledger(self.tmp, [], merge_db_records=True)
                    if i["ledger_id"] == "REQ-2026-001")
        self.assertEqual(item["status"], "제출", "서버가 꺼진 동안의 엑셀 수정이 버려지면 안 된다")
        self.assertEqual(item["created_at"], "2026-09-01 09:00:00", "created_at은 재빌드마다 초기화되면 안 된다")

    def test_protected_ids_and_tombstones_keep_db_values(self):
        self.make_excel({"2026": [["REQ-2026-001", "1", "김의원", "", "자료A", "제출", "일반", ""],
                                  ["REQ-2026-002", "2", "이의원", "", "자료B", "제출", "일반", ""]]})
        self._seed_db()
        self._seed_db(ledger_id="REQ-2026-002", requester="이의원", title="자료B", status="[삭제]")
        items = {i["ledger_id"]: i for i in load_request_ledger(self.tmp, [], merge_db_records=True,
                                                                 protected_ids={"REQ-2026-001"})}
        self.assertEqual(items["REQ-2026-001"]["status"], "작성중")
        self.assertEqual(items["REQ-2026-002"]["status"], "[삭제]")

    def test_prepare_ledger_for_rebuild_syncs_excel_edits_first(self):
        import extract_and_build_db
        excel = self.make_excel({"2026": [["REQ-2026-001", "1", "김의원", "", "자료A", "작성중", "일반", ""]]})
        svc = ExcelSyncService(base_dir=self.tmp, db_path=self.db_path, json_path=self.json_path)
        svc.sync_excel_to_db()
        wb = openpyxl.load_workbook(excel)
        wb["2026"].cell(2, 6, "제출")
        wb.save(excel)
        wb.close()
        self.json_path.unlink(missing_ok=True)
        with mock.patch.object(extract_and_build_db, "SYSTEM_DIR", self.tmp), \
                mock.patch.object(extract_and_build_db, "DB_PATH", self.db_path), \
                mock.patch.object(extract_and_build_db, "JSON_PATH", self.json_path):
            protected = extract_and_build_db.prepare_ledger_for_rebuild()
        self.assertEqual(protected, set())
        self.assertEqual(self.db_rows()["REQ-2026-001"]["status"], "제출")
        self.assertFalse(self.json_path.exists(), "재빌드 전 동기화는 JSON을 따로 쓰지 않는다")


# ---------------------------------------------------------------------------
# R3-07 / R3-08 / R3-10 / 입력 정규화
# ---------------------------------------------------------------------------
class TestLedgerWritePaths(Fixture):
    def test_non_permission_excel_failure_is_queued_and_protected(self):
        """[T-07a]"""
        self.make_excel({"2026": [["REQ-2026-001", "1", "김의원", "", "원래 제목", "작성중", "메일", ""]]})
        svc = self.svc()
        svc.sync_excel_to_db()
        with mock.patch.object(ExcelSyncService, "_apply_to_excel", side_effect=ValueError("simulated")):
            res = self.ledger().update_ledger_item("REQ-2026-001", {"status": "제출"})
        self.assertTrue(res["excel_pending"], res)
        svc.sync_excel_to_db()
        self.assertEqual(self.db_rows()["REQ-2026-001"]["status"], "제출", "웹 수정이 엑셀 값으로 되돌려지면 안 된다")

    def test_retry_limit_quarantine_keeps_protecting_db_value(self):
        """[T-07b]"""
        self.make_excel({"2026": [["REQ-2026-001", "1", "김의원", "", "원래 제목", "작성중", "메일", ""]]})
        svc = self.svc()
        svc.sync_excel_to_db()
        with mock.patch.object(ExcelSyncService, "_apply_to_excel", side_effect=ValueError("always")):
            self.ledger().update_ledger_item("REQ-2026-001", {"status": "제출"})
            for _ in range(MAX_PENDING_ATTEMPTS + 1):
                svc.flush_pending_queue()
        self.assertEqual(svc.pending_count(), 0)
        self.assertEqual(svc.failed_count(), 1)
        svc.sync_excel_to_db()
        self.assertEqual(self.db_rows()["REQ-2026-001"]["status"], "제출")

    def test_numbering_uses_id_prefix_after_year_change(self):
        """[T-08]"""
        ledger = self.ledger()
        for i in range(3):
            self.assertTrue(ledger.insert_ledger_item({"year": "2026", "title": f"t{i}", "requester": "의원"})["success"])
        self.assertTrue(ledger.update_ledger_item("REQ-2026-003", {"year": "2025"})["success"])
        res = ledger.insert_ledger_item({"year": "2026", "title": "new", "requester": "의원"})
        self.assertTrue(res["success"], res)
        self.assertEqual(res["item"]["ledger_id"], "REQ-2026-004")

    def test_excel_row_deletion_applies_to_system_request_type(self):
        """[T-10] 요청형태 '시스템'도 엑셀에서 지우면 삭제된다."""
        excel = self.make_excel({"2026": [
            ["REQ-2026-001", "1", "김의원", "", "자료A", "제출", "시스템", ""],
            ["REQ-2026-002", "2", "이의원", "", "자료B", "제출", "시스템", ""],
            ["REQ-2026-003", "3", "박의원", "", "자료C", "제출", "메일", ""],
        ]})
        svc = self.svc()
        svc.sync_excel_to_db()
        wb = openpyxl.load_workbook(excel)
        wb["2026"].delete_rows(2)
        wb.save(excel)
        wb.close()
        svc.sync_excel_to_db()
        self.assertEqual(self.db_rows()["REQ-2026-001"]["status"], "[삭제]")

    def test_web_insert_never_written_to_excel_is_not_deleted(self):
        ledger = self.ledger()
        res = ledger.insert_ledger_item({"year": "2026", "title": "엑셀 없을 때 등록", "requester": "의원"})
        self.assertTrue(res["success"])
        self.make_excel({"2026": [["REQ-2026-900", "1", "김의원", "", "자료A", "제출", "시스템", ""]]})
        self.svc().sync_excel_to_db()
        self.assertNotEqual(self.db_rows()[res["item"]["ledger_id"]]["status"], "[삭제]")

    def test_null_fields_are_not_stored_as_none_string(self):
        ledger = self.ledger()
        res = ledger.insert_ledger_item({"year": "2026", "title": "제목", "requester": "의원", "note": None})
        lid = res["item"]["ledger_id"]
        ledger.update_ledger_item(lid, {"aide": None})
        row = self.db_rows()[lid]
        self.assertEqual((row["note"], row["aide"]), ("", ""))

    def test_excel_sync_writes_json_outside_workbook_lock(self):
        self.make_excel({"2026": [["REQ-2026-001", "1", "김의원", "", "자료A", "제출", "메일", ""]]})
        svc = self.svc()
        seen = []
        real = LedgerService.sync_json_file

        def spy(service_self, async_mode=True):
            seen.append(svc.lock.locked())
            return real(service_self, async_mode=async_mode)

        with mock.patch.object(LedgerService, "sync_json_file", spy):
            svc.sync_excel_to_db()
        self.assertEqual(seen, [False])

    def test_async_json_sync_coalesces_bursts(self):
        ledger = self.ledger()
        ledger.insert_ledger_item({"year": "2026", "title": "제목", "requester": "의원"})
        LedgerService.wait_last_sync(5)
        calls = []
        real_dump = ledger_service_module.json.dump

        def slow_dump(*args, **kwargs):
            calls.append(1)
            time.sleep(0.2)
            return real_dump(*args, **kwargs)

        with mock.patch.object(ledger_service_module.json, "dump", side_effect=slow_dump):
            for _ in range(6):
                ledger.sync_json_file(async_mode=True)
            deadline = time.time() + 10
            while ledger_service_module._JSON_RUNNING.get(str(self.json_path.resolve())) and time.time() < deadline:
                time.sleep(0.05)
        self.assertLessEqual(len(calls), 2, "연속 호출은 최대 두 번의 재작성으로 합쳐져야 한다")
        self.assertTrue(json.loads(self.json_path.read_text(encoding="utf-8"))["request_ledger"])


# ---------------------------------------------------------------------------
# R3-09 watcher
# ---------------------------------------------------------------------------
class TestWatcherWhileExcelOpen(Fixture):
    def test_watcher_reads_locked_workbook_but_defers_queue(self):
        excel = self.make_excel({"2026": [["REQ-2026-001", "1", "김의원", "", "자료A", "작성중", "메일", ""]]})
        svc = self.svc()
        svc.sync_excel_to_db()
        with mock.patch.object(ExcelSyncService, "is_file_locked", return_value=True):
            self.ledger().update_ledger_item("REQ-2026-001", {"note": "웹 메모"})
            self.assertEqual(svc.pending_count(), 1)
            wb = openpyxl.load_workbook(excel)
            wb["2026"].cell(2, 5, "엑셀에서 고친 제목")
            wb.save(excel)
            wb.close()
            self.touch_future(excel)
            svc.start_file_watcher(poll_interval=0.2)
            deadline = time.time() + 10
            while time.time() < deadline and self.db_rows()["REQ-2026-001"]["title"] != "엑셀에서 고친 제목":
                time.sleep(0.1)
            svc.stop_file_watcher()
            time.sleep(0.4)
        self.assertEqual(self.db_rows()["REQ-2026-001"]["title"], "엑셀에서 고친 제목")
        self.assertEqual(svc.pending_count(), 1, "잠긴 동안에는 보류 큐를 쓰지 않는다")

    def test_save_right_after_web_write_is_not_missed(self):
        """웹 쓰기 직후 1.5초 안의 사용자 저장도 감지한다."""
        excel = self.make_excel({"2026": [["REQ-2026-001", "1", "김의원", "", "자료A", "작성중", "메일", ""]]})
        svc = self.svc()
        svc.sync_excel_to_db()
        mtime = excel.stat().st_mtime
        svc._last_known_excel_mtime = mtime
        svc._last_web_write_time = mtime
        wb = openpyxl.load_workbook(excel)
        wb["2026"].cell(2, 6, "제출")
        wb.save(excel)
        wb.close()
        os.utime(excel, (mtime + 0.5, mtime + 0.5))
        svc.start_file_watcher(poll_interval=0.2)
        deadline = time.time() + 10
        while time.time() < deadline and self.db_rows()["REQ-2026-001"]["status"] != "제출":
            time.sleep(0.1)
        svc.stop_file_watcher()
        time.sleep(0.4)
        self.assertEqual(self.db_rows()["REQ-2026-001"]["status"], "제출")


if __name__ == "__main__":
    unittest.main()
