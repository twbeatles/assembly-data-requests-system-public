# -*- coding: utf-8 -*-
"""
PROJECT_AUDIT.md (2026-09-13 회차) ISSUE-001~008 및 기능 공백 개선 회귀 테스트.

각 테스트는 감사에서 **실제로 재현한 경로**를 그대로 고정한다.
- ISSUE-001: parse_all의 워크스페이스 루트 하드코딩 → 마크다운 전멸 + 빈 DB 스왑
- ISSUE-002: sync_excel_to_db 진행 중 들어온 웹 등록이 같은 실행에서 tombstone
- ISSUE-003: flush_pending_queue와 웹 쓰기의 락 비대칭 → 워크북 lost update
- ISSUE-004: `대장ID`를 대상 연도 시트에서만 찾아 중복·유령 항목·필드 소실
- ISSUE-005: launcher_gui 완료 핸들러의 UnboundLocalError
- ISSUE-006: `[삭제]` tombstone이 파생 산출물에 재유출
- ISSUE-007: cmd.exe 인자 조립이 파일명의 `&`/`^`/`%`를 해석
- ISSUE-008: 문서 1건 실패가 파이프라인 전체를 중단

소스 문자열 grep이 아니라 **실행 가능한 검증**으로 쓴다. 예전 회차의
`test_gui_reports_pipeline_failure`는 문자열만 확인해서 실행 불가능한 코드에도
통과했고, 그 때문에 ISSUE-005가 오래 가려져 있었다.
"""

import importlib
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import symtable
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest import mock

REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = REPO_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import openpyxl

import parse_all
import system_config
from db.database_manager import DatabaseManager
from services.excel_sync_service import ExcelSyncService, MAX_PENDING_ATTEMPTS
from services.ledger_service import LedgerService

HEADERS_WITH_ID = ["대장ID", "연번", "의원", "요구자료", "제출", "요청형태", "비고"]


class LedgerFixtureMixin:
    """임시 디렉터리에 마스터 엑셀 + SQLite 대장을 세운다."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="audit_phase9_"))
        self.db_path = self.tmp / "data_requests.db"
        self.json_path = self.tmp / "data_requests.json"
        conn = sqlite3.connect(str(self.db_path))
        DatabaseManager.init_schema(conn)
        conn.close()
        ExcelSyncService._instance = None

    def tearDown(self):
        svc = ExcelSyncService._instance
        if svc is not None:
            svc.stop_file_watcher()
        ExcelSyncService._instance = None
        shutil.rmtree(self.tmp, ignore_errors=True)

    def make_excel(self, rows, sheets=("2026",)):
        path = self.tmp / "260904 국회 요구자료 목록(테스트).xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = sheets[0]
        ws.append(HEADERS_WITH_ID)
        for row in rows:
            ws.append(row)
        for extra in sheets[1:]:
            wb.create_sheet(extra).append(HEADERS_WITH_ID)
        wb.save(path)
        wb.close()
        return path

    def sync_service(self):
        svc = ExcelSyncService(base_dir=self.tmp, db_path=self.db_path, json_path=self.json_path)
        ExcelSyncService._instance = svc
        return svc

    def ledger_service(self):
        return LedgerService(self.db_path, self.json_path, self.tmp)

    def db_rows(self):
        conn = sqlite3.connect(str(self.db_path))
        conn.row_factory = sqlite3.Row
        try:
            return {r["ledger_id"]: dict(r) for r in conn.execute("SELECT * FROM request_ledger")}
        finally:
            conn.close()

    def sheet_rows(self, path, sheet):
        wb = openpyxl.load_workbook(str(path), data_only=True)
        try:
            ws = wb[sheet]
            header = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
            out = []
            for r in range(2, ws.max_row + 1):
                row = {header[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
                if any(v is not None and str(v).strip() for v in row.values()):
                    out.append(row)
            return out
        finally:
            wb.close()


# ---------------------------------------------------------------------------
# ISSUE-001 : 워크스페이스 루트 판정과 마크다운 보호
# ---------------------------------------------------------------------------
class TestWorkspaceRootDetection(unittest.TestCase):
    """[T-001] 연도 폴더 이름과 무관하게 루트를 찾고, 잘못 찾으면 아무것도 지우지 않는다."""

    def _build_workspace(self, year_dirs):
        ws = Path(tempfile.mkdtemp(prefix="audit_phase9_root_"))
        self.addCleanup(shutil.rmtree, ws, True)
        sys_dir = ws / "국회자료요구_스마트시스템(테스트팀)"
        (sys_dir / "scripts").mkdir(parents=True)
        for name in ("parse_all.py", "system_config.py"):
            shutil.copy2(SCRIPTS_DIR / name, sys_dir / "scripts" / name)
        md_root = sys_dir / "_parsed_markdown"
        for y in year_dirs:
            (ws / y).mkdir()
            (ws / y / f"{y}_요구자료.hwp").write_text("원본", encoding="utf-8")
            (md_root / y).mkdir(parents=True)
            (md_root / y / f"{y}_요구자료.hwp.md").write_text("# 전문", encoding="utf-8")
        return ws, sys_dir, md_root

    def _load_parse_all(self, sys_dir):
        """해당 시스템 폴더 기준으로 parse_all을 새로 로드한다 (모듈 전역이 경로를 정한다)."""
        for name in ("parse_all", "system_config"):
            sys.modules.pop(name, None)
        scripts = str(sys_dir / "scripts")
        sys.path.insert(0, scripts)
        try:
            module = importlib.import_module("parse_all")
            cfg = importlib.import_module("system_config")
            return module, cfg
        finally:
            sys.path.remove(scripts)

    def tearDown(self):
        # 저장소 기준 모듈로 되돌려 다른 테스트에 영향을 주지 않는다.
        for name in ("parse_all", "system_config"):
            sys.modules.pop(name, None)
        importlib.import_module("system_config")
        importlib.import_module("parse_all")

    def test_root_detection_matches_system_config_for_any_year(self):
        """2026 폴더가 없어도 parse_all과 system_config의 루트 판정이 같아야 한다."""
        for year_dirs in (["2025", "2026"], ["2027", "2028"], ["2031"]):
            with self.subTest(year_dirs=year_dirs):
                ws, sys_dir, _md = self._build_workspace(year_dirs)
                mod, cfg = self._load_parse_all(sys_dir)
                self.assertEqual(
                    mod.ROOT_DIR.resolve(),
                    cfg.find_workspace_root(mod.SYSTEM_DIR).resolve(),
                    "parse_all과 system_config의 워크스페이스 루트 판정이 달라지면 안 된다",
                )
                self.assertNotEqual(
                    mod.ROOT_DIR.resolve(), mod.SYSTEM_DIR.resolve(),
                    "연도 폴더가 있는데 루트가 시스템 폴더로 무너지면 안 된다",
                )
                self.assertEqual(len(mod.find_target_files(mod.ROOT_DIR)), len(year_dirs))

    def test_prune_keeps_everything_when_no_source_document_found(self):
        """원본을 한 건도 못 찾으면 고아 마크다운을 하나도 지우지 않는다."""
        ws, sys_dir, md_root = self._build_workspace(["2027", "2028"])
        mod, _cfg = self._load_parse_all(sys_dir)
        before = sorted(p.name for p in md_root.rglob("*.md"))

        pruned = mod.prune_orphan_markdowns(mod.SYSTEM_DIR, md_root, source_count=0)

        self.assertEqual(pruned, 0, "원본 0건이면 경로 판정 실패이지 문서 삭제가 아니다")
        self.assertEqual(sorted(p.name for p in md_root.rglob("*.md")), before)

    def test_prune_blocks_mass_deletion_over_ratio(self):
        """삭제 대상이 임계 건수·비율을 함께 넘으면 전부 보류한다."""
        base = Path(tempfile.mkdtemp(prefix="audit_phase9_prune_"))
        self.addCleanup(shutil.rmtree, base, True)
        md_dir = base / "_parsed_markdown" / "2026"
        md_dir.mkdir(parents=True)
        for i in range(10):
            (md_dir / f"ghost{i}.hwp.md").write_text("x", encoding="utf-8")

        pruned = parse_all.prune_orphan_markdowns(base, md_dir.parent, source_count=3)

        self.assertEqual(pruned, 0, "10/10건(100%) 삭제 시도는 안전장치가 막아야 한다")
        self.assertEqual(len(list(md_dir.glob("*.md"))), 10)

    def test_prune_still_removes_small_orphan_sets(self):
        """소규모 정리는 정상 동작해야 한다 (안전장치가 과하게 걸리면 안 된다)."""
        base = Path(tempfile.mkdtemp(prefix="audit_phase9_prune2_"))
        self.addCleanup(shutil.rmtree, base, True)
        year = base / "2026"
        year.mkdir()
        md_dir = base / "_parsed_markdown" / "2026"
        md_dir.mkdir(parents=True)
        for i in range(9):
            (year / f"keep{i}.hwp").write_text("원본", encoding="utf-8")
            (md_dir / f"keep{i}.hwp.md").write_text("x", encoding="utf-8")
        (md_dir / "ghost.hwp.md").write_text("x", encoding="utf-8")

        pruned = parse_all.prune_orphan_markdowns(base, md_dir.parent, source_count=9)

        self.assertEqual(pruned, 1)
        self.assertFalse((md_dir / "ghost.hwp.md").exists())
        self.assertEqual(len(list(md_dir.glob("keep*.md"))), 9)

    def test_zero_sources_with_existing_markdown_fails_the_stage(self):
        """원본 0건 + 기존 마크다운 있음 = 경로 이상. 다음 단계로 넘어가면 안 된다."""
        ws, sys_dir, md_root = self._build_workspace(["2027"])
        mod, _cfg = self._load_parse_all(sys_dir)
        setattr(mod, "ROOT_DIR", mod.SYSTEM_DIR)  # 경로 오판 상황을 강제한다

        result = mod.main()

        self.assertFalse(result["success"], "경로 이상이면 파싱 단계를 실패로 보고해야 한다")
        self.assertIn("원본", result.get("error", ""))
        self.assertTrue(list(md_root.rglob("*.md")), "마크다운이 남아 있어야 한다")


# ---------------------------------------------------------------------------
# ISSUE-008 / ISSUE-007 : 파싱 단계의 실패 처리와 인자 조립
# ---------------------------------------------------------------------------
class TestParseStageContract(unittest.TestCase):
    """[T-007][T-008] 부분 실패 분리와 cmd.exe 인자 안전성."""

    def test_partial_failure_does_not_fail_the_stage(self):
        """문서 3건 중 1건이 실패해도 단계는 성공이고, 실패는 따로 보고된다."""
        base = Path(tempfile.mkdtemp(prefix="audit_phase9_parse_"))
        self.addCleanup(shutil.rmtree, base, True)
        year = base / "2026"
        year.mkdir()
        files = []
        for i in range(3):
            f = year / f"doc{i}.hwp"
            f.write_text("원본", encoding="utf-8")
            files.append(f)

        def fake_parse(file_path, output_dir, force=False):
            failed = file_path.name == "doc1.hwp"
            return {
                "file": file_path.name, "out_file": file_path.name + ".md",
                "success": not failed, "elapsed": 0.0,
                "size_bytes": 0 if failed else 10,
                "error": "kordoc 변환 실패" if failed else "",
            }

        with mock.patch.object(parse_all, "ROOT_DIR", base), \
             mock.patch.object(parse_all, "OUTPUT_DIR", base / "_parsed_markdown"), \
             mock.patch.object(parse_all, "SYSTEM_DIR", base), \
             mock.patch.object(parse_all, "parse_single_file", side_effect=fake_parse), \
             mock.patch.object(parse_all.shutil, "which", return_value="kordoc.cmd"):
            result = parse_all.main()

        self.assertTrue(result["success"], "문서 1건 실패로 파이프라인 전체를 막으면 안 된다")
        self.assertEqual(result["partial_failures"], 1)
        self.assertEqual(result["failed_files"], ["doc1.hwp"])

    def test_majority_failure_fails_the_stage(self):
        """과반이 실패하면 단계 실패로 보고해 산출물 생성을 막는다."""
        base = Path(tempfile.mkdtemp(prefix="audit_phase9_parse2_"))
        self.addCleanup(shutil.rmtree, base, True)
        year = base / "2026"
        year.mkdir()
        for i in range(3):
            (year / f"doc{i}.hwp").write_text("원본", encoding="utf-8")

        def all_fail(file_path, output_dir, force=False):
            return {"file": file_path.name, "out_file": "", "success": False,
                    "elapsed": 0.0, "error": "kordoc 없음"}

        with mock.patch.object(parse_all, "ROOT_DIR", base), \
             mock.patch.object(parse_all, "OUTPUT_DIR", base / "_parsed_markdown"), \
             mock.patch.object(parse_all, "SYSTEM_DIR", base), \
             mock.patch.object(parse_all, "parse_single_file", side_effect=all_fail), \
             mock.patch.object(parse_all.shutil, "which", return_value="kordoc.cmd"):
            result = parse_all.main()

        self.assertFalse(result["success"])
        self.assertEqual(result["partial_failures"], 0)

    def test_missing_kordoc_is_reported_explicitly(self):
        """KorDoc 미설치는 문서별 실패가 아니라 구성요소 부재로 알려야 한다."""
        base = Path(tempfile.mkdtemp(prefix="audit_phase9_parse3_"))
        self.addCleanup(shutil.rmtree, base, True)
        year = base / "2026"
        year.mkdir()
        (year / "doc.hwp").write_text("원본", encoding="utf-8")

        with mock.patch.object(parse_all, "ROOT_DIR", base), \
             mock.patch.object(parse_all, "OUTPUT_DIR", base / "_parsed_markdown"), \
             mock.patch.object(parse_all, "SYSTEM_DIR", base), \
             mock.patch.object(parse_all.shutil, "which", return_value=None):
            result = parse_all.main()

        self.assertFalse(result["success"])
        self.assertIn("KorDoc", result.get("error", ""))

    def test_dangerous_filenames_are_routed_to_safe_copies(self):
        """cmd.exe가 해석하는 문자는 안전 복사 경로를 타야 한다."""
        for name in ("국회&대외기관.hwp", "보고서^최종.hwp", "%USERNAME%.hwp"):
            with self.subTest(name=name):
                self.assertTrue(parse_all.needs_safe_copy(Path("C:/docs") / name))
        # 실무에서 흔한 이름은 불필요한 복사를 하지 않는다.
        for name in ("자료(2026) 최종.hwp", "국회_요구자료 1-2.hwp"):
            with self.subTest(name=name):
                self.assertFalse(parse_all.needs_safe_copy(Path("C:/docs") / name))

    def test_kordoc_cmd_is_a_fully_quoted_command_string(self):
        """리스트로 넘기면 list2cmdline이 인용을 깨뜨린다. 문자열이어야 한다."""
        line = parse_all.kordoc_cmd(r"C:\docs\보고서 최종.hwp", "-o", r"C:\out\x.md")
        self.assertIsInstance(line, str, "리스트를 넘기면 cmd 인용이 깨진다")
        self.assertIn('"C:\\docs\\보고서 최종.hwp"', line)
        self.assertIn('"-o"', line)
        self.assertTrue(line.startswith("cmd.exe /s /c "))

    @unittest.skipUnless(sys.platform == "win32", "cmd.exe 동작은 Windows에서만 확인한다")
    def test_ampersand_filename_cannot_inject_a_command(self):
        """`자료&ver.hwp` 같은 이름이 부가 명령을 실행하면 안 된다 (실제 cmd 실행)."""
        tmp = Path(tempfile.mkdtemp(prefix="audit_phase9_cmd_"))
        self.addCleanup(shutil.rmtree, tmp, True)
        fake = tmp / "fake kordoc.cmd"
        fake.write_text("@echo off\r\necho ARG_IN=[%~1]\r\necho # parsed > \"%~3\"\r\n", encoding="ascii")
        src = tmp / "AAA&ver.hwp"
        src.write_text("원본", encoding="utf-8")
        out_file = tmp / "out" / "AAA&ver.hwp.md"

        with mock.patch.object(parse_all, "KORDOC_CMD", str(fake)):
            with parse_all.kordoc_paths(src, out_file) as (safe_src, safe_out):
                proc = subprocess.run(
                    parse_all.kordoc_cmd(str(safe_src), "-o", str(safe_out)),
                    capture_output=True, text=True, encoding="utf-8", errors="replace",
                )

        self.assertNotIn("Microsoft Windows", proc.stdout or "",
                         "`&` 뒤가 별도 명령으로 실행되면 안 된다")
        self.assertIn(str(safe_src), proc.stdout or "", "경로가 잘리지 않고 그대로 전달돼야 한다")
        self.assertTrue(out_file.exists(), "결과 마크다운이 원래 위치로 돌아와야 한다")


# ---------------------------------------------------------------------------
# ISSUE-004 : 워크북 전체 ID 탐색
# ---------------------------------------------------------------------------
class TestWorkbookWideIdLookup(LedgerFixtureMixin, unittest.TestCase):
    """[T-004] 연도가 바뀌어도 항목이 복제되거나 필드가 지워지지 않는다."""

    def test_year_change_moves_the_row_and_keeps_every_field(self):
        excel = self.make_excel(
            [["REQ-2026-001", "1", "이의원", "올해자료", "작성중", "일반", ""]],
            sheets=("2026",),
        )
        wb = openpyxl.load_workbook(excel)
        ws = wb.create_sheet("2025", 0)
        ws.append(HEADERS_WITH_ID)
        ws.append(["REQ-2025-001", "1", "김의원", "작년자료", "완료", "일반", "원래비고"])
        wb.save(excel)
        wb.close()

        svc = self.sync_service()
        svc.sync_excel_to_db()
        ledger = self.ledger_service()

        res = ledger.update_ledger_item("REQ-2025-001", {"year": "2026", "title": "작년자료(연도정정)"})
        self.assertTrue(res["success"])

        self.assertEqual(self.sheet_rows(excel, "2025"), [], "원래 시트의 행이 남으면 중복 ID가 된다")
        moved = [r for r in self.sheet_rows(excel, "2026") if r["대장ID"] == "REQ-2025-001"]
        self.assertEqual(len(moved), 1, "옮겨 간 행은 정확히 하나여야 한다")
        self.assertEqual(moved[0]["의원"], "김의원", "새 행을 만들 때 전체 필드를 기록해야 한다")
        self.assertEqual(moved[0]["요구자료"], "작년자료(연도정정)")

        svc._id_stamp_fingerprint = None
        svc.sync_excel_to_db()
        active = {lid for lid, r in self.db_rows().items() if r["status"] != "[삭제]"}
        self.assertEqual(active, {"REQ-2025-001", "REQ-2026-001"},
                         "연도 변경이 유령 항목을 만들면 안 된다")
        self.assertEqual(self.db_rows()["REQ-2025-001"]["requester"], "김의원",
                         "요구자명이 빈 칸으로 덮어써지면 안 된다")

    def test_delete_finds_the_row_in_any_sheet(self):
        excel = self.make_excel(
            [["REQ-2026-001", "1", "이의원", "올해자료", "작성중", "일반", ""]], sheets=("2026",)
        )
        wb = openpyxl.load_workbook(excel)
        ws = wb.create_sheet("2025", 0)
        ws.append(HEADERS_WITH_ID)
        ws.append(["REQ-2025-001", "1", "김의원", "작년자료", "완료", "일반", ""])
        wb.save(excel)
        wb.close()

        svc = self.sync_service()
        svc.sync_excel_to_db()
        # DB의 연도를 일부러 어긋나게 만들어 대상 시트 추정이 빗나가게 한다.
        conn = sqlite3.connect(str(self.db_path))
        conn.execute("UPDATE request_ledger SET year='2026' WHERE ledger_id='REQ-2025-001'")
        conn.commit()
        conn.close()

        res = self.ledger_service().delete_ledger_item("REQ-2025-001")
        self.assertTrue(res["success"])
        self.assertTrue(res["excel_synced"], "다른 시트에 있어도 삭제 표시를 해야 한다")
        row = [r for r in self.sheet_rows(excel, "2025") if r["대장ID"] == "REQ-2025-001"][0]
        self.assertEqual(row["제출"], "[삭제]")

    def test_delete_reports_failure_when_row_is_absent(self):
        """엑셀에 행이 없으면 반영했다고 보고하지 않는다.

        예전에는 대상 행을 못 찾아도 True를 돌려줘서 "✓ 즉시 반영되었습니다"가
        떴고, 사용자는 엑셀이 정리된 줄 알았다.
        """
        self.make_excel([["REQ-2026-001", "1", "이의원", "올해자료", "작성중", "일반", ""]])
        svc = self.sync_service()
        svc.sync_excel_to_db()

        # 엑셀에는 없고 DB에만 있는 항목을 만든다 (엑셀 잠금 중 등록된 상태와 같다).
        now = time.strftime("%Y-%m-%d %H:%M:%S")
        conn = sqlite3.connect(str(self.db_path))
        conn.execute(
            "INSERT INTO request_ledger (ledger_id, year, requester, title, status, "
            "request_type, created_at, updated_at) VALUES (?,?,?,?,?,?,?,?)",
            ("REQ-2026-777", "2026", "없는의원", "엑셀에 없는 항목", "작성중", "메일", now, now),
        )
        conn.commit()
        conn.close()

        res = self.ledger_service().delete_ledger_item("REQ-2026-777")
        self.assertTrue(res["success"], "DB 삭제 자체는 성공한다")
        self.assertFalse(res["excel_synced"], "엑셀에서 못 찾았으면 반영했다고 하면 안 된다")
        self.assertIn("찾지 못해", res["excel_message"])


# ---------------------------------------------------------------------------
# ISSUE-002 / ISSUE-003 : 동시성
# ---------------------------------------------------------------------------
class TestLedgerConcurrency(LedgerFixtureMixin, unittest.TestCase):
    """[T-002][T-003] 동시 실행에서 사용자의 등록이 사라지지 않는다."""

    def test_insert_during_excel_sync_is_not_tombstoned(self):
        rows = [[f"REQ-2026-{i:03d}", str(i), f"의원{i}", f"자료{i}", "작성중", "메일", ""]
                for i in range(1, 21)]
        self.make_excel(rows)
        svc = self.sync_service()
        svc.sync_excel_to_db()
        ledger = self.ledger_service()

        import extractors.ledger_parser as ledger_parser
        real_parse = ledger_parser.load_request_ledger
        started = threading.Event()

        def slow_parse(*args, **kwargs):
            result = real_parse(*args, **kwargs)
            started.set()
            time.sleep(0.8)          # 대형 워크북 파싱 시간을 흉내
            return result

        created = {}

        def do_insert():
            started.wait(5)
            created["res"] = ledger.insert_ledger_item({
                "year": "2026", "requester": "신규의원", "title": "동시 등록된 신규 자료",
                "request_type": "메일", "status": "작성중",
            })

        with mock.patch.object(ledger_parser, "load_request_ledger", side_effect=slow_parse):
            svc._id_stamp_fingerprint = None
            t_sync = threading.Thread(target=svc.sync_excel_to_db)
            t_insert = threading.Thread(target=do_insert)
            t_sync.start()
            t_insert.start()
            t_sync.join(30)
            t_insert.join(30)

        self.assertTrue(created["res"]["success"])
        new_id = created["res"]["item"]["ledger_id"]
        self.assertNotEqual(
            self.db_rows()[new_id]["status"], "[삭제]",
            "동기화 중에 등록한 항목이 같은 실행에서 삭제되면 안 된다",
        )

    def test_queue_flush_and_web_write_do_not_overwrite_each_other(self):
        """watcher의 flush와 HTTP 쓰기가 같은 워크북을 동시에 덮어쓰면 안 된다."""
        excel = self.make_excel([["REQ-2026-001", "1", "김의원", "자료A", "작성중", "일반", ""]])
        svc = self.sync_service()
        with svc._pending_lock:
            svc._enqueue_pending({
                "ledger_id": "REQ-2026-900", "year": "2026", "seq_no": "90",
                "requester": "보류의원", "title": "보류로 들어온 자료",
            }, "insert")

        # load~save 사이에 지연을 넣어 실제 대형 워크북의 인터리브를 재현한다.
        real_load = openpyxl.load_workbook
        # A barrier inside the serialized critical section can never see both
        # threads: it made a correct implementation take multiple 15s timeouts.
        def slow_load(*args, **kwargs):
            wb = real_load(*args, **kwargs)
            time.sleep(0.1)
            return wb

        import services.excel_sync_service as ess
        with mock.patch.object(ess.openpyxl, "load_workbook", side_effect=slow_load):
            t_flush = threading.Thread(target=svc.flush_pending_queue)
            t_web = threading.Thread(target=lambda: svc.sync_item_to_excel({
                "ledger_id": "REQ-2026-002", "year": "2026", "seq_no": "2",
                "requester": "웹의원", "title": "웹에서 바로 등록",
            }, action="insert"))
            t_flush.start()
            t_web.start()
            t_flush.join(30)
            t_web.join(30)
            self.assertFalse(t_flush.is_alive())
            self.assertFalse(t_web.is_alive())

        ids = {r["대장ID"] for r in self.sheet_rows(excel, "2026")}
        self.assertIn("REQ-2026-900", ids, "보류 큐에서 반영한 등록이 덮어써지면 안 된다")
        self.assertIn("REQ-2026-002", ids, "웹에서 바로 한 등록이 덮어써지면 안 된다")

    def test_workbook_writes_are_serialized_by_a_single_lock(self):
        """두 쓰기 경로가 동시에 _apply_to_excel에 들어가지 않아야 한다."""
        self.make_excel([["REQ-2026-001", "1", "김의원", "자료A", "작성중", "일반", ""]])
        svc = self.sync_service()
        with svc._pending_lock:
            svc._enqueue_pending({"ledger_id": "REQ-2026-900", "year": "2026",
                                  "requester": "보류의원", "title": "보류 자료"}, "insert")

        inside = []
        max_concurrent = []
        real_apply = svc._apply_to_excel
        guard = threading.Lock()

        def traced(path, item, action):
            with guard:
                inside.append(1)
                max_concurrent.append(len(inside))
            try:
                time.sleep(0.2)
                return real_apply(path, item, action)
            finally:
                with guard:
                    inside.pop()

        svc._apply_to_excel = traced
        t1 = threading.Thread(target=svc.flush_pending_queue)
        t2 = threading.Thread(target=lambda: svc.sync_item_to_excel(
            {"ledger_id": "REQ-2026-002", "year": "2026", "requester": "웹의원", "title": "웹 자료"},
            action="insert"))
        t1.start()
        t2.start()
        t1.join(30)
        t2.join(30)

        self.assertTrue(max_concurrent, "_apply_to_excel이 한 번도 호출되지 않았다")
        self.assertEqual(max(max_concurrent), 1,
                         "워크북 쓰기가 직렬화되지 않으면 lost update가 발생한다")


# ---------------------------------------------------------------------------
# 보류 큐 / 삭제 안전장치 / 입력 검증
# ---------------------------------------------------------------------------
class TestQueueAndGuards(LedgerFixtureMixin, unittest.TestCase):

    def test_permanently_failing_operation_is_quarantined(self):
        """영구 실패 작업이 큐 선두에서 뒤의 정상 작업을 영원히 막으면 안 된다."""
        excel = self.make_excel([["REQ-2026-001", "1", "김의원", "자료A", "작성중", "일반", ""]])
        svc = self.sync_service()
        with svc._pending_lock:
            svc._enqueue_pending({"ledger_id": "REQ-2026-BAD", "year": "2026",
                                  "requester": "실패의원", "title": "영구 실패"}, "insert")
            svc._enqueue_pending({"ledger_id": "REQ-2026-002", "year": "2026",
                                  "requester": "정상의원", "title": "뒤에 밀린 정상 작업"}, "insert")

        real_apply = svc._apply_to_excel

        def apply_but_fail_bad(path, item, action):
            if item.get("ledger_id") == "REQ-2026-BAD":
                return False
            return real_apply(path, item, action)

        svc._apply_to_excel = apply_but_fail_bad
        for _ in range(MAX_PENDING_ATTEMPTS + 2):
            svc.flush_pending_queue()

        self.assertEqual(svc.failed_count(), 1, "한도를 넘긴 작업은 격리돼야 한다")
        ids = {r["대장ID"] for r in self.sheet_rows(excel, "2026")}
        self.assertIn("REQ-2026-002", ids, "뒤의 정상 작업이 결국 반영돼야 한다")
        self.assertEqual(svc.pending_count(), 0)

    def test_small_ledger_is_protected_from_bulk_tombstoning(self):
        """10건 이하 대장도 파싱 급감 시 일괄 삭제 방어를 받아야 한다."""
        excel = self.make_excel([
            ["REQ-2026-001", "1", "김의원", "자료A", "작성중", "메일", ""],
            ["REQ-2026-002", "2", "이의원", "자료B", "작성중", "메일", ""],
            ["REQ-2026-003", "3", "박의원", "자료C", "작성중", "메일", ""],
            ["REQ-2026-004", "4", "최의원", "자료D", "작성중", "메일", ""],
        ])
        svc = self.sync_service()
        svc.sync_excel_to_db()
        self.assertEqual(len(self.db_rows()), 4)

        # 엑셀이 손상돼 1건만 파싱된 상황
        wb = openpyxl.load_workbook(excel)
        ws = wb["2026"]
        ws.delete_rows(3, 3)
        wb.save(excel)
        wb.close()
        svc._id_stamp_fingerprint = None
        svc.sync_excel_to_db()

        active = {lid for lid, r in self.db_rows().items() if r["status"] != "[삭제]"}
        self.assertEqual(len(active), 4, "소규모 대장도 일괄 삭제로부터 보호돼야 한다")

    def test_custom_ledger_id_must_match_the_expected_format(self):
        """자동 채번이 REQ-YYYY-NNN을 전제하므로 형식을 강제한다."""
        self.make_excel([["REQ-2026-001", "1", "김의원", "자료A", "작성중", "일반", ""]])
        ledger = self.ledger_service()
        bad = ledger.insert_ledger_item({"ledger_id": "임의번호-1", "year": "2026",
                                         "requester": "김의원", "title": "형식 위반"})
        self.assertFalse(bad["success"])
        self.assertIn("REQ-", bad["error"])

        good = ledger.insert_ledger_item({"ledger_id": "REQ-2026-050", "year": "2026",
                                          "requester": "김의원", "title": "형식 준수"})
        self.assertTrue(good["success"])

    def test_master_excel_is_backed_up_before_overwrite(self):
        """마스터 엑셀을 덮어쓰기 전에 백업본이 남아야 한다."""
        self.make_excel([["REQ-2026-001", "1", "김의원", "자료A", "작성중", "일반", ""]])
        # 서버 시작 순서대로 엑셀을 먼저 DB에 반영한다. 반영 전에 등록하면 DB가 엑셀의
        # REQ-2026-001을 몰라 같은 ID를 발급하고, 이제는 충돌로 거부된다(감사 R3-02).
        self.sync_service().sync_excel_to_db()
        ledger = self.ledger_service()
        res = ledger.insert_ledger_item({"year": "2026", "requester": "신규의원", "title": "백업 확인용"})
        self.assertTrue(res["success"])

        backups = list((self.tmp / ".ledger_backup").glob("*.xlsx"))
        self.assertTrue(backups, "원본을 덮어쓰기 전 백업이 있어야 한다")
        wb = openpyxl.load_workbook(str(backups[0]))
        try:
            self.assertIn("2026", wb.sheetnames, "백업본이 정상적으로 열려야 한다")
        finally:
            wb.close()


# ---------------------------------------------------------------------------
# ISSUE-006 : tombstone 이 파생 산출물로 새지 않는다
# ---------------------------------------------------------------------------
class TestTombstoneIsolation(LedgerFixtureMixin, unittest.TestCase):
    """[T-006] 삭제한 항목은 오프라인 대시보드·엑셀 산출물에 다시 나타나지 않는다."""

    def test_pipeline_json_excludes_tombstones_but_db_keeps_them(self):
        self.make_excel([
            ["REQ-2026-001", "1", "김의원", "자료A", "작성중", "일반", ""],
            ["REQ-2026-002", "2", "이의원", "자료B", "작성중", "일반", ""],
        ])
        svc = self.sync_service()
        svc.sync_excel_to_db()
        self.assertTrue(self.ledger_service().delete_ledger_item("REQ-2026-001")["success"])

        import extract_and_build_db as builder
        with mock.patch.object(builder, "SYSTEM_DIR", self.tmp), \
             mock.patch.object(builder, "ROOT_DIR", self.tmp), \
             mock.patch.object(builder, "DB_PATH", self.db_path), \
             mock.patch.object(builder, "JSON_PATH", self.json_path), \
             mock.patch.object(builder, "PARSED_DIR", self.tmp / "_parsed_markdown"), \
             mock.patch.object(builder, "DIST_MANIFEST", self.tmp / ".distribution_manifest.json"):
            result = builder.main()

        self.assertTrue(result["success"])
        payload = json.loads(self.json_path.read_text(encoding="utf-8"))
        ids = {r["ledger_id"] for r in payload["request_ledger"]}
        self.assertNotIn("REQ-2026-001", ids, "삭제 항목이 검색 산출물에 되살아나면 안 된다")
        self.assertIn("REQ-2026-002", ids)

        self.assertEqual(self.db_rows()["REQ-2026-001"]["status"], "[삭제]",
                         "tombstone 자체는 DB에 남아야 stale 엑셀이 항목을 되살리지 못한다")

    def test_slim_payload_drops_tombstones_as_a_safety_net(self):
        from template_renderer import DashboardRenderer
        payload = DashboardRenderer.build_slim_payload({
            "documents": [], "qa_items": [],
            "request_ledger": [
                {"ledger_id": "REQ-2026-001", "status": "[삭제]"},
                {"ledger_id": "REQ-2026-002", "status": "작성중"},
            ],
        })
        self.assertEqual([r["ledger_id"] for r in payload["request_ledger"]], ["REQ-2026-002"])


# ---------------------------------------------------------------------------
# ISSUE-005 : GUI 완료 핸들러
# ---------------------------------------------------------------------------
class TestLauncherGuiCompletionHandler(unittest.TestCase):
    """[T-005] 소스 문자열 grep이 아니라 실제 스코프를 검증한다."""

    def _on_finish_symbol(self, name):
        src = (SCRIPTS_DIR / "launcher_gui.py").read_text(encoding="utf-8")
        table = symtable.symtable(src, "launcher_gui.py", "exec")

        def walk(t):
            for child in t.get_children():
                if child.get_name() == "on_finish":
                    return child
                found = walk(child)
                if found:
                    return found
            return None

        on_finish = walk(table)
        self.assertIsNotNone(on_finish, "run_smart_update의 on_finish를 찾지 못했다")
        assert on_finish is not None
        return on_finish.lookup(name)

    def test_err_msg_is_bound_to_the_enclosing_scope(self):
        """nonlocal이 없으면 첫 참조에서 UnboundLocalError가 나고 결과 안내가 사라진다."""
        symbol = self._on_finish_symbol("err_msg")
        self.assertFalse(
            symbol.is_local() and not symbol.is_nonlocal(),
            "on_finish의 err_msg가 지역 변수로 잡히면 UnboundLocalError가 난다. nonlocal 선언이 필요하다.",
        )

    def test_completion_handler_pattern_executes(self):
        """같은 구조를 실제로 실행해 예외 없이 오류 경로를 타는지 확인한다."""
        seen = []

        def worker(result_data):
            err_msg = None

            def on_finish():
                nonlocal err_msg
                if not err_msg and isinstance(result_data, dict) and not result_data.get("success", True):
                    err_msg = result_data.get("error") or "파이프라인 단계 실패"
                seen.append(err_msg)

            on_finish()

        worker({"success": False, "error": "parse_all 실패"})
        worker({"success": True})
        self.assertEqual(seen, ["parse_all 실패", None])

    def test_source_declares_nonlocal_in_completion_handler(self):
        src = (SCRIPTS_DIR / "launcher_gui.py").read_text(encoding="utf-8")
        self.assertIn("nonlocal err_msg", src)


if __name__ == "__main__":
    unittest.main()
