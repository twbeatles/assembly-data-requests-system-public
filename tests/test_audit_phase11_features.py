# -*- coding: utf-8 -*-
"""검색·DB·UI/UX 개선 회귀 테스트 (감사 3회차 후속).

- 대시보드: 진행상태·마감 필터, 탭별 필터 적용 범위, 연도 칩 자동 생성, 마감 정렬, 빈 결과 안내,
  엑셀 상태 배너 문구, 최근 검색어 (Node로 실행)
- 서버: 대장 마감 판정(대시보드와 같은 규칙), 상태·마감 필터, 요약, 페이지네이션, 스키마 확인 캐시
- tools/purge_test_data.py: 테스트 데이터 판별과 정리 계획(실제 수정 없음)
"""

import datetime
import importlib.util
import json
import shutil
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

TESTS_DIR = Path(__file__).resolve().parent
SYSTEM_DIR = TESTS_DIR.parent
SCRIPTS_DIR = SYSTEM_DIR / "scripts"
for p in (str(SCRIPTS_DIR), str(TESTS_DIR)):
    if p not in sys.path:
        sys.path.insert(0, p)

from db.database_manager import DatabaseManager
from services import ledger_service as ls_module
from services.ledger_service import LedgerService, ledger_due_state
import web_server
from test_audit_phase10_dashboard import NODE, TEMPLATE, extract_js, run_node, SEARCH_FUNCS

TODAY = datetime.date(2026, 9, 14)
LEDGER_FIXTURE = [
    {"ledger_id": "REQ-2026-001", "year": "2026", "status": "작성중", "deadline": "2026-09-10", "title": "지난 마감", "requester": "김의원"},
    {"ledger_id": "REQ-2026-002", "year": "2026", "status": "작성중", "deadline": "2026-09-14", "title": "오늘 마감", "requester": "이의원"},
    {"ledger_id": "REQ-2026-003", "year": "2026", "status": "검토중", "deadline": "2026-09-20", "title": "6일 뒤", "requester": "박의원"},
    {"ledger_id": "REQ-2026-004", "year": "2026", "status": "작성중", "deadline": "2026-10-30", "title": "먼 마감", "requester": "최의원"},
    {"ledger_id": "REQ-2026-005", "year": "2026", "status": "제출", "deadline": "2026-09-01", "title": "제출 완료", "requester": "정의원"},
    {"ledger_id": "REQ-2027-001", "year": "2027", "status": "작성중", "deadline": "", "title": "마감 미정", "requester": "한의원"},
]
EXPECTED_STATES = {"REQ-2026-001": "overdue", "REQ-2026-002": "soon", "REQ-2026-003": "soon",
                   "REQ-2026-004": "later", "REQ-2026-005": "done", "REQ-2027-001": "none"}

UI_FUNCS = SEARCH_FUNCS + [
    "LEDGER_DONE_STATUSES", "DUE_SOON_DAYS", "parseLedgerDate", "startOfDay", "ledgerDueInfo",
    "matchesDueFilter", "dueBadgeHtml", "compareByDeadline", "collectFilterYears", "collectLedgerStatuses",
    "countLedgerDue", "emptyStateHints", "excelStatusMessage", "nextRecentSearches", "escapeHtml",
]


def node_body(code: str) -> str:
    return (
        "let currentMode = 'ledger';\n"
        f"const LEDGER = {json.dumps(LEDGER_FIXTURE, ensure_ascii=False)};\n"
        "const TODAY = new Date(2026, 8, 14);\n" + code
    )


@unittest.skipUnless(NODE, "node가 없어 대시보드 JS 실행 테스트를 건너뜁니다")
class TestDashboardLedgerUx(unittest.TestCase):
    def test_due_state_badge_and_counts(self):
        res = run_node(UI_FUNCS, node_body("""
const states = {}; const badges = {};
for (const it of LEDGER) { states[it.ledger_id] = ledgerDueInfo(it, TODAY).state; badges[it.ledger_id] = dueBadgeHtml(it, TODAY); }
console.log(JSON.stringify({ states, badges, counts: countLedgerDue(LEDGER, TODAY) }));
"""))
        self.assertEqual(res["states"], EXPECTED_STATES)
        self.assertIn("마감 4일 지남", res["badges"]["REQ-2026-001"])
        self.assertIn("오늘 마감", res["badges"]["REQ-2026-002"])
        self.assertIn("D-6", res["badges"]["REQ-2026-003"])
        self.assertEqual(res["badges"]["REQ-2026-005"], "")
        self.assertEqual(res["counts"], {"soon": 2, "overdue": 1, "open": 5})

    def test_ledger_filters_ignore_document_only_filters(self):
        """관리대장 탭에서 기관·주제·통계표 필터가 켜져 있어도 대장 결과가 0건이 되지 않는다."""
        res = run_node(UI_FUNCS, node_body("""
activeFilters.inst = '국회'; activeFilters.topic = '딥페이크'; activeFilters.tablesOnly = true;
const out = {};
out.all = filterItemsFromList(LEDGER, false).length;
activeFilters.status = '작성중'; out.status = filterItemsFromList(LEDGER, false).map(x => x.ledger_id).sort();
activeFilters.status = 'ALL'; activeFilters.due = 'open'; out.open = filterItemsFromList(LEDGER, false).length;
activeFilters.due = 'ALL'; activeFilters.year = '2027'; out.year = filterItemsFromList(LEDGER, false).map(x => x.ledger_id);
console.log(JSON.stringify(out));
"""))
        self.assertEqual(res["all"], 6)
        self.assertEqual(res["status"], ["REQ-2026-001", "REQ-2026-002", "REQ-2026-004", "REQ-2027-001"])
        self.assertEqual(res["open"], 5)
        self.assertEqual(res["year"], ["REQ-2027-001"])

    def test_deadline_sort_puts_overdue_first_and_done_last(self):
        res = run_node(UI_FUNCS, node_body("""
const sorted = LEDGER.slice().sort((a, b) => compareByDeadline(a, b, TODAY)).map(x => x.ledger_id);
console.log(JSON.stringify(sorted));
"""))
        self.assertEqual(res, ["REQ-2026-001", "REQ-2026-002", "REQ-2026-003", "REQ-2026-004", "REQ-2027-001", "REQ-2026-005"])

    def test_dynamic_year_and_status_chips(self):
        res = run_node(UI_FUNCS, node_body("""
const docs = [{ year: '2018' }, { year: '2026' }, { year: '' }];
const ledger = LEDGER.concat([{ ledger_id: 'X', year: '2026', status: '[삭제]' }]);
console.log(JSON.stringify({ years: collectFilterYears(docs, ledger), statuses: collectLedgerStatuses(ledger) }));
"""))
        self.assertEqual(res["years"], ["2027", "2026", "2018"])
        self.assertEqual(dict(res["statuses"]), {"작성중": 4, "검토중": 1, "제출": 1})

    def test_empty_state_hints_explain_active_filters(self):
        res = run_node(UI_FUNCS, node_body("""
currentMode = 'docs';
activeFilters.query = '딥페이크 텔레그램'; activeFilters.inst = '국회';
const a = emptyStateHints(0, 10);
activeFilters.query = '-텔레그램'; activeFilters.inst = 'ALL';
const b = emptyStateHints(0, 10);
const c = emptyStateHints(0, 0);
console.log(JSON.stringify({ a, b, c }));
"""))
        self.assertIn('검색어 "딥페이크 텔레그램"', res["a"]["active"])
        self.assertIn("기관 국회", res["a"]["active"])
        self.assertTrue(any("OR" in h for h in res["a"]["hints"]))
        self.assertTrue(any("제외어" in h for h in res["b"]["hints"]))
        self.assertTrue(any("적재된 문서가 없습니다" in h for h in res["c"]["hints"]))

    def test_excel_status_banner_and_recent_searches(self):
        res = run_node(UI_FUNCS, node_body("""
console.log(JSON.stringify({
  none: excelStatusMessage({ excel_pending: 0, excel_failed: 0 }),
  pending: excelStatusMessage({ excel_pending: 2, excel_failed: 0 }),
  failed: excelStatusMessage({ excel_pending: 1, excel_failed: 3 }),
  recent: nextRecentSearches(['a1', '텔레그램', 'b2', 'c3', 'd4', 'e5', 'f6', 'g7'], '텔레그램'),
  short: nextRecentSearches(['x1'], 'a'),
  broken: nextRecentSearches('not-a-list', '딥페이크'),
  esc: escapeHtml(2026),
}));
"""))
        self.assertIsNone(res["none"])
        self.assertEqual(res["pending"]["level"], "warn")
        self.assertEqual(res["failed"]["level"], "error")
        self.assertIn("반영하지 못한 변경 3건", res["failed"]["text"])
        self.assertEqual(res["recent"][0], "텔레그램")
        self.assertEqual(len(res["recent"]), 8)
        self.assertEqual(res["short"], ["x1"])
        self.assertEqual(res["broken"], ["딥페이크"])
        self.assertEqual(res["esc"], "2026")

    def test_template_has_single_status_filter_and_no_hardcoded_years(self):
        src = TEMPLATE.read_text(encoding="utf-8")
        self.assertEqual(src.count('id="filter-group-status"'), 1)
        self.assertNotIn('data-filter="year" data-val="2018"', src)
        self.assertIn('id="empty-state"', src)
        self.assertIn('id="server-excel-banner"', src)
        # 대장 표의 원문보기 링크는 linked_doc_id를 이스케이프해야 한다(웹 입력으로 넣을 수 있는 값).
        self.assertNotIn("DOCS_MAP['${item.linked_doc_id}']", src)
        self.assertNotIn("${item.deadline || '-'}", src)


@unittest.skipUnless(NODE, "node가 없어 대시보드 JS 실행 테스트를 건너뜁니다")
class TestDashboardHeaderConfig(unittest.TestCase):
    """오프라인(01) 대시보드 제목에 부서명이 두 번 나오고 부제에 "&amp;"가 보이던 문제."""

    def test_rendered_header_uses_raw_config_values(self):
        from template_renderer import DashboardRenderer
        r = DashboardRenderer()
        cfg = {"department_name": "기획예산팀", "agency_name": "공공기관",
               "system_title": "자료요구 검색", "system_subtitle": "전문 & 협업</script>"}
        html = r.render(r.load_template(), "", r.compress_payload({"documents": [], "qa_items": [], "request_ledger": []}), cfg)
        self.assertNotIn("/* __APP_CONFIG_JSON__ */", html)
        body = extract_js(html, "EMBEDDED_APP_CONFIG") + "\n" + extract_js(html, "appConfig") + """
const cleanTitle = (appConfig.system_title || '').replace(/^🛡️\\s*/, '');
console.log(JSON.stringify({ h1: `${appConfig.department_name} ${cleanTitle}`, sub: `${appConfig.agency_name} ${appConfig.department_name} | ${appConfig.system_subtitle}` }));
"""
        res = run_node([], body)
        self.assertEqual(res["h1"], "기획예산팀 자료요구 검색")
        self.assertEqual(res["sub"], "공공기관 기획예산팀 | 전문 & 협업</script>")
        self.assertNotIn("전문 & 협업</script>", html, "설정값이 스크립트 블록을 끊으면 안 된다")

    def test_universal_builder_uses_shared_root_detection(self):
        import system_config
        import build_universal_package
        self.assertEqual(build_universal_package.WORKSPACE_ROOT, system_config.find_workspace_root(build_universal_package.SYSTEM_DIR))


class TestServerLedgerFeatures(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="audit_phase11_"))
        self.db = self.tmp / "data_requests.db"
        conn = sqlite3.connect(str(self.db))
        DatabaseManager.init_schema(conn)
        for it in LEDGER_FIXTURE:
            conn.execute(
                "INSERT INTO request_ledger (ledger_id, year, status, deadline, title, requester, created_at, updated_at) "
                "VALUES (?, ?, ?, ?, ?, ?, '', '')",
                (it["ledger_id"], it["year"], it["status"], it["deadline"], it["title"], it["requester"]),
            )
        conn.execute("INSERT INTO request_ledger (ledger_id, year, status, title, requester) VALUES "
                     "('REQ-2026-099', '2026', '[삭제]', '삭제됨', '의원')")
        conn.commit()
        conn.close()
        self.svc = LedgerService(self.db, self.tmp / "data_requests.json", self.tmp)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_due_state_matches_dashboard_rule(self):
        states = {it["ledger_id"]: ledger_due_state(it, TODAY)[0] for it in LEDGER_FIXTURE}
        self.assertEqual(states, EXPECTED_STATES)

    def test_query_ledger_status_and_due_filters(self):
        ids = lambda rows: sorted(r["ledger_id"] for r in rows)
        self.assertEqual(ids(self.svc.query_ledger(due="overdue", today=TODAY)), ["REQ-2026-001"])
        self.assertEqual(ids(self.svc.query_ledger(due="soon", today=TODAY)), ["REQ-2026-002", "REQ-2026-003"])
        self.assertEqual(ids(self.svc.query_ledger(status="검토중")), ["REQ-2026-003"])
        self.assertEqual(len(self.svc.query_ledger(due="bogus")), 6, "알 수 없는 due 값은 무시한다")

    def test_ledger_summary(self):
        summary = self.svc.ledger_summary(today=TODAY)
        self.assertEqual(summary["total"], 6, "tombstone은 세지 않는다")
        self.assertEqual(summary["due"], {"overdue": 1, "soon": 2, "open": 5})
        self.assertEqual(summary["by_status"]["작성중"], 4)

    def test_paged_response(self):
        rows = [{"i": n} for n in range(25)]
        body = web_server.paged_response(rows, {"limit": ["10"], "offset": ["20"]})
        self.assertEqual((body["count"], body["total"], body["offset"], body["limit"]), (5, 25, 20, 10))
        self.assertEqual(web_server.paged_response(rows, {})["count"], 25)
        self.assertEqual(web_server.parse_page_params({"limit": ["99999"], "offset": ["-5"]}), (1000, 0))
        self.assertEqual(web_server.parse_page_params({"limit": ["abc"]}), (None, 0))

    def test_schema_check_runs_once_per_db_file(self):
        calls = []
        real = DatabaseManager.ensure_schema

        def counting(conn):
            calls.append(1)
            return real(conn)

        ls_module._SCHEMA_READY.clear()
        with mock.patch.object(DatabaseManager, "ensure_schema", side_effect=counting):
            self.svc.query_ledger()
            self.svc.query_documents(kw="자료")
            self.svc.get_stats()
            self.assertEqual(len(calls), 1, "같은 DB 파일에는 한 번만 DDL을 실행해야 한다")
            # 파이프라인이 DB를 새 파일로 교체한 상황 (복사본은 새 파일 식별자를 가진다)
            replacement = self.tmp / "new.db"
            shutil.copy2(self.db, replacement)
            import os
            os.replace(replacement, self.db)
            self.svc.query_ledger()
            self.assertEqual(len(calls), 2, "교체된 DB 파일은 다시 확인해야 한다")


def load_purge_tool():
    spec = importlib.util.spec_from_file_location("purge_test_data", SYSTEM_DIR / "tools" / "purge_test_data.py")
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class TestPurgeToolPlanning(unittest.TestCase):
    """운영 정리 도구의 판별·계획 로직. 실제 엑셀(COM)·DB 수정 경로는 실행하지 않는다."""

    def setUp(self):
        self.tool = load_purge_tool()
        self.tmp = Path(tempfile.mkdtemp(prefix="audit_phase11_purge_"))
        self.db = self.tmp / "data_requests.db"
        conn = sqlite3.connect(str(self.db))
        DatabaseManager.init_schema(conn)
        rows = [
            ("REQ-2025-001", "2025", "실의원", "실제 요구자료"),
            ("REQ-2025-002", "2025", "테스트의원", "테스트 제목"),
            ("REQ-2026-001", "2026", "홍길동 의원", "추가 자동 연계 테스트 자료"),
            ("REQ-2026-002", "2026", "홍길동 의원", "딥페이크 테스트 결과 보고"),  # '테스트'가 들어간 실제 제목
        ]
        for lid, year, req, title in rows:
            conn.execute("INSERT INTO request_ledger (ledger_id, year, requester, title, status) VALUES (?, ?, ?, ?, '작성중')",
                         (lid, year, req, title))
        conn.execute("INSERT INTO ledger_history (ledger_id, action, changed_at, snapshot) VALUES ('REQ-2026-009', 'EXCEL_APPLIED', '', ?)",
                     (json.dumps({"requester": "단위테스트의원", "title": "단위테스트 요구자료 제목"}, ensure_ascii=False),))
        conn.execute("INSERT INTO ledger_history (ledger_id, action, changed_at, snapshot) VALUES ('REQ-2025-001', 'UPDATE', '', '{}')")
        conn.execute("INSERT INTO documents (doc_id, linked_ledger_id) VALUES ('DOC-2026-001', 'REQ-2026-001')")
        conn.commit()
        conn.close()
        self.excel = self.tmp / "대장.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "2025"
        ws.append(["연번", "의원", "요구자료", "대장ID"])
        ws.append(["1", "실의원", "실제 요구자료", "REQ-2025-001"])
        ws.append(["001", "테스트의원", "테스트 제목", "REQ-2025-002"])
        ws2 = wb.create_sheet("2026")
        ws2.append(["연번", "의원", "요구자료", "대장ID"])
        ws2.append(["4", "홍길동 의원", "딥페이크 테스트 결과 보고", "REQ-2026-002"])
        ws2.append(["5", "새의원", "새 수기 행", None])
        ws3 = wb.create_sheet("통합목록(검색용)")
        ws3.append(["No.", "연도", "의원", "요구자료"])
        ws3.append([1, 2025, "실의원", "실제 요구자료"])
        ws3.append([2, 2025, "테스트의원", "테스트 제목"])
        wb.save(self.excel)
        wb.close()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_signature_match_is_exact(self):
        self.assertTrue(self.tool.is_test_pair("테스트의원", "테스트 제목"))
        self.assertFalse(self.tool.is_test_pair("홍길동 의원", "딥페이크 테스트 결과 보고"))
        self.assertFalse(self.tool.is_test_pair("실의원", "테스트 제목"))

    def test_database_plan(self):
        plan = self.tool.plan_database(self.db)
        self.assertEqual(plan["purge_ids"], ["REQ-2025-002", "REQ-2026-001"])
        self.assertEqual(len(plan["history_ids"]), 1, "테스트 스냅숏의 EXCEL_APPLIED 이력만 대상")
        self.assertEqual(plan["unlink_docs"], ["DOC-2026-001"])

    def test_workbook_plan_deletes_test_rows_and_stamps_missing_ids(self):
        db_plan = self.tool.plan_database(self.db)
        plan = self.tool.plan_workbook(self.excel, db_plan["remaining_records"])
        self.assertEqual(plan["delete_rows"], {"2025": [3], "통합목록(검색용)": [3]})
        self.assertTrue(plan["renumber"])
        self.assertEqual(len(plan["id_assignments"]), 1)
        sheet, row, lid = plan["id_assignments"][0]
        self.assertEqual((sheet, row), ("2026", 3))
        self.assertEqual(lid, "REQ-2026-003", "명시 ID가 있는 시트에서는 기존·삭제 ID를 새 행에 주지 않는다")

    def test_dry_run_does_not_modify_files(self):
        before = (self.db.stat().st_mtime_ns, self.excel.stat().st_mtime_ns)
        db_plan = self.tool.plan_database(self.db)
        self.tool.plan_workbook(self.excel, db_plan["remaining_records"])
        self.assertEqual((self.db.stat().st_mtime_ns, self.excel.stat().st_mtime_ns), before)


if __name__ == "__main__":
    unittest.main()
