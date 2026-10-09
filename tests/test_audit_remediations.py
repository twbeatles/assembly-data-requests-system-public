# -*- coding: utf-8 -*-
"""Regression coverage for the functional-audit remediations."""

import json
import shutil
import sqlite3
import sys
import tempfile
import threading
import unittest
from pathlib import Path

import openpyxl

SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import parse_all
from db.database_manager import DatabaseManager
from extractors.ledger_parser import load_request_ledger
from services.excel_sync_service import ExcelSyncService
from services.ledger_service import LedgerService


class TestAuditRemediations(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="audit_remediation_"))
        self.db = self.tmp / "data_requests.db"
        self.json_path = self.tmp / "data_requests.json"
        conn = sqlite3.connect(str(self.db))
        DatabaseManager.init_schema(conn)
        conn.close()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _ledger_excel(self):
        path = self.tmp / "260904 국회 요구자료 목록(테스트).xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "2026"
        ws.append(["대장ID", "연번", "의원", "요구자료", "제출", "요청형태"])
        ws.append(["REQ-2026-001", "1", "김의원", "삭제 재발 방지", "작성중", "일반"])
        wb.save(path)
        wb.close()
        return path

    def test_distribution_preserves_source_less_markdown(self):
        output = self.tmp / "_parsed_markdown" / "2026"
        output.mkdir(parents=True)
        orphan = output / "only-in-package.pdf.md"
        orphan.write_text("# 전문 sentinel", encoding="utf-8")
        self.assertEqual(parse_all.prune_orphan_markdowns(self.tmp, output.parent, preserve_missing_sources=True), 0)
        self.assertTrue(orphan.exists())
        self.assertEqual(parse_all.prune_orphan_markdowns(self.tmp, output.parent, preserve_missing_sources=False), 1)
        self.assertFalse(orphan.exists())

    def test_deleted_excel_tombstone_does_not_resurrect(self):
        self._ledger_excel()
        ledger = LedgerService(self.db, self.json_path, self.tmp)
        inserted = ledger.insert_ledger_item({"ledger_id": "REQ-2026-001", "year": "2026", "requester": "김의원", "title": "삭제 재발 방지", "request_type": "일반"})
        self.assertTrue(inserted["success"])
        self.assertTrue(ledger.delete_ledger_item("REQ-2026-001")["success"])

        sync = ExcelSyncService(self.tmp, self.db, self.json_path)
        result = sync.sync_excel_to_db()
        self.assertTrue(result["success"])
        self.assertIsNone(ledger.get_ledger_item("REQ-2026-001"))
        conn = sqlite3.connect(str(self.db))
        self.assertEqual(conn.execute("SELECT status FROM request_ledger WHERE ledger_id = ?", ("REQ-2026-001",)).fetchone()[0], "[삭제]")
        conn.close()

    def test_queue_failure_preserves_failed_operation_and_tail(self):
        excel = self.tmp / "master.xlsx"
        excel.touch()
        sync = ExcelSyncService(self.tmp, self.db, self.json_path)
        sync.find_master_excel = lambda: excel
        sync.is_file_locked = lambda _: False
        sync._pending_queue = [
            {"operation_id": str(i), "action": "insert", "item": {"ledger_id": f"REQ-{i}"}, "timestamp": 0, "attempts": 0}
            for i in range(3)
        ]
        sync._apply_to_excel = lambda *_: (_ for _ in ()).throw(RuntimeError("locked mid-flush"))
        self.assertEqual(sync.flush_pending_queue(), 0)
        self.assertEqual([op["operation_id"] for op in sync._pending_queue], ["0", "1", "2"])
        persisted = json.loads((self.tmp / ".excel_pending_queue.json").read_text(encoding="utf-8"))
        self.assertEqual([op["operation_id"] for op in persisted], ["0", "1", "2"])

    def test_excel_insert_replay_is_idempotent(self):
        path = self._ledger_excel()
        sync = ExcelSyncService(self.tmp, self.db, self.json_path)
        item = {"ledger_id": "REQ-2026-002", "year": "2026", "seq_no": "2", "requester": "박의원", "title": "재시도 안전 등록", "status": "작성중"}
        self.assertTrue(sync._apply_to_excel(path, item, "insert"))
        self.assertTrue(sync._apply_to_excel(path, item, "insert"))
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["2026"]
        ids = [ws.cell(row, 1).value for row in range(2, ws.max_row + 1)]
        wb.close()
        self.assertEqual(ids.count("REQ-2026-002"), 1)

    def test_concurrent_auto_ids_are_unique(self):
        service = LedgerService(self.db, self.json_path, self.tmp)
        gate = threading.Barrier(12)
        results = []
        result_lock = threading.Lock()

        def worker(n):
            gate.wait()
            res = service.insert_ledger_item({"year": "2026", "requester": f"의원{n}", "title": f"동시 등록 {n}"})
            with result_lock:
                results.append(res)

        threads = [threading.Thread(target=worker, args=(i,)) for i in range(12)]
        for thread in threads: thread.start()
        for thread in threads: thread.join()
        self.assertTrue(all(result["success"] for result in results), results)
        ids = [result["item"]["ledger_id"] for result in results]
        self.assertEqual(len(ids), len(set(ids)))

    def test_dynamic_year_sheets_include_future_year(self):
        path = self.tmp / "260904 국회 요구자료 목록(연도).xlsx"
        wb = openpyxl.Workbook()
        for index, year in enumerate(("2026", "2027")):
            ws = wb.active if index == 0 else wb.create_sheet()
            assert ws is not None
            ws.title = year
            ws.append(["대장ID", "연번", "의원", "요구자료"])
            ws.append([f"REQ-{year}-001", "1", "김의원", f"{year} 자료"])
        wb.save(path)
        wb.close()
        ids = {item["ledger_id"] for item in load_request_ledger(self.tmp, [], merge_db_records=False)}
        self.assertEqual(ids, {"REQ-2026-001", "REQ-2027-001"})


if __name__ == "__main__":
    unittest.main()
